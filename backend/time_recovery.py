"""Personal Excel recovery: reviewed, idempotent and never auto-approved."""
import base64
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime, time, timedelta, timezone
from types import SimpleNamespace
from fastapi import Body, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


def install_recovery(app, namespace):
    s=SimpleNamespace(**namespace)
    def digest(value):
        return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,default=str).encode()).hexdigest()
    def check_actor(actor):
        s.operational(actor)
        with s.connection() as c:
            manager=s.owner_manager(c,actor)
            if not manager or s.get(c,'users',manager)['status']!='ACTIVE':s.fail('Necesitas un Manager activo asignado para recuperar horas.')
    def fingerprint(c,actor):
        return digest([dict(actor),s.owner_manager(c,actor),*[[r for r in s.rows(c,t) if r.get('user_id')==actor['id'] or t=='ot_time_corrections' and r['overtime_id'] in {o['id'] for o in s.rows(c,'overtime') if o['user_id']==actor['id']}] for t in ('overtime','guards','timers','ot_time_corrections','time_recovery_items')]])
    def moment(value):
        try:return s.parse_at(str(value)).astimezone(s.TZ)
        except Exception:s.fail('Fecha u horario inválido en el respaldo.')
    def clean_parts(parts,allow_future=False):
        if not isinstance(parts,list) or not parts or len(parts)>128:s.fail('El registro requiere entre 1 y 128 tramos.')
        cleaned=[];previous=None
        for part in parts:
            if not isinstance(part,dict):s.fail('Cada tramo debe contener inicio y fin.')
            a=moment(part.get('start_at'));b=moment(part.get('end_at'))
            if b<=a or (not allow_future and b>datetime.now(timezone.utc)) or (b-a)>timedelta(days=32) or previous and a<previous:s.fail('Los horarios deben ser pasados, positivos, ordenados y sin traslapes.')
            cleaned.append({'start_at':a.isoformat(),'end_at':b.isoformat()});previous=b
        return cleaned
    def clean_item(item,actor):
        if str(item.get('employee_id',''))!=str(actor['employee_id']):s.fail('El Excel contiene horas de otro Employee ID. Recupera únicamente tus propias horas.')
        kind=item.get('kind')
        if kind not in ('OT','GUARD'):s.fail('Tipo de registro no compatible.')
        if kind=='GUARD' and actor['is_lead']:s.fail('El Lead sólo puede recuperar OT, no Guardias.')
        parts=clean_parts(item.get('segments'),allow_future=kind=='GUARD')
        result={'kind':kind,'segments':parts,'source_id':str(item.get('source_id',''))[:200], 'origin':str(item.get('origin',''))[:500]}
        if kind=='OT':
            for key in ('reference','description','justification','impact','in_lieu'):
                value=str(item.get(key) or '').strip()
                if len(value)>10000:s.fail('Un campo de documentación supera el tamaño permitido.')
                result[key]=value
        else:
            # Normalize fixed coverage into a snapshot; never reinterpret net hours as gross rule coverage.
            by_day={}
            for part in parts:
                a,b=moment(part['start_at']),moment(part['end_at'])
                while a<b:
                    end=min(b,datetime.combine(a.date()+timedelta(days=1),time(),s.TZ))
                    if a.second or a.microsecond or end.second or end.microsecond:s.fail('La cobertura de Guardia debe tener precisión de minutos.')
                    by_day.setdefault(a.date(),[]).append({'day':a.weekday(),'start':a.strftime('%H:%M'),'end':end.strftime('%H:%M')});a=end
            first,last=min(by_day),max(by_day)
            if (last-first).days>31:s.fail('La Guardia no puede exceder 32 días.')
            # A fixed recovery can span repeated weekdays with different coverage: store exact intervals.
            result.update(start_date=str(first),end_date=str(last),name=str(item.get('name') or 'Guardia recuperada')[:200])
        result['source_key']=digest([actor['employee_id'],kind,result['source_id']]) if result['source_id'] else digest([actor['employee_id'],kind,parts,result.get('reference','')])
        return result
    def intervals(c,actor):
        found={'OT':[],'GUARD':[]}
        for o in s.rows(c,'overtime'):
            if o['user_id']!=actor['id'] or o['status']=='CANCELLED':continue
            parts=s.effective_ot_segments(c,o) if o['execution']=='COMPLETED' else ([{'start_at':o['scheduled_start'],'end_at':o['scheduled_end']}] if o['scheduled_start'] and o['scheduled_end'] else [])
            if not parts and o['execution'] in ('RUNNING','PAUSED') and o['start_at']:
                parts=[{'start_at':o['start_at'],'end_at':datetime.now(timezone.utc).isoformat()}]
            found['OT'].extend((moment(p['start_at']),moment(p['end_at'])) for p in parts)
        for g in s.rows(c,'guards'):
            if g['user_id']==actor['id'] and g['status']!='CANCELLED':found['GUARD'].extend((moment(p['start_at']),moment(p['end_at'])) for p in guard_parts(g))
        return found
    def guard_parts(record):
        return [{'start_at':a.isoformat(),'end_at':b.isoformat()} for a,b in s.guard_coverage(record)]
    def diagnose(c,actor,items):
        existing=intervals(c,actor);seen=set();result=[]
        imported={r['source_key'] for r in s.rows(c,'time_recovery_items') if r['user_id']==actor['id']}
        for n,item in enumerate(items,1):
            parts=[(moment(p['start_at']),moment(p['end_at'])) for p in item['segments']]
            duplicate=item['source_key'] in imported or item['source_key'] in seen or all(not s.subtract_intervals(a,b,existing[item['kind']]) for a,b in parts)
            overlap=not duplicate and any(a<y and b>x for a,b in parts for x,y in existing[item['kind']])
            status='DUPLICATE' if duplicate else 'CONFLICT' if overlap else 'READY'
            documentation=item['kind']=='OT' and (not all(item.get(k) for k in ('reference','description','justification','impact','in_lieu')) or not re.fullmatch(r'(CHG|INC)[A-Z0-9-]+',item['reference'].upper()))
            result.append({'index':n,'kind':item['kind'],'reference':item.get('reference') or item.get('name',''),'start_at':item['segments'][0]['start_at'],'end_at':item['segments'][-1]['end_at'],'minutes':sum((b-a).total_seconds()/60 for a,b in parts),'status':status,'needs_documentation':documentation})
            if status=='READY':existing[item['kind']].extend(parts)
            seen.add(item['source_key'])
        return result
    def read_file(data,actor):
        try:
            encoded=data.get('content','')
            if not isinstance(encoded,str) or len(encoded)>8_000_000:raise ValueError()
            raw=base64.b64decode(encoded,validate=True)
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist())>200 or sum(x.file_size for x in archive.infolist())>25_000_000:raise ValueError()
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        except Exception:s.fail('Selecciona un respaldo .pulse o un Excel .xlsx válido (máximo 6 MB).')
        try:
            items=[];warnings=[]
            if 'Respaldo Pulse' in book.sheetnames:
                sheet=book['Respaldo Pulse']
                if sheet.max_row>2001:s.fail('El respaldo supera 2000 registros.')
                records=list(sheet.iter_rows(min_row=1,max_row=2003,values_only=False))
                if not records or [c.value for c in records[0]][:2]!=['Formato','Registro JSON']:s.fail('Encabezado del respaldo no compatible.')
                for row in records[1:]:
                    if all(c.value is None for c in row):continue
                    if row[0].value!='PULSE-TIME-1' or any(c.data_type=='f' for c in row):s.fail('Versión o contenido del respaldo no compatible.')
                    try:item=json.loads(''.join(str(c.value or '') for c in row[1:]))
                    except Exception:s.fail('Hay un registro ilegible en el respaldo.')
                    if not isinstance(item,dict):s.fail('Registro de respaldo inválido.')
                    items.append(clean_item(item,actor))
            elif 'Tiempo aprobado' in book.sheetnames:
                warnings.append('Excel anterior: sólo contiene horas aprobadas calculadas. No incluye la regla original ni toda la documentación de OT. La Guardia se recupera como cobertura neta fija por semana; no se inventa su cobertura original.')
                if book['Tiempo aprobado'].max_row>2001:s.fail('El Excel supera 2000 filas de tiempo.')
                values=book['Tiempo aprobado'].iter_rows(min_row=1,max_row=2003);header=[c.value for c in next(values)]
                required=['Fecha','Employee ID','Tipo','Inicio','Fin','Duración HH:MM','Referencia','Descripción']
                if any(k not in header for k in required):s.fail('El Excel no tiene las columnas de tiempo de Pulse.')
                if 'Resumen' in book.sheetnames:
                    meta=dict(tuple(r[:2]) for r in book['Resumen'].iter_rows(values_only=True) if len(r)>=2 and r[0])
                    if meta.get('Zona horaria',str(s.TZ))!=str(s.TZ):s.fail('El Excel usa otra zona horaria; revisa sus horarios antes de recuperarlo.')
                guards={}
                for row in values:
                    if all(c.value is None for c in row):continue
                    if any(c.data_type=='f' for c in row):s.fail('El Excel no debe contener fórmulas en sus registros.')
                    record={header[i]:c.value for i,c in enumerate(row) if i<len(header)}
                    try:
                        day=record['Fecha'];day=day.date() if isinstance(day,datetime) else day if isinstance(day,date) else date.fromisoformat(str(day))
                        clock=lambda value:value if isinstance(value,time) else time.fromisoformat(str(value))
                        a=datetime.combine(day,clock(record['Inicio']),s.TZ);b=datetime.combine(day,clock(record['Fin']),s.TZ)
                        if b<=a:b+=timedelta(days=1)
                        duration=str(record['Duración HH:MM']);h,m=map(int,duration.split(':'))
                        if m not in range(60) or abs((b-a).total_seconds()/60-(h*60+m))>1:raise ValueError()
                        kind={'Stand-by':'GUARD','Onsite Overtime':'OT'}[record['Tipo']]
                    except Exception:s.fail('Hay fechas, tipos o duraciones inconsistentes en el Excel.')
                    candidate={'employee_id':str(record['Employee ID']),'kind':kind,'segments':[{'start_at':a.isoformat(),'end_at':b.isoformat()}],'reference':record.get('Referencia'),'description':record.get('Descripción'),'origin':'Excel de tiempo aprobado anterior'}
                    if kind=='GUARD':
                        if str(candidate['employee_id'])!=str(actor['employee_id']):s.fail('El Excel contiene horas de otro Employee ID.')
                        week=day-timedelta(days=(day.weekday()+2)%7);guards.setdefault(str(week),[]).extend(candidate['segments'])
                    else:items.append(clean_item(candidate,actor))
                for week,parts in guards.items():
                    # Repeated or overlapping source rows are rejected instead of silently adding time.
                    parts.sort(key=lambda x:x['start_at']);items.append(clean_item({'employee_id':actor['employee_id'],'kind':'GUARD','name':'Guardia recuperada · '+week,'segments':parts,'origin':'Excel anterior: cobertura neta fija'},actor))
            else:s.fail('Este archivo no es un respaldo ni un Excel de tiempo aprobado de Pulse.')
            if not items:s.fail('El Excel no contiene horas para recuperar.')
            if len(items)>2000:s.fail('Recupera hasta 2000 registros por archivo.')
            view_month=None
            if 'Resumen' in book.sheetnames:
                meta={r[0]:r[1] for r in book['Resumen'].iter_rows(max_row=100,values_only=True) if len(r)>=2 and r[0]}
                candidate=str(meta.get('Mes',''))
                try:
                    if len(candidate)==7:date.fromisoformat(candidate+'-01');view_month=candidate
                except ValueError:pass
            if view_month:items[0]['view_month']=view_month
            return items,warnings
        finally:book.close()

    def append_native(wb,c,actor,month):
        if 'Respaldo Pulse' in wb.sheetnames:del wb['Respaldo Pulse']
        sheet=wb.create_sheet('Respaldo Pulse');sheet.append(['Formato','Registro JSON','Continuación 1','Continuación 2'])
        lower=datetime.combine(date.fromisoformat(month+'-01'),time(),s.TZ);upper=datetime.combine((lower.date().replace(day=28)+timedelta(days=4)).replace(day=1),time(),s.TZ)
        for table,kind in [('overtime','OT'),('guards','GUARD')]:
            for record in s.rows(c,table):
                if record['user_id']!=actor['id'] or record['status']=='CANCELLED' or kind=='OT' and record['execution']!='COMPLETED':continue
                parts=s.effective_ot_segments(c,record) if kind=='OT' else guard_parts(record)
                if not any(moment(p['start_at'])<upper and moment(p['end_at'])>lower for p in parts):continue
                item={'employee_id':actor['employee_id'],'kind':kind,'source_id':record['id'],'segments':parts,'origin':'Pulse '+s.VERSION+' · '+month}
                if kind=='OT':item.update({k:record[k] for k in ('reference','description','justification','impact','in_lieu')})
                else:item['name']=json.loads(record['rule_snapshot'])['name']
                value=json.dumps(item,ensure_ascii=False)
                if len(value)>90000:s.fail('Un registro supera el tamaño de respaldo Excel; conserva también la base de datos.')
                sheet.append(['PULSE-TIME-1',*[value[i:i+30000] for i in range(0,len(value),30000)]])
        if sheet.max_row>2001:s.fail('Hay más de 2000 registros: divide el respaldo en periodos menores.')
        sheet.freeze_panes='A2'
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=60
        if 'Resumen' in wb.sheetnames:
            wb['Resumen'].append(['Respaldo de recuperación','Horas finalizadas completas que intersectan el mes; incluye pendientes y tramos fuera del mes. OT futura y Timers activos no incluidos; las Guardias pueden incluir cobertura futura. Recuperación con nueva revisión.'])

    def append_capture(wb,c,detail):
        from openpyxl.worksheet.table import Table,TableStyleInfo
        header=['Mes/Año','Empleado','Tipo','Semana','Fecha dd/mm/aaaa','Hora_Inicio','Horas_ADP','Horas_ADP_decimal','HRS_SAP_Decimal','Hora_Fin','Ticket CHG','Ticket INC','Descripción','ID Emp','Employee Name','Employee email','1 L Manager email','Total Hours','Start Date\nmm/dd/yyyy','End Date\nmm/dd/yyyy','Week Number','Client Name','Category','Overtime Description','Justification - Reason for Overtime','Impact if it is not approved','Why in-lieu time off cannot be considered','Were the hours approved by the DIL?\nY/N','# Change/Incident']
        page=wb.create_sheet('Captura',0);page.append(header)
        people={p['id']:p for p in s.rows(c,'users')};groups={g['id']:g for g in s.rows(c,'groups')};accounts={a['id']:a for a in s.rows(c,'accounts')};ots={o['id']:o for o in s.rows(c,'overtime')}
        safe=lambda v:"'"+str(v) if str(v).startswith(('=','+','-','@')) else v
        for row in detail:
            person=people[row['user_id']];manager=people.get(s.owner_manager(c,person),{});group=groups.get(person['group_id'],{});account=accounts.get(group.get('account_id') or manager.get('account_id'),{})
            a,b=moment(row['start_at']),moment(row['end_at']);day=a.date();first=day.replace(day=1);offset=(first.weekday()+2)%7;week=(day.day-1+offset)//7+1
            record=ots.get(row['id'],{}) if row['kind']=='OT' else {};reference=row['reference'] or '';description=row['description'] or ''
            values=[day.strftime('%m/%Y'),row['name'],'Overtime' if row['kind']=='OT' else 'Standby','Semana '+str(week),day,a.time().replace(tzinfo=None),row['minutes']/1440,row['minutes']/60,row['minutes']/60,'24:00' if b.date()>a.date() and b.time()==time() else b.time().replace(tzinfo=None),reference if reference.upper().startswith('CHG') else '',reference if reference.upper().startswith('INC') else '',description,person['employee_id'],row['name'],person['email'],manager.get('email',''),row['minutes']/60,day,b.date(),week,account.get('name',''),'',description,record.get('justification',''),record.get('impact',''),record.get('in_lieu',''),'',reference]
            page.append([safe(v) if isinstance(v,str) else v for v in values])
            n=page.max_row
            for column,fmt in [(5,'dd/mm/yyyy'),(6,'hh:mm'),(7,'[h]:mm'),(8,'0.00'),(9,'0.00'),(10,'hh:mm'),(18,'0.00'),(19,'mm/dd/yyyy'),(20,'mm/dd/yyyy')]:page.cell(n,column).number_format=fmt
        for cell in page[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173E48');cell.alignment=Alignment(wrap_text=True,vertical='center')
        page.row_dimensions[1].height=44;page.freeze_panes='F2';page.auto_filter.ref=page.dimensions
        for col in page.columns:page.column_dimensions[col[0].column_letter].width=22
        for name in ['M','X','Y','Z','AA']:page.column_dimensions[name].width=45
        if detail:
            table=Table(displayName='ResultadoMensual',ref=page.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);page.add_table(table)
        if 'Resumen' in wb.sheetnames:wb['Resumen'].append(['Campos externos','Category y autorización DIL no se infieren: quedan vacíos. Semana es semana del mes con inicio sábado. Sólo horas aprobadas en Captura.'])

    @app.get('/api/time-backup.xlsx', include_in_schema=False)
    @app.get('/api/time-backup.pulse')
    def download(month:str,actor=Depends(s.current_user)):
        s.operational(actor);s.analytics_months(month)
        wb=Workbook();wb.remove(wb.active)
        with s.connection() as c:
            detail=wb.create_sheet('Tiempo aprobado');detail.append(['Fecha','Employee ID','Técnico','Grupo','Tipo','Inicio','Fin','Duración HH:MM','Horas decimales','Referencia','Descripción'])
            for r in s.accounting(c,{actor['id']},month):
                safe=lambda x:"'"+str(x) if str(x).startswith(('=','+','-','@')) else x
                n=round(r['minutes']);detail.append([r['date'],safe(r['employee_id']),safe(r['name']),'','Onsite Overtime' if r['kind']=='OT' else 'Stand-by',r['start'],r['end'],f'{n//60:02}:{n%60:02}',r['hours'],safe(r['reference']),safe(r['description'])])
        summary=wb.create_sheet('Resumen');summary.append(['Mes',month]);summary.append(['Zona horaria',str(s.TZ)]);summary.append(['Employee ID',actor['employee_id']]);summary.append(['Versión',s.VERSION]);summary.append(['Alcance','Horas finalizadas que intersectan el mes, incluidos tramos completos fuera del mes; OT futura y Timers activos excluidos; cobertura de Guardia futura permitida.']);summary.append(['Recuperación','Nunca autoaprueba. OT incompleta requiere documentación; lo demás vuelve a revisión del Manager. No sustituye el respaldo completo de la base de datos.'])
        with s.connection() as c:
            append_capture(wb,c,s.accounting(c,{actor['id']},month));append_native(wb,c,actor,month)
        for page in wb:
            page.freeze_panes='A2'
            for cell in page[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173E48')
            if page.title not in ('Captura','Respaldo Pulse'):
                for col in page.columns:page.column_dimensions[col[0].column_letter].width=28
        out=io.BytesIO();wb.save(out);out.seek(0)
        return StreamingResponse(out,media_type='application/vnd.operational-pulse.backup',headers={'Content-Disposition':f'attachment; filename="Operational_Pulse_Respaldo_{month}.pulse"'})

    @app.post('/api/time-recovery/preview')
    def preview(data:dict=Body(...),actor=Depends(s.current_user)):
        check_actor(actor);items,warnings=read_file(data,actor)
        with s.connection(True) as c:
            diagnostics=diagnose(c,actor,items);identity=s.uid();stamp=datetime.now(timezone.utc)
            c.execute('DELETE FROM time_recovery_drafts WHERE expires_at<?',(stamp.isoformat(),))
            c.execute('INSERT INTO time_recovery_drafts VALUES(?,?,?,?,?,?)',(identity,actor['id'],json.dumps(items),fingerprint(c,actor),(stamp+timedelta(minutes=30)).isoformat(),None))
        return {'id':identity,'rows':diagnostics,'warnings':warnings,'ready':sum(x['status']=='READY' for x in diagnostics),'duplicates':sum(x['status']=='DUPLICATE' for x in diagnostics),'conflicts':sum(x['status']=='CONFLICT' for x in diagnostics)}

    @app.post('/api/time-recovery/{identity}/confirm')
    def confirm(identity:str,data:dict=Body(...),actor=Depends(s.current_user)):
        check_actor(actor)
        if data.get('confirm') is not True:s.fail('Confirma que revisaste las horas y su reenvío a aprobación.')
        with s.connection(True) as c:
            draft=c.execute('SELECT * FROM time_recovery_drafts WHERE id=? AND user_id=?',(identity,actor['id'])).fetchone()
            if not draft:s.fail('La vista previa no existe para tu usuario.',404)
            if draft['result']:return json.loads(draft['result'])
            if datetime.fromisoformat(draft['expires_at'])<=datetime.now(timezone.utc):s.fail('La vista previa venció. Vuelve a revisar el archivo.',409)
            if draft['fingerprint']!=fingerprint(c,actor):s.fail('Tus registros o asignación cambiaron. Vuelve a generar la vista previa.',409)
            items=json.loads(draft['items']);diagnostics=diagnose(c,actor,items);created=[];documents=0
            for item,row in zip(items,diagnostics):
                if row['status']!='READY':continue
                rid=s.uid();stamp=s.now();kind=item['kind']
                if kind=='OT':
                    status='DOCUMENTATION' if row['needs_documentation'] else 'SUBMITTED';documents+=status=='DOCUMENTATION'
                    c.execute('INSERT INTO overtime(id,user_id,capture_type,start_at,end_at,reference,description,justification,impact,in_lieu,status,execution,revision,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(rid,actor['id'],'HISTORICAL',item['segments'][0]['start_at'],item['segments'][-1]['end_at'],item['reference'],item['description'],item['justification'],item['impact'],item['in_lieu'],status,'COMPLETED',0,stamp,stamp))
                    c.execute('INSERT INTO ot_time_corrections VALUES(?,?,?,?)',(rid,json.dumps(item['segments']),'Recuperación desde respaldo Excel; tramos originales del archivo.',stamp))
                else:
                    snapshot={'name':item['name'],'segments':[],'recovered_intervals':item['segments'],'recovery_fixed':True,'effective_from':item['start_date']}
                    c.execute('INSERT INTO guards VALUES(?,?,?,?,?,?,?,?)',(rid,actor['id'],item['start_date'],item['end_date'],json.dumps(snapshot),'SUBMITTED',None,stamp))
                c.execute('INSERT INTO time_recovery_items VALUES(?,?,?,?,?)',(actor['id'],item['source_key'],rid,identity,json.dumps({'kind':kind,'origin':item['origin'],'source_id':item['source_id']})))
                s.capture_origin(c,rid,actor);s.audit(c,actor,'RECOVER_TIME_FROM_EXCEL',rid,{'batch':identity,'source_id':item['source_id'],'origin':item['origin'],'approval_preserved':False});created.append(rid)
            result={'created':len(created),'documentation':documents,'submitted':len(created)-documents,'skipped':len(items)-len(created),'ids':created,'month':items[0].get('view_month') or min(x['segments'][0]['start_at'][:7] for x in items)}
            c.execute('UPDATE time_recovery_drafts SET result=? WHERE id=?',(json.dumps(result),identity));s.audit(c,actor,'RECOVER_TIME_BATCH',identity,result)
        return result

    return append_capture, append_native
