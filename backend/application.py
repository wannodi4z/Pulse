"""Operational Pulse: server-owned authorization, SQLite transactions and time accounting."""
from __future__ import annotations

import asyncio
import base64
import hashlib
import hmac
import io
import json
import os
import secrets
import re
import sqlite3
import uuid
import unicodedata
from contextlib import contextmanager, asynccontextmanager, suppress
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from fastapi import Body, Depends, FastAPI, HTTPException, Request, Response
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from cryptography.fernet import Fernet, InvalidToken
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

VERSION = '1.0.10'
ROOT = Path(__file__).resolve().parent
if ROOT.name == 'backend':
    ROOT = ROOT.parent
DB = Path(os.environ.get('PULSE_DB_PATH', ROOT / 'operational_pulse.db'))
RECOVERY_KEY = Path(os.environ.get('PULSE_RECOVERY_KEY_PATH', DB.with_suffix('.key')))
TZ = ZoneInfo(os.environ.get('PULSE_TIMEZONE', 'America/Mexico_City'))
COUNTRY_CODES = frozenset('AD AE AF AG AI AL AM AO AQ AR AS AT AU AW AX AZ BA BB BD BE BF BG BH BI BJ BL BM BN BO BQ BR BS BT BV BW BY BZ CA CC CD CF CG CH CI CK CL CM CN CO CR CU CV CW CX CY CZ DE DJ DK DM DO DZ EC EE EG EH ER ES ET FI FJ FK FM FO FR GA GB GD GE GF GG GH GI GL GM GN GP GQ GR GS GT GU GW GY HK HM HN HR HT HU ID IE IL IM IN IO IQ IR IS IT JE JM JO JP KE KG KH KI KM KN KP KR KW KY KZ LA LB LC LI LK LR LS LT LU LV LY MA MC MD ME MF MG MH MK ML MM MN MO MP MQ MR MS MT MU MV MW MX MY MZ NA NC NE NF NG NI NL NO NP NR NU NZ OM PA PE PF PG PH PK PL PM PN PR PS PT PW PY QA RE RO RS RU RW SA SB SC SD SE SG SH SI SJ SK SL SM SN SO SR SS ST SV SX SY SZ TC TD TF TG TH TJ TK TL TM TN TO TR TT TV TW TZ UA UG UM US UY UZ VA VC VE VG VI VN VU WF WS YE YT ZA ZM ZW'.split())
PLATFORMS = {'MF':'Mainframe','OPEN':'OPEN / Distribuidos','OTHER':'Otra plataforma'}
SERVICE_CATALOG = {'MF':['CICS-MQ','DB2','ZOS','IMS','STORAGE'], 'OPEN':['LINUX','UNIX','ORACLE','WINDOWS','SQL-SERVER','POSTGRESQL','MYSQL','DB2','NETWORK','CLOUD','STORAGE','SECURITY'], 'OTHER':['GENERAL']}

SESSION_SECONDS = 8 * 3600
@asynccontextmanager
async def lifespan(app):
    async def clock_loop():
        while True:
            await asyncio.to_thread(reconcile_planned_timers)
            await asyncio.sleep(1)
    task = asyncio.create_task(clock_loop())
    try:
        yield
    finally:
        task.cancel()
        with suppress(asyncio.CancelledError):
            await task


app = FastAPI(title='Operational Pulse', version=VERSION, lifespan=lifespan)


def now():
    return datetime.now(timezone.utc).isoformat(timespec='microseconds')


def uid():
    return uuid.uuid4().hex


def fail(message, status=422):
    raise HTTPException(status_code=status, detail=message)


def password_hash(value):
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac('sha256', value.encode(), salt, 260000)
    return base64.b64encode(salt + digest).decode()


def password_matches(value, stored):
    try:
        raw = base64.b64decode(stored)
        return hmac.compare_digest(raw[16:], hashlib.pbkdf2_hmac('sha256', value.encode(), raw[:16], 260000))
    except (ValueError, TypeError):
        return False


@contextmanager
def connection(write=False):
    c = sqlite3.connect(DB, timeout=20)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    try:
        if write:
            c.execute('BEGIN IMMEDIATE')
            reconcile=globals().get('pause_due_timers')
            if reconcile:
                reconcile(c,now())
        yield c
        c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()


def init_db():
    DB.parent.mkdir(parents=True, exist_ok=True)
    with connection(True) as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS accounts(id TEXT PRIMARY KEY, code TEXT UNIQUE COLLATE NOCASE NOT NULL, name TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS groups(id TEXT PRIMARY KEY, account_id TEXT NOT NULL REFERENCES accounts(id), manager_id TEXT REFERENCES users(id), code TEXT UNIQUE COLLATE NOCASE NOT NULL, name TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS retired_groups(group_id TEXT PRIMARY KEY REFERENCES groups(id), retired_at TEXT NOT NULL, reason TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY, role TEXT NOT NULL CHECK(role IN ('ADMIN','MANAGER','USER')), username TEXT UNIQUE COLLATE NOCASE NOT NULL, employee_id TEXT UNIQUE COLLATE NOCASE, first TEXT NOT NULL, last TEXT NOT NULL, email TEXT NOT NULL, account_id TEXT REFERENCES accounts(id), group_id TEXT REFERENCES groups(id), is_lead INTEGER NOT NULL DEFAULT 0, status TEXT NOT NULL, password_hash TEXT, activation_hash TEXT, created_at TEXT NOT NULL);
        CREATE UNIQUE INDEX IF NOT EXISTS one_lead ON users(group_id) WHERE is_lead=1;
        CREATE TABLE IF NOT EXISTS session_activity(token_hash TEXT PRIMARY KEY, last_activity TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS sessions(token_hash TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), expires_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS delegations(id TEXT PRIMARY KEY, manager_id TEXT NOT NULL REFERENCES users(id), lead_id TEXT NOT NULL REFERENCES users(id), group_ids TEXT NOT NULL, start_at TEXT NOT NULL, end_at TEXT NOT NULL, revoked_at TEXT);
        CREATE TABLE IF NOT EXISTS rules(id TEXT PRIMARY KEY, manager_id TEXT NOT NULL REFERENCES users(id), name TEXT NOT NULL, effective_from TEXT NOT NULL, segments TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS windows(id TEXT PRIMARY KEY, group_id TEXT NOT NULL REFERENCES groups(id), created_by TEXT NOT NULL REFERENCES users(id), scheduled_start TEXT NOT NULL, scheduled_end TEXT NOT NULL, reference TEXT NOT NULL, description TEXT NOT NULL, status TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS overtime(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), window_id TEXT REFERENCES windows(id), capture_type TEXT NOT NULL, scheduled_start TEXT, scheduled_end TEXT, start_at TEXT, end_at TEXT, reference TEXT NOT NULL DEFAULT '', description TEXT NOT NULL DEFAULT '', justification TEXT NOT NULL DEFAULT '', impact TEXT NOT NULL DEFAULT '', in_lieu TEXT NOT NULL DEFAULT '', status TEXT NOT NULL, execution TEXT NOT NULL, revision INTEGER NOT NULL DEFAULT 0, approved_by TEXT REFERENCES users(id), created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE UNIQUE INDEX IF NOT EXISTS unique_actual_ot ON overtime(user_id,start_at,end_at) WHERE start_at IS NOT NULL AND end_at IS NOT NULL;
        CREATE UNIQUE INDEX IF NOT EXISTS unique_participation ON overtime(user_id,window_id) WHERE window_id IS NOT NULL;
        CREATE TABLE IF NOT EXISTS timers(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), overtime_id TEXT NOT NULL REFERENCES overtime(id), start_at TEXT NOT NULL, end_at TEXT);
        CREATE TABLE IF NOT EXISTS ot_extensions(id TEXT PRIMARY KEY, overtime_id TEXT NOT NULL REFERENCES overtime(id), requested_by TEXT NOT NULL REFERENCES users(id), previous_end TEXT NOT NULL, requested_end TEXT NOT NULL, reason TEXT NOT NULL, status TEXT NOT NULL, decided_by TEXT REFERENCES users(id), decision_reason TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, decided_at TEXT);
        CREATE UNIQUE INDEX IF NOT EXISTS one_pending_extension ON ot_extensions(overtime_id) WHERE status='SUBMITTED';
        CREATE UNIQUE INDEX IF NOT EXISTS one_active_timer ON timers(user_id) WHERE end_at IS NULL;
        CREATE TABLE IF NOT EXISTS guards(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), start_date TEXT NOT NULL, end_date TEXT NOT NULL, rule_snapshot TEXT NOT NULL, status TEXT NOT NULL, approved_by TEXT REFERENCES users(id), created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS recoveries(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), status TEXT NOT NULL, code_hash TEXT, expires_at TEXT, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS recovery_code_vault(recovery_id TEXT PRIMARY KEY REFERENCES recoveries(id), encrypted_code BLOB NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS audit(id TEXT PRIMARY KEY, actor_id TEXT REFERENCES users(id), action TEXT NOT NULL, entity_id TEXT NOT NULL, detail TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS transfer_requests(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), requested_by TEXT NOT NULL REFERENCES users(id), source_group_id TEXT NOT NULL REFERENCES groups(id), destination TEXT NOT NULL, reason TEXT NOT NULL, status TEXT NOT NULL, created_at TEXT NOT NULL, decided_at TEXT, decided_by TEXT REFERENCES users(id), decision_reason TEXT NOT NULL DEFAULT '', destination_group_id TEXT REFERENCES groups(id));
        CREATE UNIQUE INDEX IF NOT EXISTS one_pending_transfer ON transfer_requests(user_id) WHERE status='SUBMITTED';
        CREATE TABLE IF NOT EXISTS manager_transfers(id TEXT PRIMARY KEY,user_id TEXT NOT NULL REFERENCES users(id),source_manager_id TEXT NOT NULL REFERENCES users(id),target_manager_id TEXT NOT NULL REFERENCES users(id),source_group_id TEXT NOT NULL REFERENCES groups(id),reason TEXT NOT NULL,status TEXT NOT NULL,created_at TEXT NOT NULL,decided_at TEXT,decision_reason TEXT NOT NULL DEFAULT '',destination_group_id TEXT REFERENCES groups(id));
        CREATE TABLE IF NOT EXISTS ot_group_activities(source_id TEXT PRIMARY KEY REFERENCES overtime(id), window_id TEXT UNIQUE NOT NULL REFERENCES windows(id), authorized_by TEXT REFERENCES users(id), snapshot TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS ot_time_corrections(overtime_id TEXT PRIMARY KEY REFERENCES overtime(id), segments TEXT NOT NULL, reason TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS time_recovery_drafts(id TEXT PRIMARY KEY,user_id TEXT NOT NULL REFERENCES users(id),items TEXT NOT NULL,fingerprint TEXT NOT NULL,expires_at TEXT NOT NULL,result TEXT);
        CREATE TABLE IF NOT EXISTS time_recovery_items(user_id TEXT NOT NULL REFERENCES users(id),source_key TEXT NOT NULL,record_id TEXT NOT NULL,batch_id TEXT NOT NULL,metadata TEXT NOT NULL,PRIMARY KEY(user_id,source_key));
        CREATE TABLE IF NOT EXISTS record_origins(record_id TEXT PRIMARY KEY, snapshot TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS login_attempts(identity TEXT PRIMARY KEY, failures INTEGER NOT NULL, locked_until TEXT);
        ''')
        # Add metadata without renaming legacy keys or changing hierarchy/records.
        for table,columns in {'accounts':{'country_code':"TEXT NOT NULL DEFAULT ''",'base_code':"TEXT NOT NULL DEFAULT ''"},'users':{'manager_id':'TEXT REFERENCES users(id)'},'groups':{'platform_code':"TEXT NOT NULL DEFAULT ''",'service_key':"TEXT NOT NULL DEFAULT ''",'standard_code':"TEXT NOT NULL DEFAULT ''"}}.items():
            present={row['name'] for row in c.execute('PRAGMA table_info('+table+')')}
            for column,definition in columns.items():
                if column not in present:c.execute('ALTER TABLE '+table+' ADD COLUMN '+column+' '+definition)
        c.execute("UPDATE accounts SET base_code=code WHERE base_code=''")
        c.execute("CREATE UNIQUE INDEX IF NOT EXISTS country_account_key ON accounts(country_code,base_code COLLATE NOCASE) WHERE country_code!='' AND base_code!=''")
        c.execute("CREATE UNIQUE INDEX IF NOT EXISTS global_group_key ON groups(standard_code COLLATE NOCASE) WHERE standard_code!=''")
        c.execute("CREATE TABLE IF NOT EXISTS specialties(id TEXT PRIMARY KEY,account_id TEXT NOT NULL REFERENCES accounts(id),key TEXT NOT NULL,name TEXT NOT NULL,UNIQUE(account_id,key))")
        c.execute('DROP INDEX IF EXISTS one_direct_lead')
        # Upgrade the previous Group Lead into the Manager's direct backup once.
        for lead in c.execute("SELECT u.*,g.manager_id AS owner FROM users u JOIN groups g ON g.id=u.group_id WHERE u.is_lead=1 AND g.manager_id IS NOT NULL").fetchall():
            for table in ('overtime','guards'):
                for record in c.execute('SELECT id FROM '+table+' WHERE user_id=?',(lead['id'],)).fetchall():
                    g=dict(c.execute('SELECT * FROM groups WHERE id=?',(lead['group_id'],)).fetchone())
                    a=dict(c.execute('SELECT * FROM accounts WHERE id=?',(g['account_id'],)).fetchone())
                    c.execute('INSERT OR IGNORE INTO record_origins VALUES(?,?)',(record['id'],json.dumps({'group_id':g['id'],'group_name':g['name'],'group_code':g['code'],'account_name':a['name'],'account_code':a['code'],'manager_id':lead['owner'],'inferred':True})))
            c.execute('UPDATE users SET manager_id=?,group_id=NULL WHERE id=?',(lead['owner'],lead['id']))

        # Adopt manual revocation for the latest currently effective legacy backup.
        # Expired, revoked and future assignments are never reactivated.
        stamp = now()
        legacy = c.execute("SELECT d.* FROM delegations d JOIN users m ON m.id=d.manager_id WHERE d.revoked_at IS NULL AND d.start_at<=? AND (d.end_at='' OR d.end_at>?) AND m.status='ACTIVE' ORDER BY d.rowid DESC", (stamp, stamp)).fetchall()
        migrated = set()
        for d in legacy:
            if d['manager_id'] in migrated:
                continue
            migrated.add(d['manager_id'])
            if d['end_at']:
                c.execute("UPDATE delegations SET end_at='' WHERE id=?", (d['id'],))
                c.execute('INSERT INTO audit VALUES(?,?,?,?,?,?)', (uid(), None, 'DELEGATION_MANUAL_REVOCATION', d['id'], json.dumps({'previous_end_at':d['end_at']}), stamp))
        # Only a genuinely new database receives bootstrap administrators. Existing data is never reset.
        if c.execute('SELECT count(*) FROM users').fetchone()[0] == 0 and c.execute('SELECT count(*) FROM audit').fetchone()[0] == 0:
            for n, surname in [(1, 'Principal'), (2, 'Respaldo')]:
                c.execute('INSERT INTO users(id,role,username,first,last,email,status,password_hash,created_at) VALUES(?,?,?,?,?,?,?,?,?)',
                          (uid(), 'ADMIN', f'admin{n}', 'Administrador', surname, f'admin{n}@operationalpulse.local', 'ACTIVE', password_hash(f'admin{n}2026'), now()))


def install_example_database():
    """The test distribution initializes only a missing default DB, never an existing one."""
    example=ROOT/'fixtures'/'banorte_pruebas.sqlite3'
    if 'PULSE_DB_PATH' in os.environ or DB.exists() or not example.is_file():return
    import shutil
    temporary=DB.with_name('.pulse-example-'+uuid.uuid4().hex+'.sqlite3')
    try:
        shutil.copyfile(example,temporary)
        try:os.link(temporary,DB)  # Exclusive and atomic: an existing database always wins.
        except FileExistsError:pass
    finally:temporary.unlink(missing_ok=True)


install_example_database()
init_db()


def account_display_code(account):
    return account['country_code']+'-'+(account.get('base_code') or account['code']) if account.get('country_code') else account['code']


def group_display_code(group, account):
    if group.get('standard_code'):return group['standard_code']
    if not account.get('country_code'):return group['code']
    prefix=account_display_code(account)
    if group['code'].startswith(prefix+'-'):return group['code']
    return account['country_code']+'-'+group['code'] if group['code'].startswith(account['base_code']+'-') else prefix+'-'+group['code']


def rows(c, table):
    result=[dict(r) for r in c.execute(f'SELECT * FROM {table}')]
    if table=='accounts':
        for a in result:a['display_code']=account_display_code(a)
    elif table=='groups':
        retired={r[0] for r in c.execute('SELECT group_id FROM retired_groups')}
        result=[g for g in result if g['id'] not in retired]
        accounts={a['id']:dict(a) for a in c.execute('SELECT * FROM accounts')}
        for g in result:
            a=accounts[g['account_id']];g['country_code']=a['country_code'];g['display_code']=group_display_code(g,a)
    return result


def get(c, table, identity):
    row = c.execute(f'SELECT * FROM {table} WHERE id=?', (identity,)).fetchone()
    if row is None:
        fail('No se encontró el registro.', 404)
    return dict(row)


def audit(c, actor, action, entity, detail=''):
    c.execute('INSERT INTO audit VALUES(?,?,?,?,?,?)', (uid(), actor['id'] if actor else None, action, entity, json.dumps(detail, ensure_ascii=False), now()))


def public_user(u):
    return {k: v for k, v in dict(u).items() if k not in ('password_hash', 'activation_hash')}


def current_user(request: Request):
    token = request.cookies.get('pulse_session', '')
    hashed = hashlib.sha256(token.encode()).hexdigest()
    with connection() as c:
        activity=c.execute('SELECT last_activity FROM session_activity WHERE token_hash=?',(hashed,)).fetchone()
        if not activity or (datetime.now(timezone.utc)-parse_at(activity['last_activity'])).total_seconds()>=300:
            fail('Tu sesión terminó por 5 minutos de inactividad. Inicia sesión nuevamente.',401)
        row = c.execute('SELECT u.* FROM users u JOIN sessions s ON s.user_id=u.id WHERE s.token_hash=? AND s.expires_at>? AND u.status=?', (hashed, now(), 'ACTIVE')).fetchone()
        if row is None:
            fail('Inicia sesión para continuar.', 401)
        return dict(row)


@app.post('/api/session/activity')
def session_activity(request: Request, actor=Depends(current_user)):
    hashed=hashlib.sha256(request.cookies.get('pulse_session','').encode()).hexdigest()
    with connection(True) as c:
        c.execute('INSERT OR REPLACE INTO session_activity VALUES(?,?)',(hashed,datetime.now(timezone.utc).isoformat()))
    return {'ok':True}


def admin(user=Depends(current_user)):
    if user['role'] != 'ADMIN':
        fail('Esta acción requiere una cuenta administradora.', 403)
    return user


def effective_delegations(c):
    """One effective backup per Manager, including legacy databases with duplicate rows.

    Preserve history and select the most recently assigned active record; do not
    delete or reseed records while reading a database from an older version.
    """
    stamp = now()
    selected = {}
    for row in c.execute("SELECT d.* FROM delegations d JOIN users m ON m.id=d.manager_id WHERE d.revoked_at IS NULL AND d.start_at<=? AND (d.end_at='' OR d.end_at>?) AND m.status='ACTIVE' ORDER BY d.rowid DESC", (stamp, stamp)):
        selected.setdefault(row['manager_id'], dict(row))
    return list(selected.values())


def scope(c, user):
    if user['role'] == 'ADMIN':
        return {g['id'] for g in rows(c, 'groups')}
    if user['role'] == 'MANAGER':
        return {r[0] for r in c.execute('SELECT id FROM groups WHERE manager_id=?', (user['id'],))}
    ids = set()
    if user['is_lead']:
        for d in effective_delegations(c):
            if d['lead_id'] != user['id']:
                continue
            allowed = {r[0] for r in c.execute('SELECT id FROM groups WHERE manager_id=?', (d['manager_id'],))}
            if user['manager_id'] == d['manager_id']:
                ids.update(allowed)
    return ids


def operational(user):
    if user['role'] != 'USER':
        fail('Esta función es para Técnicos.', 403)


def owner_manager(c,user):
    if user.get('is_lead') and user.get('manager_id'):return user['manager_id']
    row=c.execute('SELECT manager_id FROM groups WHERE id=?',(user.get('group_id'),)).fetchone()
    return row[0] if row else None


def visible_users(c,actor):
    managers={actor['id']} if actor['role']=='MANAGER' else {d['manager_id'] for d in effective_delegations(c) if d['lead_id']==actor['id'] and actor['is_lead'] and actor['manager_id']==d['manager_id']}
    return {p['id'] for p in rows(c,'users') if actor['role']=='ADMIN' or p['id']==actor['id'] or p['role']=='USER' and owner_manager(c,p) in managers}


def authorized_target(c, actor, target, approval=False):
    if actor['role'] == 'ADMIN' and not approval:
        return
    if target['id'] not in visible_users(c,actor) or (approval and target['id'] == actor['id']) or (approval and actor['role']=='ADMIN'):
        fail('No tienes autoridad sobre este registro.', 403)


def text(data, key, required=True):
    value = str(data.get(key, '')).strip()
    if required and not value:
        fail('Completa los campos obligatorios.')
    if len(value) > 4000:
        fail('El texto es demasiado largo.')
    return value


def parse_at(value):
    try:
        result = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if result.tzinfo is None:
            result = result.replace(tzinfo=TZ)
        return result.astimezone(timezone.utc)
    except (ValueError, TypeError):
        fail('La fecha y hora no son válidas.')


def interval(data, a, b):
    start, end = parse_at(data.get(a)), parse_at(data.get(b))
    if end <= start:
        fail('El fin debe ser posterior al inicio.')
    if end - start > timedelta(days=31):
        fail('El intervalo no puede exceder 31 días.')
    return start.isoformat(), end.isoformat()


def valid_reference(data):
    import re
    ref = text(data, 'reference').upper()
    if not re.fullmatch(r'(CHG|INC)[A-Z0-9-]+', ref):
        fail('Usa una referencia completa CHG o INC.')
    return ref


def invalidate_sessions(c, identity):
    c.execute('DELETE FROM sessions WHERE user_id=?', (identity,))


@app.middleware('http')
async def security_headers(request: Request, call_next):
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        origin = request.headers.get('origin')
        if origin and origin.rstrip('/') != str(request.base_url).rstrip('/'):
            return Response('Origen no permitido.', status_code=403)
        if request.headers.get('sec-fetch-site') == 'cross-site':
            return Response('Origen no permitido.', status_code=403)
    if request.url.path.startswith('/api/'):
        await asyncio.to_thread(reconcile_planned_timers)
    response = await call_next(request)
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    response.headers['X-Operational-Pulse-Build'] = VERSION
    return response


@app.exception_handler(sqlite3.IntegrityError)
async def integrity_error(request, exc):
    return __import__('fastapi').responses.JSONResponse(status_code=409, content={'detail': 'El registro está duplicado o tiene relaciones que impiden guardar el cambio.'})


@app.get('/api/health')
def health():
    with connection() as c:
        c.execute('SELECT 1').fetchone()
    return {'status': 'ok', 'version': VERSION, 'timezone': str(TZ)}


@app.post('/api/login')
def login(response: Response, request: Request, data: dict = Body(...)):
    identity = text(data, 'identity').lower()
    password = text(data, 'password')
    with connection(True) as c:
        attempt = c.execute('SELECT * FROM login_attempts WHERE identity=?', (identity,)).fetchone()
        if attempt and attempt['locked_until'] and attempt['locked_until'] > now():
            fail('Demasiados intentos. Espera 15 minutos.', 429)
        u = c.execute('SELECT * FROM users WHERE username=? COLLATE NOCASE OR employee_id=? COLLATE NOCASE', (identity, identity)).fetchone()
        good = u and u['status'] == 'ACTIVE' and password_matches(password, u['password_hash'])
        if not good:
            failures = (attempt['failures'] if attempt and (not attempt['locked_until'] or attempt['locked_until'] > now()) else 0) + 1
            lock = (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat() if failures >= 8 else None
            c.execute('INSERT INTO login_attempts VALUES(?,?,?) ON CONFLICT(identity) DO UPDATE SET failures=excluded.failures,locked_until=excluded.locked_until', (identity, failures, lock))
        else:
            c.execute('DELETE FROM login_attempts WHERE identity=?', (identity,))
            token = secrets.token_urlsafe(40)
            c.execute('DELETE FROM sessions WHERE expires_at<=?', (now(),))
            c.execute('INSERT INTO sessions VALUES(?,?,?)', (hashlib.sha256(token.encode()).hexdigest(), u['id'], (datetime.now(timezone.utc) + timedelta(seconds=SESSION_SECONDS)).isoformat()))
            c.execute('INSERT OR REPLACE INTO session_activity VALUES(?,?)',(hashlib.sha256(token.encode()).hexdigest(),datetime.now(timezone.utc).isoformat()))
            audit(c, dict(u), 'LOGIN', u['id'])
    if not good:
        fail('Usuario o contraseña incorrectos, o cuenta pendiente de activación.', 401)
    response.set_cookie('pulse_session', token, max_age=SESSION_SECONDS, httponly=True, samesite='strict', secure=request.url.scheme == 'https')
    return {'user': public_user(u)}


@app.post('/api/logout')
def logout(request: Request, response: Response):
    token = request.cookies.get('pulse_session', '')
    with connection(True) as c:
        c.execute('DELETE FROM sessions WHERE token_hash=?', (hashlib.sha256(token.encode()).hexdigest(),))
    response.delete_cookie('pulse_session')
    return {'ok': True}


@app.post('/api/activate')
def activate(data: dict = Body(...)):
    identity, code, password = text(data, 'identity'), text(data, 'code'), text(data, 'password')
    if len(password) < 10:
        fail('La contraseña debe tener al menos 10 caracteres.')
    with connection(True) as c:
        u = c.execute('SELECT * FROM users WHERE username=? COLLATE NOCASE', (identity,)).fetchone()
        if not u or u['status'] != 'PENDING_ACTIVATION' or not password_matches(code, u['activation_hash']):
            fail('La cuenta o el código de activación no son válidos.')
        c.execute("UPDATE users SET status='ACTIVE',password_hash=?,activation_hash=NULL WHERE id=?", (password_hash(password), u['id']))
        audit(c, dict(u), 'ACTIVATE', u['id'])
    return {'ok': True}


def pause_due_timers(c, stamp):
    query = "SELECT t.id,t.overtime_id,o.scheduled_end FROM timers t JOIN overtime o ON o.id=t.overtime_id WHERE t.end_at IS NULL AND o.capture_type='PLANNED' AND o.scheduled_end<=?"
    for t in c.execute(query,(stamp,)).fetchall():
        c.execute('UPDATE timers SET end_at=? WHERE id=?',(t['scheduled_end'],t['id']))
        c.execute("UPDATE overtime SET end_at=?,execution='PAUSED',updated_at=? WHERE id=?",(t['scheduled_end'],stamp,t['overtime_id']))
        audit(c,None,'PAUSE_PLANNED_WINDOW',t['overtime_id'],{'timer_id':t['id'],'authorized_end':t['scheduled_end']})


def reconcile_planned_timers():
    stamp = now()
    query = "SELECT 1 FROM timers t JOIN overtime o ON o.id=t.overtime_id WHERE t.end_at IS NULL AND o.capture_type='PLANNED' AND o.scheduled_end<=?"
    with connection() as c:
        if not c.execute(query,(stamp,)).fetchone():
            return
    with connection(True) as c:
        pause_due_timers(c,stamp)


def cancel_pending_extensions(c, actor, overtime_id, reason):
    for e in c.execute("SELECT id FROM ot_extensions WHERE overtime_id=? AND status='SUBMITTED'",(overtime_id,)).fetchall():
        c.execute("UPDATE ot_extensions SET status='CANCELLED',decided_at=?,decision_reason=? WHERE id=?",(now(),reason,e['id']))
        audit(c,actor,'CANCEL_OT_EXTENSION',overtime_id,{'extension_id':e['id'],'reason':reason})


def finish_paused(c, actor, o):
    cancel_pending_extensions(c,actor,o['id'],'El registro de tiempo trabajado fue finalizado.')
    c.execute("UPDATE overtime SET execution='COMPLETED',status='SUBMITTED',approved_by=NULL,updated_at=? WHERE id=?",(now(),o['id']))
    audit(c,actor,'FINISH_PAUSED_OT',o['id'])


@app.get('/api/state')
def state(user=Depends(current_user)):
    with connection() as c:
        group_ids = scope(c, user)
        people = rows(c, 'users')
        visible = visible_users(c,user)
        if user['role'] == 'ADMIN':
            visible = {p['id'] for p in people}
        groups = rows(c, 'groups')
        own_ids = group_ids | ({user['group_id']} if user['group_id'] else set())
        shown_groups = groups if user['role'] == 'ADMIN' else [g for g in groups if g['id'] in own_ids]
        accounts = rows(c, 'accounts')
        visible_entities = visible | {r['id'] for table in ('overtime', 'guards', 'timers') for r in rows(c, table) if r['user_id'] in visible}
        result = {'user': public_user(user), 'can_approve': user['role'] == 'MANAGER' or bool(group_ids) and user['role'] == 'USER', 'scope': list(group_ids), 'version': VERSION, 'timezone': str(TZ), 'server_time': now(),
                  'specialties':[r for r in rows(c,'specialties') if user['role']=='ADMIN' or any(g['account_id']==r['account_id'] for g in shown_groups)],'country_codes':sorted(COUNTRY_CODES),'platforms':PLATFORMS,'service_catalog':SERVICE_CATALOG,
                  'accounts': accounts if user['role'] == 'ADMIN' else [a for a in accounts if user['role']=='MANAGER' and user['account_id']==a['id'] or any(g['account_id'] == a['id'] for g in shown_groups) or user['is_lead'] and get(c,'users',user['manager_id'])['account_id']==a['id']],
                  'groups': shown_groups, 'people': [public_user(p) for p in people if p['id'] in visible or p['role'] == 'MANAGER' and (any(g['manager_id'] == p['id'] for g in shown_groups) or user['manager_id']==p['id'])],
                  'overtime': [r for r in rows(c, 'overtime') if r['user_id'] in visible],
                  'guards': [r for r in rows(c, 'guards') if r['user_id'] in visible],
                  'timers': [r for r in rows(c, 'timers') if r['user_id'] in visible and not r['end_at']],
                  'timer_segments': [r for r in rows(c,'timers') if r['user_id'] in visible],
                  'extensions': [e for e in rows(c,'ot_extensions') if get(c,'overtime',e['overtime_id'])['user_id'] in visible],
                  'windows': [r for r in rows(c, 'windows') if user['role'] == 'ADMIN' or r['group_id'] in own_ids],
                  'rules': [dict(r, segments=json.loads(r['segments'])) for r in rows(c, 'rules') if user['role'] == 'ADMIN' or r['manager_id'] == user['id'] or user['is_lead'] and r['manager_id']==user['manager_id'] or any(g['manager_id'] == r['manager_id'] for g in shown_groups)],
                  'delegations': [dict(d, group_ids=[g['id'] for g in rows(c, 'groups') if g['manager_id']==d['manager_id']] if not d['revoked_at'] else json.loads(d['group_ids'])) for d in rows(c, 'delegations') if user['role'] == 'ADMIN' or d['manager_id'] == user['id'] or d['lead_id'] == user['id']],
                  'transfer_requests':[r for r in rows(c,'transfer_requests') if user['role']=='ADMIN' or r['requested_by']==user['id']],
                  'manager_directory':[{'id':p['id'],'first':p['first'],'last':p['last'],'employee_id':p['employee_id'],'account_name':get(c,'accounts',p['account_id'])['name'],'account_id':p['account_id'],'account_code':account_display_code(get(c,'accounts',p['account_id'])),'country_code':get(c,'accounts',p['account_id'])['country_code']} for p in people if user['role']=='MANAGER' and p['role']=='MANAGER' and p['status']=='ACTIVE' and p['id']!=user['id']],
                  'manager_transfers':[dict(r,user_name=(get(c,'users',r['user_id'])['first']+' '+get(c,'users',r['user_id'])['last']),employee_id=get(c,'users',r['user_id'])['employee_id'],source_name=get(c,'users',r['source_manager_id'])['first']+' '+get(c,'users',r['source_manager_id'])['last'],target_name=get(c,'users',r['target_manager_id'])['first']+' '+get(c,'users',r['target_manager_id'])['last']) for r in rows(c,'manager_transfers') if user['role']=='MANAGER' and user['id'] in (r['source_manager_id'],r['target_manager_id'])],
                  'recoveries': [r for r in rows(c, 'recoveries') if recovery_handler_allowed(c, user, get(c, 'users', r['user_id']))],
                  'audit': [a for a in rows(c, 'audit') if user['role'] == 'ADMIN' or a['actor_id'] in visible or a['entity_id'] in visible_entities][-300:]}
        result['group_activities']=[]
        if user['role']=='USER' and not user['is_lead'] and user['group_id']:
            for link in rows(c,'ot_group_activities'):
                w=get(c,'windows',link['window_id'])
                if w['group_id']!=user['group_id'] or parse_at(w['scheduled_end'])<=datetime.now(timezone.utc) or w['status'] not in ('SUBMITTED','APPROVED','REJECTED'):continue
                author=get(c,'users',w['created_by'])
                result['group_activities'].append({**w,'source_id':link['source_id'],'author_name':author['first']+' '+author['last']})
        people_by_id={p['id']:p for p in people}
        for request in result['transfer_requests']:
            target=people_by_id.get(request['user_id'],{});author=people_by_id.get(request['requested_by'],{})
            request['user_name']=(target.get('first','')+' '+target.get('last','')).strip()
            request['employee_id']=target.get('employee_id','')
            request['requested_by_name']=(author.get('first','')+' '+author.get('last','')).strip()
        origins={r['record_id']:json.loads(r['snapshot']) for r in rows(c,'record_origins')}
        for record in result['overtime']+result['guards']:record['origin_info']=origins.get(record['id'])
        for o in result['overtime']:
            o['actual_segments']=effective_ot_segments(c,o) if o['execution']=='COMPLETED' else []
            o['time_corrected']=bool(c.execute('SELECT 1 FROM ot_time_corrections WHERE overtime_id=?',(o['id'],)).fetchone())
            o['worked_seconds']=sum((parse_at(t['end_at'])-parse_at(t['start_at'])).total_seconds() for t in o['actual_segments']) if o['execution']=='COMPLETED' else sum(max(0,(parse_at(t['end_at'])-parse_at(t['start_at'])).total_seconds()) for t in result['timer_segments'] if t['overtime_id']==o['id'] and t['end_at'])
        recovered={r['record_id']:json.loads(r['metadata']) for r in rows(c,'time_recovery_items')}
        for record in result['overtime']+result['guards']:record['recovery_info']=recovered.get(record['id'])
        actor_ids={event['actor_id'] for event in result['audit']}
        result['audit_people']=[{key:person[key] for key in ('id','first','last','employee_id','role')} for person in people if person['id'] in actor_ids]
        for g in result['guards']:
            g['rule_snapshot'] = json.loads(g['rule_snapshot'])
            if user['role']=='USER' and g['user_id']==user['id']:
                g['timeline_days']=guard_timeline_days(g)
                g['rule_changed']=g['status']!='CANCELLED' and guard_rule_changed(c,g)
        for r in result['recoveries']:
            r.pop('code_hash', None)
            r['code_available'] = bool(r['status'] == 'ISSUED' and r['expires_at'] and parse_at(r['expires_at']) > datetime.now(timezone.utc) and c.execute('SELECT 1 FROM recovery_code_vault WHERE recovery_id=?', (r['id'],)).fetchone())
        return result


def revoke_authority(c, actor, reason, *, group_ids=(), manager_ids=(), lead_ids=()):
    group_ids, manager_ids, lead_ids = set(group_ids), set(manager_ids), set(lead_ids)
    stamp = now()
    for d in rows(c, 'delegations'):
        if d['revoked_at'] or (d['end_at'] and parse_at(d['end_at']) <= parse_at(stamp)):
            continue
        if d['manager_id'] in manager_ids or d['lead_id'] in lead_ids or group_ids.intersection(json.loads(d['group_ids'])):
            c.execute('UPDATE delegations SET revoked_at=? WHERE id=?', (stamp, d['id']))
            audit(c, actor, 'REVOKE_DELEGATION', d['id'], {'reason': reason})


def reopen_recoveries(c, actor, user_ids, reason):
    for identity in set(user_ids):
        for r in c.execute("SELECT id FROM recoveries WHERE user_id=? AND status='ISSUED'", (identity,)).fetchall():
            c.execute("UPDATE recoveries SET status='PENDING',code_hash=NULL,expires_at=NULL WHERE id=?", (r['id'],))
            c.execute('DELETE FROM recovery_code_vault WHERE recovery_id=?', (r['id'],))
            audit(c, actor, 'RECOVERY_REASSIGNED', r['id'], {'reason': reason})


def security_issued_by(c, identity):
    result = []
    for r in c.execute("SELECT * FROM recoveries WHERE status='ISSUED'"):
        last = c.execute("SELECT actor_id FROM audit WHERE action='ISSUE_RECOVERY' AND entity_id=? ORDER BY rowid DESC LIMIT 1", (r['id'],)).fetchone()
        if last and last['actor_id'] == identity:
            result.append(dict(r))
    return result


def admin_person_impact(c, identity):
    user = get(c, 'users', identity)
    stamp = now()
    today = parse_at(stamp).astimezone(TZ).date().isoformat()
    owned = [g for g in rows(c, 'groups') if g['manager_id'] == identity]
    for g in owned:
        g['members'] = [public_user(p) for p in rows(c, 'users') if p['role'] == 'USER' and p['group_id'] == g['id']]
    timers = [r for r in rows(c, 'timers') if r['user_id'] == identity and not r['end_at']]
    overtime = [r for r in rows(c, 'overtime') if r['user_id'] == identity]
    documents = [r for r in overtime if r['status'] == 'DOCUMENTATION' or r['execution'] == 'RUNNING' and r['capture_type'] == 'EMERGENCY']
    scheduled = [r for r in overtime if r['execution'] == 'SCHEDULED' and r['status'] not in ('CANCELLED','REJECTED')]
    guards = []
    for r in rows(c, 'guards'):
        if r['user_id'] != identity or r['status'] not in ('SUBMITTED','APPROVED') or date.fromisoformat(r['end_date']) + timedelta(days=1) < date.fromisoformat(today):
            continue
        snapshot = json.loads(r['rule_snapshot'])
        if snapshot.get('coverage_end_at') and parse_at(snapshot['coverage_end_at']) <= parse_at(stamp):
            continue
        guards.append(r)
    delegations = [r for r in rows(c, 'delegations') if not r['revoked_at'] and (not r['end_at'] or parse_at(r['end_at']) > parse_at(stamp)) and (r['manager_id'] == identity or r['lead_id'] == identity)]
    issued = security_issued_by(c, identity)
    direct_leads=[public_user(p) for p in rows(c,'users') if p['is_lead'] and p['manager_id']==identity and p['status']!='DISABLED']
    payload = {'direct_leads':direct_leads,'person': public_user(user), 'groups': owned, 'timers': timers, 'documents': documents, 'scheduled': scheduled, 'guards': guards, 'delegations': delegations, 'issued_codes': [{'id': r['id'], 'user_id': r['user_id']} for r in issued],
               'pending_time': sum(r['status']=='SUBMITTED' for r in overtime) + c.execute("SELECT count(*) FROM guards WHERE user_id=? AND status='SUBMITTED'", (identity,)).fetchone()[0]}
    fingerprint = hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(',',':'), ensure_ascii=False).encode()).hexdigest()
    candidates = [public_user(p) for p in rows(c, 'users') if p['role']=='MANAGER' and p['status']=='ACTIVE' and p['id']!=identity]
    return {**payload, 'expected': fingerprint, 'candidates': candidates, 'history_preserved': True}


def authorize_offboard(c,actor,identity):
    user=get(c,'users',identity)
    if actor['role']=='ADMIN':return
    if actor['role']=='MANAGER' and user['role']=='USER' and not user['is_lead'] and user['group_id'] and get(c,'groups',user['group_id'])['manager_id']==actor['id']:return
    fail('Sólo puedes dar de baja Técnicos de tu equipo.',403)


@app.get('/api/admin/people/{identity}/impact')
def lifecycle_preview(identity: str, actor=Depends(current_user)):
    with connection() as c:
        c.execute('BEGIN')
        authorize_offboard(c,actor,identity)
        return admin_person_impact(c, identity)


@app.post('/api/admin/people/{identity}/offboard')
def offboard_person(identity: str, data: dict=Body(...), actor=Depends(current_user)):
    reason = text(data, 'reason')
    if data.get('confirm') is not True:
        fail('Confirma que revisaste el impacto de la baja.')
    with connection(True) as c:
        authorize_offboard(c,actor,identity)
        impact = admin_person_impact(c, identity)
        if not hmac.compare_digest(str(data.get('expected','')), impact['expected']):
            fail('La situación cambió desde la revisión. Actualiza el impacto antes de confirmar.', 409)
        user = get(c, 'users', identity)
        if identity == actor['id']:
            fail('No puedes dar de baja tu propia cuenta administradora.')
        if user['role']=='ADMIN' and user['status']=='ACTIVE' and c.execute("SELECT count(*) FROM users WHERE role='ADMIN' AND status='ACTIVE'").fetchone()[0] <= 1:
            fail('Debe permanecer al menos un administrador activo.')
        if user['status']=='DISABLED' and not any((impact['groups'],impact['timers'],impact['documents'],impact['scheduled'],impact['guards'],impact['delegations'],impact['issued_codes'],user['is_lead'])):
            return {'ok':True,'already_disabled':True}
        transfers = data.get('transfers', {})
        if not isinstance(transfers, dict) or set(transfers) != {g['id'] for g in impact['groups']}:
            fail('Asigna un Manager de destino a cada Grupo antes de dar de baja al Manager.')
        if impact['direct_leads']:
            replacement=get(c,'users',data.get('lead_manager_id'))
            if replacement['role']!='MANAGER' or replacement['status']!='ACTIVE' or replacement['account_id']!=user['account_id'] or replacement['id']==identity:fail('Asigna un Manager activo de la Cuenta para el Lead.')
            if c.execute("SELECT 1 FROM users WHERE is_lead=1 AND manager_id=? AND status!='DISABLED'",(replacement['id'],)).fetchone():fail('El Manager de destino ya tiene un Lead.')
            for backup in impact['direct_leads']:
                if c.execute('SELECT 1 FROM timers WHERE user_id=? AND end_at IS NULL',(backup['id'],)).fetchone():fail('Detén el Timer del Lead antes de trasladarlo.')
                revoke_authority(c,actor,reason,lead_ids=[backup['id']]);invalidate_sessions(c,backup['id'])
                c.execute('UPDATE users SET manager_id=? WHERE id=?',(replacement['id'],backup['id']))
                audit(c,actor,'TRANSFER_LEAD',backup['id'],{'from_manager':identity,'to_manager':replacement['id'],'reason':reason})
        for g in impact['groups']:
            destination = get(c,'users',str(transfers[g['id']]))
            if destination['id']==identity or destination['role']!='MANAGER' or destination['status']!='ACTIVE' or destination['account_id']!=g['account_id']:
                fail('Cada destino debe ser un Manager activo de la misma Cuenta. Crea y activa al reemplazo primero.')
        if impact['timers'] and data.get('stop_timers') is not True:
            fail('Hay un Timer activo. Confirma su cierre administrativo para continuar.')
        provided = data.get('documents', [])
        if not isinstance(provided,list):
            fail('Revisa la documentación pendiente.')
        document_data = {str(d.get('id')):d for d in provided if isinstance(d,dict)}
        if set(document_data) != {r['id'] for r in impact['documents']}:
            fail('Completa la documentación de todas las OT pendientes antes de confirmar la baja.')
        prepared = {}
        for r in impact['documents']:
            d = document_data[r['id']]
            prepared[r['id']] = {'reference':valid_reference(d), **{key:text(d,key) for key in ('description','justification','impact','in_lieu')}}
        stamp = now()
        for g in impact['groups']:
            destination_id = transfers[g['id']]
            c.execute('UPDATE groups SET manager_id=? WHERE id=?',(destination_id,g['id']))
            revoke_authority(c,actor,reason,group_ids=[g['id']])
            reopen_recoveries(c,actor,[p['id'] for p in g['members']],reason)
            audit(c,actor,'TRANSFER_GROUP',g['id'],{'from_manager':identity,'to_manager':destination_id,'members':len(g['members']),'reason':reason})
        validate_manager_leads(c,list(transfers.values()))
        for t in impact['timers']:
            o = get(c,'overtime',t['overtime_id'])
            closed_at=o['scheduled_end'] if o['capture_type']=='PLANNED' and parse_at(stamp)>parse_at(o['scheduled_end']) else stamp
            c.execute('UPDATE timers SET end_at=? WHERE id=?',(closed_at,t['id']))
            status = 'DOCUMENTATION' if o['capture_type']=='EMERGENCY' else 'SUBMITTED'
            c.execute("UPDATE overtime SET end_at=?,execution='COMPLETED',status=?,approved_by=NULL,updated_at=? WHERE id=?",(closed_at,status,stamp,o['id']))
            cancel_pending_extensions(c,actor,o['id'],'Baja administrativa del solicitante.')
            audit(c,actor,'ADMIN_STOP_TIMER_ON_OFFBOARD',o['id'],{'timer_id':t['id'],'user_id':identity,'reason':reason})
        for paused in c.execute("SELECT * FROM overtime WHERE user_id=? AND execution='PAUSED'",(identity,)).fetchall():
            finish_paused(c,actor,dict(paused))
        for oid,d in prepared.items():
            c.execute("UPDATE overtime SET reference=?,description=?,justification=?,impact=?,in_lieu=?,status='SUBMITTED',approved_by=NULL,revision=revision+1,updated_at=? WHERE id=?",(d['reference'],d['description'],d['justification'],d['impact'],d['in_lieu'],stamp,oid))
            audit(c,actor,'ADMIN_DOCUMENT_ON_OFFBOARD',oid,{'user_id':identity,'reason':reason})
        for o in impact['scheduled']:
            cancel_pending_extensions(c,actor,o['id'],'Baja administrativa del solicitante.')
            c.execute("UPDATE overtime SET status='CANCELLED',approved_by=NULL,updated_at=? WHERE id=?",(stamp,o['id']))
            audit(c,actor,'CANCEL_SCHEDULE_ON_OFFBOARD',o['id'],{'previous_status':o['status'],'reason':reason})
        for guard in impact['guards']:
            snapshot=json.loads(guard['rule_snapshot'])
            began=any(a<parse_at(stamp) for a,b in guard_coverage(guard))
            if not began:
                c.execute("UPDATE guards SET status='CANCELLED' WHERE id=?",(guard['id'],))
            else:
                snapshot=json.loads(guard['rule_snapshot']);snapshot['coverage_end_at']=stamp
                c.execute('UPDATE guards SET rule_snapshot=? WHERE id=?',(json.dumps(snapshot),guard['id']))
            audit(c,actor,'CLOSE_GUARD_ON_OFFBOARD',guard['id'],{'coverage_end_at':stamp,'reason':reason})
        revoke_authority(c,actor,reason,manager_ids=[identity],lead_ids=[identity])
        reopen_recoveries(c,actor,[r['user_id'] for r in impact['issued_codes'] if r['user_id']!=identity],reason)
        c.execute("UPDATE recoveries SET status='CANCELLED',code_hash=NULL WHERE user_id=? AND status IN ('PENDING','ISSUED')",(identity,))
        c.execute('DELETE FROM recovery_code_vault WHERE recovery_id IN (SELECT id FROM recoveries WHERE user_id=?)',(identity,))
        c.execute("UPDATE users SET status='DISABLED',is_lead=0,activation_hash=NULL WHERE id=?",(identity,))
        c.execute("UPDATE manager_transfers SET status='CANCELLED',decided_at=?,decision_reason=? WHERE user_id=? AND status='SUBMITTED'",(stamp,'Baja del Técnico: '+reason,identity))
        invalidate_sessions(c,identity)
        audit(c,actor,'OFFBOARD_PERSON',identity,{'reason':reason,'groups_transferred':len(impact['groups']),'timers_closed':len(impact['timers']),'documents_submitted':len(prepared),'guards_closed':len(impact['guards'])})
    return {'ok':True,'groups_transferred':len(impact['groups']),'history_preserved':True}


@app.post('/api/admin/people/{identity}/reactivate')
def reactivate_person(identity: str, data: dict=Body(...), actor=Depends(admin)):
    reason=text(data,'reason')
    with connection(True) as c:
        user=get(c,'users',identity)
        if user['status']!='DISABLED':
            fail('Sólo se reactiva una cuenta en baja.')
        if user['role']=='MANAGER' and c.execute('SELECT 1 FROM groups WHERE manager_id=?',(identity,)).fetchone():
            fail('Regulariza primero los Grupos que todavía dependen de este Manager en baja.')
        revive_lead=bool(user['role']=='USER' and user['manager_id'] and not user['group_id'])
        if revive_lead:
            m=get(c,'users',user['manager_id'])
            if m['status']!='ACTIVE' or c.execute("SELECT 1 FROM users WHERE manager_id=? AND is_lead=1 AND status!='DISABLED' AND id<>?",(m['id'],identity)).fetchone():fail('El Manager del Lead no está disponible o ya tiene otro backup. Reasigna primero a la persona a un Grupo como Técnico.')
        code='ACT-'+secrets.token_hex(5).upper()
        c.execute("UPDATE users SET status='PENDING_ACTIVATION',is_lead=?,password_hash=NULL,activation_hash=? WHERE id=?",(int(revive_lead),password_hash(code),identity))
        invalidate_sessions(c,identity)
        revoke_authority(c,actor,reason,manager_ids=[identity],lead_ids=[identity])
        audit(c,actor,'REACTIVATE_PERSON',identity,{'reason':reason})
    return {'id':identity,'activation_code':code,'username':user['username']}


def normalized_key(value):
    raw=''.join(ch for ch in unicodedata.normalize('NFD',str(value)) if not unicodedata.combining(ch)).upper()
    result=re.sub('[^A-Z0-9]+','-',raw).strip('-')
    result={'Z-OS':'ZOS','SO-LINUX':'LINUX','SO-UNIX':'UNIX','BD-ORACLE':'ORACLE'}.get(result,result)
    if not result or len(result)>24:fail('Usa una clave de hasta 24 letras, números o guiones.')
    return result


def standard_group_code(c, account, platform, service, exclude=''):
    if account['country_code'] not in COUNTRY_CODES:fail('Primero define el país de la Cuenta en Organización.')
    if platform and platform not in PLATFORMS:fail('Plataforma no válida.')
    stem=account_display_code(account)+('-'+platform if platform else '')+'-'+service;number=1
    while True:
        candidate=stem if number==1 else stem+'-'+str(number).zfill(2)
        if not c.execute('SELECT 1 FROM groups WHERE (code=? COLLATE NOCASE OR standard_code=? COLLATE NOCASE) AND id!=?',(candidate,candidate,exclude)).fetchone():return candidate
        number+=1


@app.post('/api/accounts')
@app.put('/api/accounts/{identity}')
def save_account(data: dict = Body(...), identity: str = '', actor=Depends(admin)):
    name=text(data,'name')
    with connection(True) as c:
        old=get(c,'accounts',identity) if identity else None
        country=str(data.get('country_code',old['country_code'] if old else '')).upper().strip()
        if country and country not in COUNTRY_CODES:fail('Selecciona un país válido (ISO de dos letras).')
        supplied=data.get('base_code') or data.get('code') or (old['base_code'] if old else '')
        if old and supplied==old['code']:supplied=old['base_code'] or old['code']
        base=normalized_key(supplied)
        if country and base.startswith(country+'-'):base=base[len(country)+1:]
        if old:
            if c.execute('SELECT 1 FROM groups WHERE account_id=?',(identity,)).fetchone() and ((old['country_code'] and country!=old['country_code']) or base!=old['base_code']):
                fail('Una Cuenta con Grupos conserva su país y clave. Traslada primero los Grupos para cambiar esa identidad.')
            legacy=old['code']!=account_display_code(old) or not old['country_code']
            code=old['code'] if legacy else country+'-'+base if country else base
            c.execute('UPDATE accounts SET code=?,name=?,country_code=?,base_code=? WHERE id=?',(code,name,country,base,identity))
        else:
            identity=uid();code=country+'-'+base if country else base
            c.execute('INSERT INTO accounts(id,code,name,country_code,base_code) VALUES(?,?,?,?,?)',(identity,code,name,country,base))
        audit(c,actor,'SAVE_ACCOUNT',identity,{'country_code':country,'base_code':base})
    return {'id': identity}


@app.delete('/api/accounts/{identity}')
def delete_account(identity: str, actor=Depends(admin)):
    with connection(True) as c:
        get(c, 'accounts', identity)
        if c.execute('SELECT 1 FROM groups WHERE account_id=?', (identity,)).fetchone():
            fail('La Cuenta contiene Grupos. Muévelos o elimínalos primero.')
        if c.execute('SELECT 1 FROM users WHERE account_id=?', (identity,)).fetchone():
            fail('La Cuenta contiene Managers. Muévelos primero.')
        c.execute('DELETE FROM accounts WHERE id=?', (identity,))
        audit(c, actor, 'DELETE_ACCOUNT', identity)
    return {'ok': True}


def validate_manager_leads(c, manager_ids):
    """Validate the resulting ownership inside the same write transaction; never erase legacy data."""
    for manager_id in set(manager_ids) - {None, ''}:
        holders = c.execute("SELECT id FROM users WHERE manager_id=? AND is_lead=1 AND status!='DISABLED'", (manager_id,)).fetchall()
        if len(holders) > 1:
            fail('Este Manager ya tiene un Lead. Retira una designación antes de reunir los equipos; sólo puede haber un backup por Manager.')


@app.post('/api/groups')
@app.put('/api/groups/{identity}')
def save_group(data: dict = Body(...), identity: str = '', actor=Depends(current_user)):
    aid, mid = text(data, 'account_id'), data.get('manager_id') or None
    with connection(True) as c:
        account=get(c, 'accounts', aid)
        old = get(c, 'groups', identity) if identity else None
        if identity and c.execute('SELECT 1 FROM retired_groups WHERE group_id=?',(identity,)).fetchone():fail('El Grupo ya fue eliminado.',409)
        if actor['role']=='MANAGER':
            if aid!=actor['account_id'] or mid not in (None,actor['id']) or old and old['manager_id']!=actor['id']:fail('Sólo puedes administrar Grupos de tu propio equipo.',403)
            mid=actor['id']
        elif actor['role']!='ADMIN':fail('No tienes permiso para administrar Grupos.',403)
        platform=str(data.get('platform_code',old['platform_code'] if old else '')).upper()
        service=normalized_key(data.get('service_key') or old['service_key']) if data.get('service_key') or old and old['service_key'] else ''
        standard=old['standard_code'] if old else ''
        if service:
            if platform and platform not in PLATFORMS:fail('Plataforma no válida.')
            changed=not old or old['account_id']!=aid or old['platform_code']!=platform or old['service_key']!=service or not standard
            if changed:standard=standard_group_code(c,account,platform,service,identity)
            code=standard if not old or old['code']==old['standard_code'] else old['code']
        else:
            code=text(data,'code').upper() if data.get('code') else old['code'] if old else ''
            if not code:fail('Selecciona la plataforma y especialidad del Grupo.')
        if mid:
            m = get(c, 'users', mid)
            if m['role'] != 'MANAGER' or m['status']=='DISABLED' or m['account_id'] != aid:
                fail('El Manager debe pertenecer a la misma Cuenta que el Grupo.')
        if not mid and identity and c.execute('SELECT 1 FROM users WHERE group_id=? AND is_lead=1',(identity,)).fetchone():
            fail('Retira la capacidad Lead antes de dejar su Grupo sin Manager.')
        if old:
            members=c.execute('SELECT id,status FROM users WHERE group_id=?',(identity,)).fetchall()
            if old['manager_id']!=mid and members:
                transfer_reason=text(data,'transfer_reason')
                if any(p['status']!='DISABLED' for p in members) and (not mid or get(c,'users',mid)['status']!='ACTIVE'):
                    fail('Un Grupo con personas activas necesita un Manager de destino activo.')
                reopen_recoveries(c,actor,[p['id'] for p in members],transfer_reason)
                audit(c,actor,'TRANSFER_GROUP',identity,{'from_manager':old['manager_id'],'to_manager':mid,'reason':transfer_reason})
            c.execute('UPDATE groups SET code=?,name=?,account_id=?,manager_id=?,platform_code=?,service_key=?,standard_code=? WHERE id=?', (code,text(data,'name'),aid,mid,platform,service,standard,identity))
            if old['manager_id'] != mid:
                # A change of Manager never silently transfers delegated approval rights.
                for d in rows(c, 'delegations'):
                    if identity in json.loads(d['group_ids']) and not d['revoked_at']:
                        c.execute('UPDATE delegations SET revoked_at=? WHERE id=?', (now(), d['id']))
        else:
            identity = uid()
            c.execute('INSERT INTO groups(id,account_id,manager_id,code,name,platform_code,service_key,standard_code) VALUES(?,?,?,?,?,?,?,?)', (identity,aid,mid,code,text(data,'name'),platform,service,standard))
        if service:
            c.execute('INSERT OR IGNORE INTO specialties VALUES(?,?,?,?)',(uid(),aid,service,str(data.get('specialty_name') or service).strip()[:100]))
        validate_manager_leads(c,[mid])
        audit(c, actor, 'SAVE_GROUP', identity)
    return {'id': identity}


def group_removal_impact(c,identity,actor):
    g=get(c,'groups',identity)
    if actor['role']!='ADMIN' and (actor['role']!='MANAGER' or g['manager_id']!=actor['id']):fail('Sólo puedes eliminar Grupos de tu equipo.',403)
    if c.execute('SELECT 1 FROM retired_groups WHERE group_id=?',(identity,)).fetchone():fail('El Grupo ya fue eliminado.',409)
    members=[public_user(p) for p in rows(c,'users') if p['group_id']==identity]
    if actor['role']=='ADMIN' and members:fail('El Manager debe resolver los Técnicos de este Grupo antes de eliminarlo.',403)
    expected=hashlib.sha256(json.dumps({'group':g,'members':sorted(p['id'] for p in members)},sort_keys=True).encode()).hexdigest()
    destinations=[x for x in rows(c,'groups') if x['id']!=identity and x['manager_id']==g['manager_id'] and x['account_id']==g['account_id']]
    return {'group':g,'members':members,'destinations':destinations,'expected':expected}


@app.get('/api/groups/{identity}/removal-impact')
def preview_group_removal(identity:str,actor=Depends(current_user)):
    with connection() as c:return group_removal_impact(c,identity,actor)


@app.delete('/api/groups/{identity}')
def delete_group(identity: str, data:dict=Body(...), actor=Depends(current_user)):
    reason=text(data,'reason')
    if data.get('confirm') is not True:fail('Confirma que revisaste el destino de los Técnicos.')
    with connection(True) as c:
        impact=group_removal_impact(c,identity,actor)
        if not hmac.compare_digest(str(data.get('expected','')),impact['expected']):fail('El Grupo cambió. Revisa de nuevo los Técnicos y el destino.',409)
        destination=None
        if impact['members']:
            destination=next((g for g in impact['destinations'] if g['id']==data.get('destination_group_id')),None)
            if not destination:fail('Elige otro Grupo del mismo Manager para recibir a los Técnicos.')
            for user in impact['members']:
                if c.execute('SELECT 1 FROM timers WHERE user_id=? AND end_at IS NULL',(user['id'],)).fetchone():fail('Detén los Timers activos del Grupo antes de mover a sus Técnicos.')
                if user['is_lead']:fail('El Lead debe estar asignado al Manager antes de eliminar este Grupo.')
            for user in impact['members']:
                for table in ('overtime','guards'):
                    for record in c.execute('SELECT id FROM '+table+' WHERE user_id=?',(user['id'],)).fetchall():capture_origin(c,record['id'],user,inferred=True)
                c.execute('UPDATE users SET group_id=? WHERE id=?',(destination['id'],user['id']))
                invalidate_sessions(c,user['id'])
                audit(c,actor,'TRANSFER_PERSON',user['id'],{'from_group':identity,'to_group':destination['id'],'reason':reason})
            reopen_recoveries(c,actor,[p['id'] for p in impact['members']],reason)
        stamp=now()
        c.execute("UPDATE manager_transfers SET status='CANCELLED',decided_at=?,decision_reason=? WHERE status='SUBMITTED' AND (source_group_id=? OR destination_group_id=?)",(stamp,'Grupo eliminado: '+reason,identity,identity))
        c.execute("UPDATE transfer_requests SET status='CANCELLED',decided_at=?,decision_reason=? WHERE status='SUBMITTED' AND (source_group_id=? OR destination_group_id=?)",(stamp,'Grupo eliminado: '+reason,identity,identity))
        # Retain the original row for windows and transfer history, outside the active catalogue.
        c.execute('INSERT INTO retired_groups VALUES(?,?,?)',(identity,stamp,reason))
        c.execute('UPDATE groups SET manager_id=NULL WHERE id=?',(identity,))
        audit(c,actor,'DELETE_GROUP',identity,{'reason':reason,'moved':len(impact['members']),'destination_group_id':destination['id'] if destination else None,'group':impact['group']})
    return {'ok':True,'moved':len(impact['members']),'history_preserved':True}


@app.post('/api/people')
@app.put('/api/people/{identity}')
def save_person(data: dict = Body(...), identity: str = '', actor=Depends(current_user)):
    with connection(True) as c:
        old = get(c, 'users', identity) if identity else None
        if actor['role']=='MANAGER':
            if old and (old['role']!='USER' or old['is_lead'] or old['status']=='DISABLED' or not old['group_id'] or get(c,'groups',old['group_id'])['manager_id']!=actor['id']):
                fail('Sólo puedes administrar Técnicos vigentes de tus propios Grupos.',403)
            if data.get('role','USER')!='USER' or data.get('is_lead') or data.get('manager_id'):
                fail('El Manager no puede crear Administradores, Managers ni Leads.',403)
            destination=get(c,'groups',data.get('group_id') or (old['group_id'] if old else None))
            if destination['manager_id']!=actor['id'] or destination['account_id']!=actor['account_id']:
                fail('El Grupo debe pertenecer a tu equipo. No se permite cambiar de Manager por esta vía.',403)
        elif actor['role']=='ADMIN':
            if data.get('_transfer_only'):fail('Admin no autoriza movimientos de Técnicos.',403)
            if (old and old['role']=='USER' and not old['is_lead'] and not data.get('is_lead')) or (not old and data.get('role')=='USER' and not data.get('is_lead')):
                fail('Los Técnicos los administra su Manager; Admin gestiona Administradores, Managers y Leads.',403)
        else:
            fail('No tienes permiso para administrar personas.',403)
        if data.get('_transfer_only'):
            if not old or old['role']!='USER' or old['status']=='DISABLED' or old['is_lead']:
                fail('Sólo se transfieren Técnicos activos o pendientes de activación.')
            if text(data,'expected_group_id')!=old['group_id']:
                fail('El Técnico cambió de Grupo desde que abriste el traslado. Actualiza la organización.',409)
            # Read editable identity fields under the same write lock; moving never replays a stale profile.
            data={**old,'group_id':text(data,'group_id'),'transfer_reason':text(data,'transfer_reason'),'transfer_request_id':data.get('transfer_request_id')}
        role = old['role'] if old else text(data, 'role')
        if role not in ('ADMIN', 'MANAGER', 'USER'):
            fail('Rol no válido.')
        aid = data.get('account_id') if role == 'MANAGER' else None
        gid = data.get('group_id') if role == 'USER' else None
        lead = bool(data.get('is_lead')) and role == 'USER'
        if aid:
            get(c, 'accounts', aid)
        elif role == 'MANAGER':
            fail('El Manager debe pertenecer a una Cuenta.')
        lead_manager=None
        if role=='USER':
            if lead:
                lead_manager=data.get('manager_id') or (get(c,'groups',gid)['manager_id'] if gid else old.get('manager_id') if old else None)
                m=get(c,'users',lead_manager)
                if m['role']!='MANAGER' or m['status']=='DISABLED':fail('Selecciona un Manager vigente para el Lead.')
                if c.execute("SELECT 1 FROM users WHERE manager_id=? AND is_lead=1 AND status!='DISABLED' AND id<>?",(lead_manager,identity)).fetchone():fail('Este Manager ya tiene un Lead.')
                gid=None
            else:get(c,'groups',gid)
        group_ids = list(dict.fromkeys(data.get('group_ids', [g['id'] for g in c.execute('SELECT id FROM groups WHERE manager_id=?',(identity,))] if old else []))) if role == 'MANAGER' else []
        for g_id in group_ids:
            g = get(c, 'groups', g_id)
            if g['account_id'] != aid:
                fail('Un Manager no puede administrar Grupos de distintas Cuentas.')
            if g['manager_id'] and g['manager_id'] != identity:
                fail('El Grupo ya pertenece a otro Manager.')
        if old and role == 'MANAGER':
            for g in c.execute('SELECT * FROM groups WHERE manager_id=?', (identity,)):
                if g['account_id'] != aid and g['id'] not in group_ids:
                    fail('Mueve o desasigna los Grupos actuales antes de cambiar la Cuenta del Manager.')
        status = data.get('status', old['status'] if old else 'PENDING_ACTIVATION')
        if status not in ('ACTIVE', 'DISABLED', 'PENDING_ACTIVATION'):
            fail('Estado no válido.')
        if old and status!=old['status']:
            fail('Utiliza Baja o Reactivar para cambiar el acceso de una persona.')
        if status=='DISABLED' and lead:
            fail('Una cuenta en baja no puede recibir capacidad Lead; reactívala primero.')
        if old and role=='USER' and (gid!=old['group_id'] or lead_manager!=old.get('manager_id')):
            if c.execute('SELECT 1 FROM timers WHERE user_id=? AND end_at IS NULL',(identity,)).fetchone():
                fail('Detén el Timer antes de transferir al Técnico de Grupo.')
            transfer_reason=text(data,'transfer_reason')
            request_id=data.get('transfer_request_id')
            if request_id:
                request=get(c,'transfer_requests',request_id)
                if request['status']!='SUBMITTED' or request['user_id']!=identity or request['source_group_id']!=old['group_id']:
                    fail('La solicitud cambió o el Técnico ya fue trasladado. Revisa la organización.',409)
            for table in ('overtime','guards'):
                for record in c.execute('SELECT id FROM '+table+' WHERE user_id=?',(identity,)).fetchall():
                    capture_origin(c,record['id'],old,inferred=True)
            if request_id:
                c.execute("UPDATE transfer_requests SET status='APPROVED',decided_by=?,decided_at=?,decision_reason=?,destination_group_id=? WHERE id=?",(actor['id'],now(),transfer_reason,gid,request_id))
            old_owner=owner_manager(c,old);new_owner=lead_manager or get(c,'groups',gid)['manager_id']
            if old['status']!='DISABLED' and (not new_owner or get(c,'users',new_owner)['status']!='ACTIVE'):
                fail('El Grupo de destino necesita un Manager activo.')
            revoke_authority(c,actor,transfer_reason,lead_ids=[identity])
            if old_owner!=new_owner:
                reopen_recoveries(c,actor,[identity],transfer_reason)
            invalidate_sessions(c,identity)
            c.execute("UPDATE manager_transfers SET status='CANCELLED',decided_at=?,decision_reason=? WHERE user_id=? AND status='SUBMITTED'",(now(),'El Técnico cambió de Grupo antes de resolver el traslado.',identity))
            audit(c,actor,'TRANSFER_PERSON',identity,{'from_group':old['group_id'],'to_group':gid,'from_manager':old_owner,'to_manager':new_owner,'reason':transfer_reason})
        if old and role=='MANAGER':
            for owned in c.execute('SELECT id FROM groups WHERE manager_id=?',(identity,)).fetchall():
                if owned['id'] not in group_ids and c.execute("SELECT 1 FROM users WHERE group_id=? AND status!='DISABLED'",(owned['id'],)).fetchone():
                    fail('Transfiere el Grupo con su equipo antes de retirarlo del Manager.')
        if old and old['role'] == 'ADMIN' and status != 'ACTIVE':
            active_count = c.execute("SELECT count(*) FROM users WHERE role='ADMIN' AND status='ACTIVE'").fetchone()[0]
            if active_count <= 1 or old['id'] == actor['id']:
                fail('No puedes desactivar tu propia cuenta ni al último administrador activo.')
        username = text(data, 'username') if role == 'ADMIN' else text(data, 'employee_id').upper()
        employee_id = None if role == 'ADMIN' else username
        code = None
        if old:
            if status == 'ACTIVE' and not old['password_hash']:
                fail('La persona debe activar su cuenta y definir su contraseña primero.')
            if role == 'USER' and (gid != old['group_id'] or lead_manager!=old.get('manager_id') or not lead):
                c.execute('UPDATE delegations SET revoked_at=? WHERE lead_id=? AND revoked_at IS NULL', (now(), identity))
            c.execute('UPDATE users SET username=?,employee_id=?,first=?,last=?,email=?,account_id=?,group_id=?,is_lead=?,status=? WHERE id=?',
                      (username, employee_id, text(data, 'first'), text(data, 'last'), text(data, 'email'), aid, gid, int(lead), status, identity))
            if status != 'ACTIVE':
                invalidate_sessions(c, identity)
        else:
            identity = uid()
            code = 'ACT-' + secrets.token_hex(5).upper()
            c.execute('INSERT INTO users(id,role,username,employee_id,first,last,email,account_id,group_id,is_lead,status,password_hash,activation_hash,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)', (identity, role, username, employee_id, text(data, 'first'), text(data, 'last'), text(data, 'email'), aid, gid, int(lead), 'PENDING_ACTIVATION', None, password_hash(code), now()))
        if role=='USER':c.execute('UPDATE users SET manager_id=? WHERE id=?',(lead_manager,identity))
        if role == 'MANAGER':
            c.execute('UPDATE groups SET manager_id=NULL WHERE manager_id=?', (identity,))
            for g_id in group_ids:
                c.execute('UPDATE groups SET manager_id=? WHERE id=?', (identity, g_id))
            for d in c.execute('SELECT * FROM delegations WHERE manager_id=? AND revoked_at IS NULL', (identity,)).fetchall():
                if not set(json.loads(d['group_ids'])).issubset(set(group_ids)):
                    c.execute('UPDATE delegations SET revoked_at=? WHERE id=?', (now(), d['id']))
        if role=='USER' and lead:
            validate_manager_leads(c,[lead_manager])
        elif role=='MANAGER':
            validate_manager_leads(c,[identity])
        audit(c, actor, 'SAVE_PERSON', identity)
    return {'id': identity, 'activation_code': code}


def capture_origin(c, record_id, user, inferred=False):
    if not user['group_id']:
        mid=owner_manager(c,user);m=get(c,'users',mid);a=get(c,'accounts',m['account_id'])
        snapshot={'group_id':None,'group_name':'Lead del Manager','group_code':'','account_name':a['name'],'account_code':account_display_code(a),'manager_id':mid,'inferred':inferred}
        c.execute('INSERT OR IGNORE INTO record_origins VALUES(?,?)',(record_id,json.dumps(snapshot)));return
    g=dict(get(c,'groups',user['group_id']));a=dict(get(c,'accounts',g['account_id']))
    snapshot={'group_id':g['id'],'group_name':g['name'],'group_code':group_display_code(g,a),'account_name':a['name'],'account_code':account_display_code(a),'manager_id':g['manager_id'],'inferred':inferred}
    c.execute('INSERT OR IGNORE INTO record_origins VALUES(?,?)',(record_id,json.dumps(snapshot)))


@app.post('/api/admin/organization/transfers')
def organize_technicians(data:dict=Body(...),actor=Depends(admin)):
    fail('Admin no autoriza movimientos de Técnicos. Corresponden a sus Managers.',403)


@app.post('/api/transfer-requests')
def request_transfer(data:dict=Body(...),actor=Depends(current_user)):
    fail('Usa los movimientos de tu equipo o solicita traslado al Manager de destino.',410)


@app.post('/api/admin/transfer-requests/{identity}/reject')
def reject_transfer(identity:str,data:dict=Body(...),actor=Depends(admin)):
    fail('Admin no decide movimientos de Técnicos.',403)


@app.post('/api/admin/people/{identity}/transfer')
def transfer_technician(identity: str, data: dict=Body(...), actor=Depends(admin)):
    fail('Admin no autoriza movimientos de Técnicos.',403)


@app.post('/api/team/transfers')
def create_manager_transfer(data:dict=Body(...),actor=Depends(current_user)):
    if actor['role']!='MANAGER':fail('Sólo el Manager puede solicitar un traslado.',403)
    with connection(True) as c:
        user=get(c,'users',text(data,'user_id'));target=get(c,'users',text(data,'target_manager_id'));reason=text(data,'reason')
        if user['role']!='USER' or user['is_lead'] or user['status']=='DISABLED' or not user['group_id'] or get(c,'groups',user['group_id'])['manager_id']!=actor['id']:fail('El Técnico debe pertenecer a tu equipo.',403)
        if data.get('expected_group_id') and data['expected_group_id']!=user['group_id']:fail('El Técnico cambió de Grupo. Vuelve a abrir su ficha.',409)
        if target['role']!='MANAGER' or target['status']!='ACTIVE' or target['id']==actor['id']:fail('Selecciona otro Manager activo.')
        if len(reason)>2000:fail('Resume el motivo en 2000 caracteres.')
        if c.execute("SELECT 1 FROM manager_transfers WHERE user_id=? AND status='SUBMITTED'",(user['id'],)).fetchone():fail('Este Técnico ya tiene un traslado pendiente.',409)
        identity=uid();c.execute('INSERT INTO manager_transfers(id,user_id,source_manager_id,target_manager_id,source_group_id,reason,status,created_at) VALUES(?,?,?,?,?,?,?,?)',(identity,user['id'],actor['id'],target['id'],user['group_id'],reason,'SUBMITTED',now()))
        audit(c,actor,'REQUEST_MANAGER_TRANSFER',identity,{'user_id':user['id'],'target_manager_id':target['id']})
    return {'id':identity}


@app.post('/api/team/transfers/{identity}/decision')
def decide_manager_transfer(identity:str,data:dict=Body(...),actor=Depends(current_user)):
    if actor['role']!='MANAGER':fail('Sólo un Manager puede resolver este traslado.',403)
    decision=text(data,'decision');reason=text(data,'reason')
    if decision not in ('APPROVED','REJECTED','CANCELLED'):fail('Decisión no válida.')
    with connection(True) as c:
        request=get(c,'manager_transfers',identity)
        expected=request['source_manager_id'] if decision=='CANCELLED' else request['target_manager_id']
        if actor['id']!=expected:fail('Sólo el responsable correspondiente puede resolver el traslado.',403)
        if request['status']!='SUBMITTED':fail('El traslado ya fue resuelto.',409)
        if decision=='APPROVED':
            user=get(c,'users',request['user_id']);destination=get(c,'groups',text(data,'group_id'))
            if destination['manager_id']!=actor['id'] or destination['account_id']!=actor['account_id']:fail('Elige uno de tus Grupos.',403)
            if user['role']!='USER' or user['is_lead'] or user['status']=='DISABLED' or user['group_id']!=request['source_group_id'] or get(c,'groups',user['group_id'])['manager_id']!=request['source_manager_id']:fail('La asignación del Técnico cambió. Solicita un traslado nuevo.',409)
            if get(c,'users',request['source_manager_id'])['status']!='ACTIVE':fail('El Manager de origen ya no está activo.',409)
            if c.execute('SELECT 1 FROM timers WHERE user_id=? AND end_at IS NULL',(user['id'],)).fetchone():fail('El Técnico debe detener su Timer antes de aceptar el traslado.')
            for table in ('overtime','guards'):
                for record in c.execute('SELECT id FROM '+table+' WHERE user_id=?',(user['id'],)).fetchall():capture_origin(c,record['id'],user,inferred=True)
            revoke_authority(c,actor,reason,lead_ids=[user['id']]);reopen_recoveries(c,actor,[user['id']],reason);invalidate_sessions(c,user['id'])
            c.execute('UPDATE users SET group_id=? WHERE id=?',(destination['id'],user['id']))
            c.execute('UPDATE manager_transfers SET destination_group_id=? WHERE id=?',(destination['id'],identity))
            audit(c,actor,'TRANSFER_PERSON',user['id'],{'from_group':user['group_id'],'to_group':destination['id'],'from_manager':request['source_manager_id'],'to_manager':actor['id'],'reason':reason,'request_id':identity})
        c.execute('UPDATE manager_transfers SET status=?,decided_at=?,decision_reason=? WHERE id=?',(decision,now(),reason,identity))
        audit(c,actor,'DECIDE_MANAGER_TRANSFER',identity,{'decision':decision,'reason':reason})
    return {'ok':True,'status':decision}


@app.post('/api/team/people/{identity}/move')
def move_team_technician(identity:str,data:dict=Body(...),actor=Depends(current_user)):
    if actor['role']!='MANAGER':fail('Sólo el Manager administra los movimientos de su equipo.',403)
    return save_person({**{key:data.get(key) for key in ('group_id','expected_group_id','transfer_reason')},'_transfer_only':True},identity,actor)


@app.post('/api/people/{identity}/activation')
def new_activation(identity: str, actor=Depends(current_user)):
    with connection(True) as c:
        u = get(c, 'users', identity)
        if actor['role']=='MANAGER':
            if u['role']!='USER' or u['is_lead'] or not u['group_id'] or get(c,'groups',u['group_id'])['manager_id']!=actor['id']:fail('Sólo puedes generar códigos de tus Técnicos.',403)
        elif actor['role']!='ADMIN' or u['role']=='USER' and not u['is_lead']:
            fail('Este acceso corresponde a otro responsable.',403)
        if u['status'] != 'PENDING_ACTIVATION':
            fail('La cuenta ya está activa. Usa recuperación de acceso.')
        code = 'ACT-' + secrets.token_hex(5).upper()
        c.execute('UPDATE users SET activation_hash=? WHERE id=?', (password_hash(code), identity))
        audit(c, actor, 'REGENERATE_ACTIVATION', identity)
    return {'activation_code': code}


@app.put('/api/profile')
def profile(data: dict = Body(...), actor=Depends(current_user)):
    with connection(True) as c:
        c.execute('UPDATE users SET first=?,last=?,email=? WHERE id=?', (text(data, 'first'), text(data, 'last'), text(data, 'email'), actor['id']))
        audit(c, actor, 'PROFILE', actor['id'])
    return {'ok': True}


@app.put('/api/password')
def change_password(data: dict = Body(...), actor=Depends(current_user)):
    if not password_matches(text(data, 'current'), actor['password_hash']):
        fail('La contraseña actual no es correcta.')
    password = text(data, 'password')
    if len(password) < 10:
        fail('Usa al menos 10 caracteres.')
    with connection(True) as c:
        c.execute('UPDATE users SET password_hash=? WHERE id=?', (password_hash(password), actor['id']))
        invalidate_sessions(c, actor['id'])
        audit(c, actor, 'CHANGE_PASSWORD', actor['id'])
    return {'ok': True, 'reauthenticate': True}


def recovery_handler_allowed(c, actor, target):
    if actor['id'] == target['id'] or actor['role'] == 'USER':
        return False
    if target['role'] in ('ADMIN','MANAGER') or target['is_lead']:
        return actor['role']=='ADMIN'
    manager_id = None
    if target['role'] == 'USER' and target['group_id']:
        manager_row = c.execute("SELECT u.id FROM groups g JOIN users u ON u.id=g.manager_id WHERE g.id=? AND u.role='MANAGER' AND u.status='ACTIVE'", (target['group_id'],)).fetchone()
        manager_id = manager_row['id'] if manager_row else None
    if manager_id:
        return actor['role'] == 'MANAGER' and actor['id'] == manager_id
    return actor['role'] == 'ADMIN'


def recovery_cipher(create=False):
    """Authenticated encryption; key stays outside SQLite and outside the source ZIP."""
    try:
        try:
            key = RECOVERY_KEY.read_bytes()
        except FileNotFoundError:
            if not create:
                fail('No está disponible la clave del código guardado. Regenera el código para poder consultarlo.', 409)
            RECOVERY_KEY.parent.mkdir(parents=True, exist_ok=True)
            candidate = Fernet.generate_key()
            try:
                descriptor = os.open(RECOVERY_KEY, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            except FileExistsError:
                key = RECOVERY_KEY.read_bytes()
            else:
                with os.fdopen(descriptor, 'wb') as output:
                    output.write(candidate)
                key = candidate
        return Fernet(key)
    except (OSError, ValueError):
        fail('No se pudo acceder a la clave local de recuperación. Revisa el archivo de clave y sus permisos.', 503)


@app.post('/api/recoveries')
def request_recovery(data: dict = Body(...)):
    with connection(True) as c:
        u = c.execute('SELECT * FROM users WHERE username=? COLLATE NOCASE', (text(data, 'identity'),)).fetchone()
        if u and u['status'] == 'ACTIVE' and not c.execute("SELECT 1 FROM recoveries WHERE user_id=? AND status IN ('PENDING','ISSUED')", (u['id'],)).fetchone():
            c.execute('INSERT INTO recoveries VALUES(?,?,?,?,?,?)', (uid(), u['id'], 'PENDING', None, None, now()))
    return {'message': 'Si la cuenta está activa, se registró la solicitud. Contacta a tu Manager o a Administración.'}


@app.post('/api/recoveries/{identity}/issue')
def issue_recovery(identity: str, actor=Depends(current_user)):
    with connection(True) as c:
        r = get(c, 'recoveries', identity)
        if r['status'] in ('USED', 'CANCELLED'):
            fail('Esta solicitud ya fue utilizada o sustituida.')
        target = get(c, 'users', r['user_id'])
        if not recovery_handler_allowed(c, actor, target):
            fail('Esta recuperación corresponde a otro responsable; nadie puede emitir su propio código.', 403)
        if target['status'] != 'ACTIVE':
            fail('La cuenta debe estar activa. Administración debe revisar su estado.')
        code = 'REC-' + secrets.token_hex(6).upper()
        expires = (datetime.now(timezone.utc) + timedelta(minutes=15)).isoformat()
        encrypted = recovery_cipher(create=True).encrypt(code.encode())
        c.execute("UPDATE recoveries SET status='CANCELLED',code_hash=NULL WHERE user_id=? AND id<>? AND status IN ('PENDING','ISSUED')", (target['id'], identity))
        c.execute('DELETE FROM recovery_code_vault WHERE recovery_id IN (SELECT id FROM recoveries WHERE user_id=?)', (target['id'],))
        c.execute("UPDATE recoveries SET status='ISSUED',code_hash=?,expires_at=? WHERE id=?", (password_hash(code), expires, identity))
        c.execute('INSERT INTO recovery_code_vault VALUES(?,?,?)', (identity, encrypted, now()))
        audit(c, actor, 'ISSUE_RECOVERY', identity)
    return {'code': code, 'expires_at': expires, 'username': target['username']}


@app.post('/api/recoveries/{identity}/view')
def view_recovery_code(identity: str, actor=Depends(current_user)):
    with connection(True) as c:
        r = get(c, 'recoveries', identity)
        target = get(c, 'users', r['user_id'])
        if not recovery_handler_allowed(c, actor, target):
            fail('Sólo el responsable autorizado puede consultar este código.', 403)
        if r['status'] != 'ISSUED' or not r['expires_at'] or parse_at(r['expires_at']) <= datetime.now(timezone.utc):
            fail('El código ya no está vigente. Genera uno nuevo.')
        stored = c.execute('SELECT encrypted_code FROM recovery_code_vault WHERE recovery_id=?', (identity,)).fetchone()
        if not stored:
            fail('El código fue generado en una versión anterior. Regénéralo para poder consultarlo.', 409)
        try:
            code = recovery_cipher().decrypt(stored['encrypted_code']).decode()
        except InvalidToken:
            fail('No se puede consultar el código con la clave disponible. Regénéralo.', 409)
        audit(c, actor, 'VIEW_RECOVERY_CODE', identity)
    return {'code': code, 'expires_at': r['expires_at'], 'username': target['username']}


@app.post('/api/recoveries/complete')
def complete_recovery(data: dict = Body(...)):
    password = text(data, 'password')
    if len(password) < 10:
        fail('Usa al menos 10 caracteres.')
    with connection(True) as c:
        u = c.execute('SELECT * FROM users WHERE username=? COLLATE NOCASE', (text(data, 'identity'),)).fetchone()
        requests = c.execute("SELECT * FROM recoveries WHERE user_id=? AND status='ISSUED' AND expires_at>?", (u['id'] if u else '', now())).fetchall()
        r = next((r for r in requests if password_matches(text(data, 'code'), r['code_hash'])), None)
        if not r:
            fail('Código inválido o vencido.')
        c.execute('UPDATE users SET password_hash=? WHERE id=?', (password_hash(password), u['id']))
        c.execute("UPDATE recoveries SET status='USED',code_hash=NULL WHERE id=?", (r['id'],))
        c.execute("UPDATE recoveries SET status='CANCELLED',code_hash=NULL WHERE user_id=? AND id<>? AND status IN ('PENDING','ISSUED')", (u['id'], r['id']))
        c.execute('DELETE FROM recovery_code_vault WHERE recovery_id IN (SELECT id FROM recoveries WHERE user_id=?)', (u['id'],))
        invalidate_sessions(c, u['id'])
        audit(c, dict(u), 'RESET_PASSWORD', u['id'])
    return {'ok': True}


@app.post('/api/delegations')
def delegate(data: dict = Body(...), actor=Depends(current_user)):
    if actor['role'] != 'MANAGER':
        fail('Sólo el Manager puede delegar su autoridad.', 403)
    start, end = now(), ''
    with connection(True) as c:
        lead = get(c, 'users', text(data, 'lead_id'))
        allowed = scope(c, actor)
        ids = sorted(allowed)
        if not lead['is_lead'] or lead['status'] != 'ACTIVE' or lead['manager_id']!=actor['id'] or not ids:
            fail('Selecciona un Lead activo de tu equipo. La delegación cubre todos tus Grupos.')
        # The write transaction serializes concurrent assignments. Changing backup
        # explicitly revokes the existing active/scheduled assignment atomically.
        outstanding = c.execute("SELECT * FROM delegations WHERE manager_id=? AND revoked_at IS NULL AND (end_at='' OR end_at>?)", (actor['id'], now())).fetchall()
        if outstanding and data.get('replace_current') is not True:
            fail('Ya hay una persona delegada. Usa la asignación del equipo para sustituirla.', 409)
        for previous in outstanding:
            c.execute('UPDATE delegations SET revoked_at=? WHERE id=?', (now(), previous['id']))
            audit(c, actor, 'REVOKE_DELEGATION', previous['id'], {'reason': 'REPLACED', 'new_lead_id': lead['id']})
        identity = uid()
        c.execute('INSERT INTO delegations VALUES(?,?,?,?,?,?,NULL)', (identity, actor['id'], lead['id'], json.dumps(ids), start, end))
        audit(c, actor, 'DELEGATE', identity)
    return {'id': identity}


@app.delete('/api/delegations/{identity}')
def revoke(identity: str, actor=Depends(current_user)):
    with connection(True) as c:
        d = get(c, 'delegations', identity)
        if d['manager_id'] != actor['id']:
            fail('Sólo el Manager que delegó puede revocar.', 403)
        if not d['revoked_at'] and (not d['end_at'] or parse_at(d['end_at']) > datetime.now(timezone.utc)):
            # Older versions could save multiple backups. Explicit revocation must
            # not reactivate an older row when the current one is revoked.
            targets = c.execute("SELECT id FROM delegations WHERE manager_id=? AND revoked_at IS NULL AND (end_at='' OR end_at>?)", (actor['id'], now())).fetchall()
        else:
            targets = [{'id': identity}]
        for target in targets:
            c.execute('UPDATE delegations SET revoked_at=? WHERE id=?', (now(), target['id']))
            audit(c, actor, 'REVOKE_DELEGATION', target['id'])
    return {'ok': True}


@app.post('/api/rules')
@app.put('/api/rules/{identity}')
def add_rule(data: dict = Body(...), identity: str = '', actor=Depends(current_user)):
    if actor['role'] != 'MANAGER':
        fail('Sólo el Manager crea reglas de Guardia.', 403)
    segments = data.get('segments', [])
    if not isinstance(segments, list) or not segments:
        fail('Agrega al menos un segmento de Guardia.')
    if len(segments) > 128:
        fail('La regla admite hasta 128 intervalos semanales.')
    normalized = []
    for s in segments:
        try:
            if int(s['day']) not in range(7):
                raise ValueError()
            start_time, end_time = time.fromisoformat(s['start']), time.fromisoformat(s['end'])
            if start_time.tzinfo or end_time.tzinfo or start_time.second or end_time.second or start_time.microsecond or end_time.microsecond:
                raise ValueError()
            normalized.append({'day': int(s['day']), 'start': start_time.strftime('%H:%M'), 'end': end_time.strftime('%H:%M')})
        except (KeyError, ValueError, TypeError):
            fail('Revisa los días y horas de la regla.')
    weekly=[]
    for segment in normalized:
        minute=lambda value:int(value[:2])*60+int(value[3:])
        start=segment['day']*1440+minute(segment['start'])
        duration=(minute(segment['end'])-minute(segment['start']))%1440 or 1440
        end=start+duration
        weekly.extend([(start,min(end,10080))]+([(0,end-10080)] if end>10080 else []))
    previous_end=-1
    for start,end in sorted(weekly):
        if start<previous_end:fail('Hay intervalos de Guardia traslapados, incluso entre días. Ajusta la cobertura antes de guardar.')
        previous_end=end
    try:
        effective = date.fromisoformat(text(data, 'effective_from')).isoformat()
    except ValueError:
        fail('Fecha de vigencia inválida.')
    with connection(True) as c:
        if identity:
            old=get(c,'rules',identity)
            if old['manager_id']!=actor['id']:
                fail('Sólo el Manager propietario puede editar esta regla.',403)
            c.execute('UPDATE rules SET name=?,effective_from=?,segments=? WHERE id=?',(text(data,'name'),effective,json.dumps(normalized),identity))
            audit(c,actor,'EDIT_RULE',identity,{'previous':old,'new_segments':normalized})
        else:
            identity=uid()
            c.execute('INSERT INTO rules VALUES(?,?,?,?,?)', (identity, actor['id'], text(data, 'name'), effective, json.dumps(normalized)))
            audit(c, actor, 'CREATE_RULE', identity)
    return {'id': identity}


@app.delete('/api/rules/{identity}')
def delete_rule(identity: str, actor=Depends(current_user)):
    if actor['role']!='MANAGER':
        fail('Sólo el Manager propietario elimina reglas de Guardia.',403)
    with connection(True) as c:
        rule=get(c,'rules',identity)
        if rule['manager_id']!=actor['id']:
            fail('Sólo el Manager propietario elimina esta regla.',403)
        c.execute('DELETE FROM rules WHERE id=?',(identity,))
        audit(c,actor,'DELETE_RULE',identity,{'previous':rule,'existing_requests_preserved':True})
    return {'ok':True}


def guard_revision_token(record):
    return hashlib.sha256(json.dumps(dict(record),sort_keys=True,ensure_ascii=False).encode()).hexdigest()


def guard_rule_changed(c,guard):
    snapshot=guard['rule_snapshot']
    if snapshot.get('recovery_fixed'):return False
    current=c.execute('SELECT * FROM rules WHERE id=?',(snapshot.get('id'),)).fetchone()
    if not current:return True
    def content(rule):
        segments=rule['segments'] if isinstance(rule['segments'],list) else json.loads(rule['segments'])
        return (rule['name'],rule['effective_from'],sorted((int(s['day']),s['start'],s['end']) for s in segments))
    return content(snapshot)!=content(dict(current))


@app.get('/api/guards/{identity}/revision')
def guard_revision_preview(identity:str,actor=Depends(current_user)):
    operational(actor)
    with connection() as c:
        guard=get(c,'guards',identity)
        if actor['is_lead'] or guard['user_id']!=actor['id']:fail('Sólo puedes editar tus Guardias.',403)
        if guard['status']=='CANCELLED':fail('La Guardia está cancelada.')
        manager=get(c,'groups',actor['group_id'])['manager_id']
        rules=[{**r,'segments':json.loads(r['segments']),'token':guard_revision_token(r)} for r in rows(c,'rules') if r['manager_id']==manager]
        return {'guard':{**guard,'rule_snapshot':json.loads(guard['rule_snapshot'])},'expected':guard_revision_token(guard),'rules':rules}


@app.put('/api/guards/{identity}')
def revise_guard(identity:str,data:dict=Body(...),actor=Depends(current_user)):
    operational(actor)
    if actor['is_lead']:fail('El Lead no solicita Guardias.',403)
    reason=text(data,'reason')
    if data.get('confirm') is not True:fail('Confirma el reenvío de la Guardia a aprobación.')
    try:a,b=date.fromisoformat(text(data,'start_date')),date.fromisoformat(text(data,'end_date'))
    except ValueError:fail('Fechas inválidas.')
    if b<a or (b-a).days>31:fail('La Guardia debe durar entre 1 y 32 días.')
    with connection(True) as c:
        guard=get(c,'guards',identity)
        if guard['user_id']!=actor['id']:fail('Sólo puedes editar tus Guardias.',403)
        if guard['status']=='CANCELLED':fail('La Guardia está cancelada.')
        if not hmac.compare_digest(str(data.get('expected','')),guard_revision_token(guard)):fail('La Guardia cambió. Vuelve a abrir la revisión.',409)
        rule=get(c,'rules',text(data,'rule_id'))
        if rule['manager_id']!=get(c,'groups',actor['group_id'])['manager_id']:fail('Elige una regla de tu Manager.',403)
        if not hmac.compare_digest(str(data.get('rule_token','')),guard_revision_token(rule)):fail('La regla cambió mientras la revisabas. Abre nuevamente el formulario.',409)
        if rule['effective_from']>a.isoformat():fail('El inicio es anterior a la vigencia de la regla seleccionada.')
        previous=json.loads(guard['rule_snapshot']);snapshot=dict(rule,segments=json.loads(rule['segments']))
        if previous.get('coverage_end_at'):snapshot['coverage_end_at']=previous['coverage_end_at']
        c.execute("UPDATE guards SET start_date=?,end_date=?,rule_snapshot=?,status='SUBMITTED',approved_by=NULL WHERE id=?",(str(a),str(b),json.dumps(snapshot),identity))
        audit(c,actor,'REVISE_GUARD',identity,{'reason':reason,'previous':guard,'new_snapshot':snapshot,'start_date':str(a),'end_date':str(b)})
    return {'id':identity,'status':'SUBMITTED'}


@app.post('/api/guards')
def add_guard(data: dict = Body(...), actor=Depends(current_user)):
    operational(actor)
    if actor['is_lead']:fail('El Lead registra OT; no puede solicitar Guardias.',403)
    try:
        a, b = date.fromisoformat(text(data, 'start_date')), date.fromisoformat(text(data, 'end_date'))
    except ValueError:
        fail('Fechas inválidas.')
    if b < a or (b - a).days > 31:
        fail('La Guardia debe durar entre 1 y 32 días.')
    identity = uid()
    with connection(True) as c:
        rule = get(c, 'rules', text(data, 'rule_id'))
        g = get(c, 'groups', actor['group_id'])
        if rule['manager_id'] != g['manager_id']:
            fail('Selecciona una regla del Manager de tu Grupo.')
        if rule['effective_from'] > a.isoformat():
            fail(f"La regla {rule['name']} aplica desde {rule['effective_from']}; tu Guardia empieza el {a.isoformat()}, antes de su vigencia.")
        snapshot = dict(rule, segments=json.loads(rule['segments']))
        c.execute('INSERT INTO guards VALUES(?,?,?,?,?,?,?,?)', (identity, actor['id'], str(a), str(b), json.dumps(snapshot), 'SUBMITTED', None, now()))
        capture_origin(c,identity,actor)
        audit(c, actor, 'REQUEST_GUARD', identity)
    return {'id': identity}


def sync_group_ot(c,source,decision=None,approver=None):
    owner=get(c,'users',source['user_id'])
    if source['capture_type']!='PLANNED' or source['execution'] not in ('SCHEDULED','RUNNING') or not owner['group_id'] or owner['is_lead']:return
    link=c.execute('SELECT * FROM ot_group_activities WHERE source_id=?',(source['id'],)).fetchone()
    if source['window_id'] and not link:return
    status=decision or source['status']
    if status=='APPROVED' and not (approver or source['approved_by']):status='SUBMITTED'
    snapshot={key:source[key] for key in ('reference','description','justification','impact','in_lieu','scheduled_start','scheduled_end')}
    if link:
        c.execute('UPDATE windows SET scheduled_start=?,scheduled_end=?,reference=?,description=?,status=? WHERE id=?',(source['scheduled_start'],source['scheduled_end'],source['reference'],source['description'],status,link['window_id']))
        c.execute('UPDATE ot_group_activities SET authorized_by=?,snapshot=? WHERE source_id=?',(approver or source['approved_by'] if status=='APPROVED' else None,json.dumps(snapshot),source['id']))
    else:
        wid=uid();c.execute('INSERT INTO windows VALUES(?,?,?,?,?,?,?,?,?)',(wid,owner['group_id'],owner['id'],source['scheduled_start'],source['scheduled_end'],source['reference'],source['description'],status,now()))
        c.execute('INSERT INTO ot_group_activities VALUES(?,?,?,?)',(source['id'],wid,(approver or source['approved_by']) if status=='APPROVED' else None,json.dumps(snapshot)))
        c.execute('UPDATE overtime SET window_id=? WHERE id=?',(wid,source['id']))


# Existing future plans gain the same group visibility; personal history is retained.
with connection(True) as _migration:
    for _plan in rows(_migration,'overtime'):
        if _plan['window_id'] is None and _plan['capture_type']=='PLANNED' and _plan['execution'] in ('SCHEDULED','RUNNING') and parse_at(_plan['scheduled_end'])>datetime.now(timezone.utc):sync_group_ot(_migration,_plan)


@app.post('/api/windows')
def add_window(data: dict = Body(...), actor=Depends(current_user)):
    fail('La programación se solicita desde Mi tiempo por cada Técnico; el responsable revisa y autoriza.', 403)


def classify_ot_period(start, end):
    """Classify on the server clock; in-progress work requires an actual Timer."""
    current = datetime.now(timezone.utc)
    if parse_at(end) <= current:
        return 'HISTORICAL'
    if parse_at(start) > current:
        return 'PLANNED'
    fail('El periodo está en curso. Inicia el Timer de OT de emergencia para registrar tiempo real, o captura la OT histórica cuando termine.')


def insert_ot(c, actor, data, window=None):
    identity = uid()
    capture = data.get('capture_type', 'PLANNED')
    if capture not in ('PLANNED', 'HISTORICAL', 'EMERGENCY'):
        fail('Tipo de OT inválido.')
    if window:
        capture = 'PLANNED'
        a, b = window['scheduled_start'], window['scheduled_end']
        ref, description = window['reference'], window['description']
    elif capture != 'EMERGENCY':
        a, b = interval(data, 'scheduled_start', 'scheduled_end')
        capture = classify_ot_period(a, b)
        ref, description = valid_reference(data), text(data, 'description')
    else:
        a = b = None
        ref = description = ''
    if capture == 'HISTORICAL' and parse_at(b) > datetime.now(timezone.utc):
        fail('La OT histórica debe haber finalizado.')
    actual_a, actual_b = (a, b) if capture == 'HISTORICAL' else (None, None)
    execution = 'COMPLETED' if capture == 'HISTORICAL' else 'SCHEDULED' if capture == 'PLANNED' else 'RUNNING'
    status = 'APPROVED' if window else 'DOCUMENTATION' if capture == 'EMERGENCY' else 'SUBMITTED'
    c.execute('INSERT INTO overtime VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
              (identity, actor['id'], window['id'] if window else None, capture, a, b, actual_a, actual_b, ref, description, text(data, 'justification', False), text(data, 'impact', False), text(data, 'in_lieu', False), status, execution, 0, None, now(), now()))
    capture_origin(c,identity,actor)
    return identity


@app.post('/api/overtime')
def add_overtime(data: dict = Body(...), actor=Depends(current_user)):
    operational(actor)
    if data.get('capture_type') == 'EMERGENCY':
        fail('Inicia la OT de emergencia con el Timer.')
    with connection(True) as c:
        identity = insert_ot(c, actor, data)
        sync_group_ot(c,get(c,'overtime',identity))
        audit(c, actor, 'CREATE_OT', identity)
    return {'id': identity}


@app.post('/api/windows/{identity}/join')
def join_window(identity: str, actor=Depends(current_user)):
    operational(actor)
    with connection(True) as c:
        w = get(c, 'windows', identity)
        if w['group_id'] != actor['group_id'] or w['status'] != 'APPROVED' or parse_at(w['scheduled_end']) <= datetime.now(timezone.utc):
            fail('Esta ventana no está disponible para tu Grupo.')
        existing=c.execute('SELECT id FROM overtime WHERE user_id=? AND window_id=?',(actor['id'],identity)).fetchone()
        if existing:
            if c.execute('SELECT 1 FROM ot_group_activities WHERE window_id=?',(identity,)).fetchone():return {'id':existing['id']}
            fail('Ya participas en esta actividad.',409)
        link=c.execute('SELECT * FROM ot_group_activities WHERE window_id=?',(identity,)).fetchone()
        snapshot=json.loads(link['snapshot']) if link else {}
        otid = insert_ot(c, actor, snapshot, w)
        if link:c.execute('UPDATE overtime SET approved_by=? WHERE id=?',(link['authorized_by'],otid))
        audit(c, actor, 'JOIN_SHARED_OT', otid)
    return {'id': otid}


def effective_ot_segments(c,o):
    correction=c.execute('SELECT segments FROM ot_time_corrections WHERE overtime_id=?',(o['id'],)).fetchone()
    if correction:return json.loads(correction['segments'])
    timers=[{'start_at':r['start_at'],'end_at':r['end_at']} for r in c.execute('SELECT start_at,end_at FROM timers WHERE overtime_id=? AND end_at IS NOT NULL ORDER BY start_at',(o['id'],))]
    if o['capture_type']=='PLANNED' and timers:return timers
    return [{'start_at':o['start_at'],'end_at':o['end_at']}] if o['start_at'] and o['end_at'] else []


def revise_actual_ot_time(c,actor,o,data):
    if 'actual_start' not in data and 'actual_end' not in data and 'actual_segments' not in data:return
    a,b=interval(data,'actual_start','actual_end')
    if parse_at(b)>datetime.now(timezone.utc):fail('El tiempo trabajado debe haber terminado; no puede incluir horas futuras.')
    previous=effective_ot_segments(c,o)
    supplied=data.get('actual_segments')
    if supplied is not None:
        if not isinstance(supplied,list) or not supplied or len(supplied)>128:fail('Revisa los tramos de trabajo.')
        corrected=[]
        for segment in supplied:
            if not isinstance(segment,dict):fail('Revisa los tramos de trabajo.')
            sa,sb=interval(segment,'start_at','end_at');corrected.append({'start_at':sa,'end_at':sb})
    elif o['capture_type']=='PLANNED' and len(previous)>1:
        corrected=[dict(s) for s in previous];corrected[0]['start_at']=a;corrected[-1]['end_at']=b
    else:corrected=[{'start_at':a,'end_at':b}]
    if parse_at(corrected[0]['start_at'])!=parse_at(a) or parse_at(corrected[-1]['end_at'])!=parse_at(b):fail('El inicio y fin deben coincidir con los extremos de los tramos.')
    end=None
    for segment in corrected:
        start,stop=parse_at(segment['start_at']),parse_at(segment['end_at'])
        if stop<=start or start<parse_at(a) or stop>parse_at(b) or end and start<end:fail('Los tramos deben ser positivos, estar ordenados y no traslaparse. Revisa también los horarios intermedios.')
        end=stop
    normalized=lambda parts:[(parse_at(p['start_at']),parse_at(p['end_at'])) for p in parts]
    if normalized(previous)==normalized(corrected):return
    reason=text(data,'time_change_reason')
    if len(reason)>2000:fail('Resume el motivo del cambio en 2000 caracteres.')
    c.execute('INSERT OR REPLACE INTO ot_time_corrections VALUES(?,?,?,?)',(o['id'],json.dumps(corrected),reason,now()))
    c.execute('UPDATE overtime SET start_at=?,end_at=? WHERE id=?',(a,b,o['id']))
    audit(c,actor,'CORRECT_OT_TIME',o['id'],{'reason':reason,'previous_start':o['start_at'],'previous_end':o['end_at'],'previous_segments':previous,'new_start':a,'new_end':b,'new_segments':corrected,'original_timers_preserved':True})


@app.put('/api/overtime/{identity}')
def edit_ot(identity: str, data: dict = Body(...), actor=Depends(current_user)):
    with connection(True) as c:
        o = get(c, 'overtime', identity)
        if o['user_id'] != actor['id']:
            fail('Sólo el Técnico propietario puede editar su OT.', 403)
        if o['execution'] in ('RUNNING','PAUSED'):
            fail('Detén o finaliza el Timer antes de documentar.')
        if 'expected_revision' in data:
            try:expected=int(data['expected_revision'])
            except (ValueError,TypeError):fail('Revisión de OT inválida.')
            if expected!=o['revision']:fail('La OT cambió desde que abriste el formulario. Revisa la versión actual antes de guardar.',409)
        if o['execution']=='SCHEDULED' and c.execute('SELECT 1 FROM ot_group_activities WHERE source_id=?',(identity,)).fetchone() and c.execute('SELECT 1 FROM overtime WHERE window_id=? AND id!=?',(o['window_id'],identity)).fetchone():fail('Hay integrantes inscritos en esta actividad. Conserva su ventana y registra una nueva OT si necesitas otra programación.',409)
        ref = valid_reference(data)
        if o['execution'] == 'SCHEDULED':
            a, b = interval(data, 'scheduled_start', 'scheduled_end')
            capture = classify_ot_period(a, b)
            c.execute('UPDATE overtime SET scheduled_start=?,scheduled_end=?,capture_type=?,execution=?,start_at=?,end_at=? WHERE id=?', (a, b, capture, 'COMPLETED' if capture=='HISTORICAL' else 'SCHEDULED', a if capture=='HISTORICAL' else None, b if capture=='HISTORICAL' else None, identity))
        if o['execution']=='COMPLETED':revise_actual_ot_time(c,actor,o,data)
        c.execute("UPDATE overtime SET reference=?,description=?,justification=?,impact=?,in_lieu=?,status='SUBMITTED',approved_by=NULL,revision=revision+1,updated_at=? WHERE id=?", (ref, text(data, 'description'), text(data, 'justification'), text(data, 'impact'), text(data, 'in_lieu'), now(), identity))
        updated=get(c,'overtime',identity)
        if updated['capture_type']=='PLANNED':sync_group_ot(c,updated)
        elif o['window_id'] and c.execute('SELECT 1 FROM ot_group_activities WHERE source_id=?',(identity,)).fetchone():c.execute("UPDATE windows SET status='CANCELLED' WHERE id=?",(o['window_id'],))
        audit(c, actor, 'EDIT_OT_REOPEN_APPROVAL', identity, {'previous_status': o['status'], 'revision': o['revision'] + 1})
    return {'ok': True}


@app.post('/api/timers/start')
def start_timer(data: dict = Body(...), actor=Depends(current_user)):
    operational(actor)
    with connection(True) as c:
        if c.execute('SELECT 1 FROM timers WHERE user_id=? AND end_at IS NULL', (actor['id'],)).fetchone():
            fail('Ya tienes un Timer activo.')
        otid = data.get('overtime_id')
        if otid:
            o = get(c, 'overtime', otid)
            if o['user_id'] != actor['id'] or o['status'] != 'APPROVED' or o['execution'] not in ('SCHEDULED','PAUSED'):
                fail('La OT debe ser tuya, estar aprobada y pendiente de iniciar.')
            if parse_at(o['scheduled_start']) > parse_at(now()):
                fail('El Timer estará disponible desde '+o['scheduled_start']+'.')
            if parse_at(o['scheduled_end']) <= parse_at(now()):
                fail('La ventana autorizada terminó. Solicita una extensión antes de iniciar o reanudar.')
        else:
            otid = insert_ot(c, actor, {'capture_type': 'EMERGENCY'})
        identity, stamp = uid(), now()
        if data.get('overtime_id') and parse_at(stamp)>=parse_at(o['scheduled_end']):
            fail('La ventana autorizada terminó. Solicita una extensión.')
        c.execute('INSERT INTO timers VALUES(?,?,?,?,NULL)', (identity, actor['id'], otid, stamp))
        c.execute("UPDATE overtime SET execution='RUNNING',start_at=COALESCE(start_at,?),end_at=NULL,updated_at=? WHERE id=?", (stamp, stamp, otid))
        audit(c, actor, 'START_TIMER', otid)
    return {'id': identity}


@app.post('/api/timers/{identity}/stop')
def stop_timer(identity: str, actor=Depends(current_user)):
    with connection(True) as c:
        t = get(c, 'timers', identity)
        if t['user_id'] != actor['id']:
            fail('Sólo el Técnico propietario puede detener el Timer.', 403)
        if t['end_at']:
            o=get(c,'overtime',t['overtime_id'])
            if o['execution']=='PAUSED':
                finish_paused(c,actor,o)
                return {'ok':True}
            fail('El Timer ya está detenido.')
        stamp = now()
        o = get(c, 'overtime', t['overtime_id'])
        if o['capture_type']=='PLANNED' and parse_at(stamp)>parse_at(o['scheduled_end']):
            stamp=o['scheduled_end']
        c.execute('UPDATE timers SET end_at=? WHERE id=?', (stamp, identity))
        status = 'DOCUMENTATION' if o['capture_type'] == 'EMERGENCY' else 'SUBMITTED'
        c.execute("UPDATE overtime SET end_at=?,execution='COMPLETED',status=?,approved_by=NULL,updated_at=? WHERE id=?", (stamp, status, stamp, o['id']))
        cancel_pending_extensions(c,actor,o['id'],'El Técnico detuvo el Timer.')
        audit(c, actor, 'STOP_TIMER', o['id'])
    return {'ok': True}


@app.post('/api/overtime/{identity}/finish')
def finish_overtime(identity: str, actor=Depends(current_user)):
    with connection(True) as c:
        o=get(c,'overtime',identity)
        if o['user_id']!=actor['id']:
            fail('Sólo el Técnico propietario finaliza su OT.',403)
        if o['execution']!='PAUSED':
            fail('La OT no está pausada.')
        finish_paused(c,actor,o)
    return {'ok':True}


@app.post('/api/overtime/{identity}/extensions')
def request_extension(identity: str, data: dict=Body(...), actor=Depends(current_user)):
    operational(actor)
    end=parse_at(text(data,'requested_end'))
    with connection(True) as c:
        o=get(c,'overtime',identity)
        if o['user_id']!=actor['id']:
            fail('Sólo el Técnico propietario solicita la extensión.',403)
        if o['capture_type']!='PLANNED' or o['status']!='APPROVED' or o['execution'] not in ('RUNNING','PAUSED') or not o['start_at']:
            fail('La extensión se solicita desde un Timer que ya inició o está pausado.')
        if end<=parse_at(o['scheduled_end']) or end<=parse_at(now()) or end-parse_at(o['scheduled_start'])>timedelta(days=31):
            fail('El nuevo fin debe ser futuro, posterior al autorizado y dentro de 31 días del inicio.')
        if c.execute("SELECT 1 FROM ot_extensions WHERE overtime_id=? AND status='SUBMITTED'",(identity,)).fetchone():
            fail('Ya existe una extensión pendiente para esta OT.',409)
        ext=uid()
        c.execute('INSERT INTO ot_extensions VALUES(?,?,?,?,?,?,?,NULL,?,?,NULL)',(ext,identity,actor['id'],o['scheduled_end'],end.isoformat(),text(data,'reason'),'SUBMITTED','',now()))
        audit(c,actor,'REQUEST_OT_EXTENSION',identity,{'extension_id':ext,'from':o['scheduled_end'],'to':end.isoformat(),'reason':text(data,'reason')})
    return {'id':ext}


@app.post('/api/extensions/{identity}/decision')
def decide_extension(identity: str, data: dict=Body(...), actor=Depends(current_user)):
    decision=text(data,'decision')
    if decision not in ('APPROVED','REJECTED'):
        fail('Decisión inválida.')
    reason=text(data,'reason',False)
    if decision=='REJECTED' and not reason:
        fail('Agrega el motivo del rechazo.')
    with connection(True) as c:
        e=get(c,'ot_extensions',identity);o=get(c,'overtime',e['overtime_id']);target=get(c,'users',o['user_id'])
        authorized_target(c,actor,target,approval=True)
        if actor['role']=='ADMIN' or target['status']!='ACTIVE':
            fail('La extensión corresponde al Manager o Lead delegado de un Técnico activo.',403)
        if e['status']!='SUBMITTED' or o['status']!='APPROVED' or o['execution'] not in ('SCHEDULED','RUNNING','PAUSED'):
            fail('La extensión ya no está pendiente.',409)
        if o['scheduled_end']!=e['previous_end']:
            fail('La ventana autorizada cambió. Solicita una extensión nueva.',409)
        if decision=='APPROVED':
            if parse_at(e['requested_end'])<=parse_at(now()):
                fail('La extensión solicitada ya terminó. Se necesita una solicitud nueva.')
            c.execute('UPDATE overtime SET scheduled_end=?,updated_at=? WHERE id=?',(e['requested_end'],now(),o['id']))
        c.execute('UPDATE ot_extensions SET status=?,decided_by=?,decision_reason=?,decided_at=? WHERE id=?',(decision,actor['id'],reason,now(),identity))
        audit(c,actor,'DECIDE_OT_EXTENSION',o['id'],{'extension_id':identity,'decision':decision,'reason':reason,'from':e['previous_end'],'to':e['requested_end']})
    return {'ok':True}


@app.delete('/api/extensions/{identity}')
def cancel_extension(identity: str, actor=Depends(current_user)):
    with connection(True) as c:
        e=get(c,'ot_extensions',identity)
        if e['requested_by']!=actor['id']:
            fail('Sólo el solicitante cancela la extensión.',403)
        if e['status']!='SUBMITTED':
            fail('La extensión ya no está pendiente.',409)
        cancel_pending_extensions(c,actor,e['overtime_id'],'Cancelada por el Técnico.')
    return {'ok':True}


@app.post('/api/approvals/{kind}/{identity}')
def decide(kind: str, identity: str, data: dict = Body(...), actor=Depends(current_user)):
    table = {'ot': 'overtime', 'guard': 'guards'}.get(kind)
    if not table:
        fail('Tipo de registro inválido.')
    decision = text(data, 'decision')
    if decision not in ('APPROVED', 'REJECTED'):
        fail('Decisión inválida.')
    with connection(True) as c:
        r = get(c, table, identity)
        authorized_target(c, actor, get(c, 'users', r['user_id']), approval=True)
        if actor['role'] == 'ADMIN':
            fail('Las aprobaciones corresponden al Manager o Lead delegado.', 403)
        if r['status'] != 'SUBMITTED' or kind == 'ot' and r['execution'] in ('RUNNING','PAUSED'):
            fail('El registro no está pendiente de aprobación.')
        if decision == 'REJECTED' and not text(data, 'reason', False):
            fail('Agrega el motivo del rechazo.')
        c.execute(f'UPDATE {table} SET status=?,approved_by=? WHERE id=?', (decision, actor['id'], identity))
        if kind=='ot' and r['execution']=='SCHEDULED':sync_group_ot(c,get(c,'overtime',identity),decision,actor['id'])
        audit(c, actor, decision, identity, {'kind': kind, 'reason': text(data, 'reason', False)})
    return {'ok': True}


# All accounting/export/capture views share this implementation, including cross-midnight splits.
def guard_coverage(guard):
    snapshot=guard['rule_snapshot'];snapshot=json.loads(snapshot) if isinstance(snapshot,str) else snapshot
    if snapshot.get('recovered_intervals'):
        cutoff=parse_at(snapshot['coverage_end_at']) if snapshot.get('coverage_end_at') else None
        return [(parse_at(p['start_at']).astimezone(TZ),min(parse_at(p['end_at']),cutoff).astimezone(TZ) if cutoff else parse_at(p['end_at']).astimezone(TZ)) for p in snapshot['recovered_intervals'] if not cutoff or parse_at(p['start_at'])<cutoff]
    intervals=[];day,last=date.fromisoformat(guard['start_date']),date.fromisoformat(guard['end_date'])
    while day<=last:
        for segment in snapshot.get('segments',[]):
            if int(segment['day'])!=day.weekday():continue
            start=datetime.combine(day,time.fromisoformat(segment['start']),TZ);end=datetime.combine(day,time.fromisoformat(segment['end']),TZ)
            if end<=start:end+=timedelta(days=1)
            if snapshot.get('coverage_end_at'):end=min(end,parse_at(snapshot['coverage_end_at']).astimezone(TZ))
            if end>start:intervals.append((start,end))
        day+=timedelta(days=1)
    return intervals


def guard_timeline_days(guard):
    """Daily coverage from immutable rules or exact recovered intervals."""
    intervals=guard_coverage(guard)
    days={}
    for start,end in union_intervals(intervals):
        while start<end:
            stop=min(end,datetime.combine(start.date()+timedelta(days=1),time(),TZ))
            days.setdefault(start.date().isoformat(),[]).append({'start_at':start.isoformat(),'end_at':stop.isoformat()});start=stop
    return [{'date':day,'start_at':parts[0]['start_at'],'end_at':parts[-1]['end_at'],'segments':parts} for day,parts in sorted(days.items())]


def union_intervals(intervals):
    result = []
    for a, b in sorted(intervals):
        if b <= a:
            continue
        if result and a <= result[-1][1]:
            result[-1] = (result[-1][0], max(b, result[-1][1]))
        else:
            result.append((a, b))
    return result


def subtract_intervals(a, b, cuts):
    parts = [(a, b)]
    for ca, cb in union_intervals(cuts):
        remaining = []
        for pa, pb in parts:
            if cb <= pa or ca >= pb:
                remaining.append((pa, pb))
            else:
                if pa < ca:
                    remaining.append((pa, ca))
                if cb < pb:
                    remaining.append((cb, pb))
        parts = remaining
    return parts


def accounting(c, visible, month=None, record_ids=None):
    if month:
        try:
            first = date.fromisoformat(month + '-01')
        except ValueError:
            fail('Mes inválido.')
        lower = datetime.combine(first, time(), TZ)
        upper = datetime.combine((first.replace(day=28) + timedelta(days=4)).replace(day=1), time(), TZ)
    people = {u['id']: u for u in rows(c, 'users')}
    approved = [o for o in rows(c, 'overtime') if o['user_id'] in visible and o['status'] == 'APPROVED' and o['execution'] == 'COMPLETED' and o['start_at'] and o['end_at']]
    all_approved=approved
    approved=[o for o in approved if record_ids is None or o['id'] in record_ids]
    ots = {}
    segments={}
    for t in rows(c,'timers'):
        if t['end_at']:
            segments.setdefault(t['overtime_id'],[]).append((parse_at(t['start_at']),parse_at(t['end_at'])))
    corrections={r['overtime_id']:json.loads(r['segments']) for r in rows(c,'ot_time_corrections')}
    def actual_parts(o):
        if o['id'] in corrections:return [(parse_at(p['start_at']),parse_at(p['end_at'])) for p in corrections[o['id']]]
        return union_intervals(segments.get(o['id'],[(parse_at(o['start_at']),parse_at(o['end_at']))])) if o['capture_type']=='PLANNED' else [(parse_at(o['start_at']),parse_at(o['end_at']))]
    for o in all_approved:
        ots.setdefault(o['user_id'], []).extend(actual_parts(o))
    output = []

    def append_parts(user_id, kind, a, b, ref, description, entity_id, capture_type=''):
        a, b = a.astimezone(TZ), b.astimezone(TZ)
        if month:
            a, b = max(a, lower), min(b, upper)
        while a < b:
            midnight = datetime.combine(a.date() + timedelta(days=1), time(), TZ)
            stop = min(b, midnight)
            minutes = round((stop.astimezone(timezone.utc) - a.astimezone(timezone.utc)).total_seconds() / 60, 4)
            if minutes:
                p = people[user_id]
                output.append({'id': entity_id, 'user_id': user_id, 'employee_id': p['employee_id'], 'name': f"{p['first']} {p['last']}", 'group_id': p['group_id'], 'kind': kind, 'capture_type': capture_type if kind == 'OT' else 'GUARD', 'date': str(a.date()), 'start': a.strftime('%H:%M'), 'end': stop.strftime('%H:%M'), 'start_at': a.isoformat(), 'end_at': stop.isoformat(), 'minutes': minutes, 'hours': round(minutes / 60, 4), 'reference': ref, 'description': description})
            a = stop

    # Overlapping approved OT is counted once, even when two different records overlap.
    for user_id, intervals in ots.items():
        prior = []
        for o in sorted([o for o in approved if o['user_id'] == user_id], key=lambda o: (o['start_at'], o['id'])):
            for a,b in actual_parts(o):
                for pa, pb in subtract_intervals(a, b, prior):
                    append_parts(user_id, 'OT', pa, pb, o['reference'], o['description'], o['id'], o['capture_type'])
                prior.append((a, b))
    guard_intervals = {}
    for g in rows(c, 'guards'):
        if g['user_id'] not in visible or g['status'] != 'APPROVED' or record_ids is not None and g['id'] not in record_ids:
            continue
        guard_intervals.setdefault(g['user_id'], []).extend(guard_coverage(g))
    for user_id, intervals in guard_intervals.items():
        for a, b in union_intervals(intervals):
            for pa, pb in subtract_intervals(a, b, ots.get(user_id, [])):
                append_parts(user_id, 'GUARD', pa, pb, '', 'Guardia neta', '')
    return sorted(output, key=lambda r: (r['date'], r['start'], r['name']))


def selected_users(c, actor, user_id=None, group_id=None):
    allowed = scope(c, actor)
    visible = visible_users(c,actor)
    if actor['role'] == 'ADMIN':
        visible = {u['id'] for u in rows(c, 'users') if u['role'] == 'USER'}
    if user_id:
        if user_id not in visible:
            fail('El Técnico no está en tu alcance.', 403)
        visible &= {user_id}
    if group_id:
        visible &= {u['id'] for u in rows(c, 'users') if u['group_id'] == group_id}
    return visible


def claim_accounting(c,visible,month):
    try:first=date.fromisoformat(month+'-01')
    except ValueError:fail('Mes inválido.')
    last=(first.replace(day=28)+timedelta(days=4)).replace(day=1)-timedelta(days=1)
    start=first-timedelta(days=(first.weekday()-5)%7)
    end=last+timedelta(days=(4-last.weekday())%7)
    months={start.strftime('%Y-%m'),month,end.strftime('%Y-%m')}
    detail=[row for key in sorted(months) for row in accounting(c,visible,key) if start.isoformat()<=row['date']<=end.isoformat()]
    detail.sort(key=lambda r:(r['date'],r['start'],r['kind'],r['id']))
    total=lambda kind,selected:round(sum(r['minutes'] for r in selected if r['kind']==kind),4)
    monthly=[r for r in detail if r['date'].startswith(month)]
    return {'rows':detail,'ot_minutes':total('OT',detail),'guard_minutes':total('GUARD',detail),'month_ot_minutes':total('OT',monthly),'month_guard_minutes':total('GUARD',monthly),'period_start':start.isoformat(),'period_end':end.isoformat(),'timezone':str(TZ),'view':'claim','basis':'APPROVED_ACTUAL'}


@app.get('/api/accounting')
def capture(month: str, user_id: str | None = None, group_id: str | None = None, view: str | None = None, actor=Depends(current_user)):
    with connection() as c:
        visible=selected_users(c,actor,user_id,group_id)
        if view=='claim':return claim_accounting(c,visible,month)
        if view:fail('Vista de captura inválida.')
        result = accounting(c, visible, month)
    return {'rows': result, 'ot_minutes': round(sum(r['minutes'] for r in result if r['kind'] == 'OT'), 4), 'guard_minutes': round(sum(r['minutes'] for r in result if r['kind'] == 'GUARD'), 4), 'timezone': str(TZ)}


ANALYTICS_TYPES = ('EMERGENCY', 'PLANNED', 'HISTORICAL', 'GUARD')


def analytics_months(month, months=None):
    selected=sorted(set((months.split(',') if months else [month])+[month]))
    if len(selected)>24:
        fail('Selecciona hasta 24 meses para la comparativa.')
    for item in selected:
        try:
            if len(item)!=7 or item[4]!='-':
                raise ValueError()
            date.fromisoformat(item+'-01')
        except ValueError:
            fail('Selecciona meses válidos en formato AAAA-MM.')
    return selected


def analytics_records(c, visible, selected_months, query='', status='', record_type=''):
    def normalized(value):
        return ''.join(ch for ch in unicodedata.normalize('NFD',str(value)) if not unicodedata.combining(ch)).casefold()
    if len(query)>500:fail('La búsqueda admite hasta 500 caracteres.')
    tokens=normalized(query).split()
    if status and status not in ('SUBMITTED','APPROVED','REJECTED','DOCUMENTATION','RUNNING','PAUSED','CANCELLED'):
        fail('Estado de búsqueda inválido.')
    if record_type and record_type not in (*ANALYTICS_TYPES,'OT'):
        fail('Tipo de búsqueda inválido.')
    people={p['id']:p for p in rows(c,'users')};groups={g['id']:g for g in rows(c,'groups')}
    periods=[]
    for month in selected_months or []:
        first=date.fromisoformat(month+'-01');last=(first.replace(day=28)+timedelta(days=4)).replace(day=1)
        periods.append((datetime.combine(first,time(),TZ),datetime.combine(last,time(),TZ)))
    extensions={}
    for e in rows(c,'ot_extensions'):
        extensions.setdefault(e['overtime_id'],[]).append(e['reason']+' '+e['decision_reason']+' extensión')
    result=[]
    for kind,table in [('ot','overtime'),('guard','guards')]:
        for r in rows(c,table):
            if r['user_id'] not in visible:
                continue
            p=people[r['user_id']];g=groups.get(p['group_id'],{})
            category=r['capture_type'] if kind=='ot' else 'GUARD'
            effective_status=r.get('execution') if r.get('execution') in ('RUNNING','PAUSED') else r['status']
            if status and effective_status!=status:
                continue
            if record_type and not (category==record_type or record_type=='OT' and category!='GUARD'):
                continue
            if kind=='ot':
                start=r['start_at'] or r['scheduled_start'] or r['created_at'];end=r['end_at'] or r['scheduled_end'] or now()
                reference=r['reference'];description=r['description']
            else:
                snapshot=json.loads(r['rule_snapshot']);reference=snapshot['name'];description='Guardia'
                intervals=guard_coverage(r)
                start=min(a for a,b in intervals).isoformat() if intervals else r['start_date']+'T00:00:00'
                end=max(b for a,b in intervals).isoformat() if intervals else (date.fromisoformat(r['end_date'])+timedelta(days=1)).isoformat()+'T00:00:00'
            a,b=parse_at(start),parse_at(end)
            if selected_months is not None and not any(a<upper and (b>lower or a==b and a>=lower) for lower,upper in periods):continue
            aliases={'PLANNED':'OT programada','HISTORICAL':'OT histórica','EMERGENCY':'OT emergente emergencia','GUARD':'Guardia'}
            text=' '.join(str(v or '') for v in [r['id'],reference,description,r.get('justification'),p['first'],p['last'],p['employee_id'],g.get('name'),g.get('code'),g.get('display_code'),start,end,aliases[category],*extensions.get(r['id'],[])])
            haystack=normalized(text)
            if not all(token in haystack for token in tokens):continue
            result.append({'id':r['id'],'kind':kind,'capture_type':category,'user_id':p['id'],'name':p['first']+' '+p['last'],'employee_id':p['employee_id'],'group_id':p['group_id'],'group_name':g.get('name','Lead del Manager' if p['is_lead'] else ''),'reference':reference,'description':description,'from_at':start,'to_at':end,'status':r['status'],'display_status':effective_status,'execution':r.get('execution'),'created_at':r['created_at']})
    return sorted(result,key=lambda r:(r['created_at'],r['id']),reverse=True)


def analytics_report(c, actor, month, user_id=None, group_id=None, personal=False, query='', status='', record_type='', months=None, period='MONTH'):
    """Aggregate accounting rows with separate personal/team authorization."""
    if personal:
        operational(actor)
        allowed_groups = {actor['group_id']} if actor['group_id'] else set()
        visible = {actor['id']}
        user_id, group_id = actor['id'], None
    else:
        allowed_groups = scope(c, actor)
        if actor['role'] != 'MANAGER' and not (actor['role'] == 'USER' and allowed_groups):
            fail('La analítica del equipo requiere un Manager o un Lead con delegación vigente.', 403)
        if group_id and group_id not in allowed_groups:
            fail('El Grupo no está en tu alcance.', 403)
        visible = {p['id'] for p in rows(c,'users') if p['role']=='USER' and p['id'] in visible_users(c,actor)}
        if user_id:
            if user_id not in visible:
                fail('El Técnico no está en tu alcance delegado.', 403)
            visible &= {user_id}
        if group_id:
            visible &= {p['id'] for p in rows(c, 'users') if p['group_id'] == group_id}
    try:
        if len(month) != 7 or month[4] != '-':
            raise ValueError()
        date.fromisoformat(month + '-01')
    except ValueError:
        fail('Selecciona un mes válido en formato AAAA-MM.')
    if period not in ('MONTH','ALL'):
        fail('Periodo de búsqueda inválido.')
    selected_months=analytics_months(month,months)
    matched=analytics_records(c,visible,None if period=='ALL' else selected_months,query,status,record_type)
    matching_ids={r['id'] for r in matched}
    if period=='ALL':
        detail=accounting(c,visible,None,matching_ids)
        found=set(row['date'][:7] for row in detail)
        for r in matched:
            a=parse_at(r['from_at']).astimezone(TZ);b=parse_at(r['to_at']).astimezone(TZ)
            cursor=a.date().replace(day=1);last=(b-timedelta(microseconds=1) if b>a else b).date().replace(day=1)
            while cursor<=last:
                found.add(cursor.strftime('%Y-%m'));cursor=(cursor.replace(day=28)+timedelta(days=4)).replace(day=1)
        selected_months=sorted(found) or [month]
        period_rows={key:[r for r in detail if r['date'].startswith(key)] for key in selected_months}
    else:
        period_rows={key:accounting(c,visible,key,matching_ids) for key in selected_months}
        detail=[row for key in selected_months for row in period_rows[key]]
    if query or status or record_type:
        visible &= {r['user_id'] for r in matched}
    people = [p for p in rows(c, 'users') if p['id'] in visible]
    shown_groups = [g for g in rows(c, 'groups') if g['id'] in allowed_groups and (not group_id or g['id'] == group_id)]
    if user_id or query or status or record_type:
        shown_groups = [g for g in shown_groups if any(p['group_id'] == g['id'] for p in people)]
    group_names = {g['id']: g['name'] for g in shown_groups}

    def empty_totals():
        return {**dict.fromkeys(ANALYTICS_TYPES, 0), 'ot_minutes': 0, 'total_minutes': 0}

    def add(target, record):
        category = record['capture_type']
        if category not in ANALYTICS_TYPES:
            fail('El registro tiene un tipo de captura no reconocido.', 409)
        target[category] += record['minutes']
        target['total_minutes'] += record['minutes']
        if category != 'GUARD':
            target['ot_minutes'] += record['minutes']

    def rounded(totals):
        return {key: round(value, 4) for key, value in totals.items()}

    summary = empty_totals()
    by_group = {g['id']: {'id': g['id'], 'code': g.get('display_code',g['code']), 'name': g['name'], 'user_count': sum(p['group_id'] == g['id'] for p in people), 'minutes': empty_totals()} for g in shown_groups}
    by_user = {p['id']: {'id': p['id'], 'name': f"{p['first']} {p['last']}", 'employee_id': p['employee_id'], 'group_id': p['group_id'], 'group_name': group_names.get(p['group_id'], 'Lead del Manager' if p['is_lead'] else ''), 'minutes': empty_totals()} for p in people}
    by_day = {}
    for record in detail:
        record['group_name'] = group_names.get(record['group_id'], 'Lead del Manager')
        add(summary, record)
        add(by_user[record['user_id']]['minutes'], record)
        if record['group_id'] in by_group:
            add(by_group[record['group_id']]['minutes'], record)
        add(by_day.setdefault(record['date'], empty_totals()), record)
    for item in [*by_group.values(), *by_user.values()]:
        item['minutes'] = rounded(item['minutes'])
    comparisons=[]
    for key in selected_months:
        totals=empty_totals()
        for row in period_rows[key]:add(totals,row)
        comparisons.append({'month':key,'minutes':rounded(totals)})
    hourly=[{'hour':h,'minutes':empty_totals()} for h in range(24)]
    for record in detail:
        cursor=parse_at(record['start_at']).astimezone(TZ)
        stop=parse_at(record['end_at']).astimezone(TZ)
        while cursor<stop:
            boundary=cursor.replace(minute=0,second=0,microsecond=0)+timedelta(hours=1)
            end=min(boundary,stop)
            add(hourly[cursor.hour]['minutes'],{**record,'minutes':(end-cursor).total_seconds()/60})
            cursor=end
    for item in hourly:item['minutes']=rounded(item['minutes'])
    counts={}
    for record in matched:counts[record['display_status']]=counts.get(record['display_status'],0)+1
    return {'month': month, 'months':selected_months,'comparison':comparisons,'records':matched,'record_counts':counts,'record_count':len(matched),'timezone': str(TZ), 'basis': 'APPROVED_ACTUAL', 'filters': {'user_id': user_id, 'group_id': group_id,'q':query,'status':status,'record_type':record_type,'period':period},
            'mode': 'PERSONAL' if personal else 'TEAM', 'summary': rounded(summary), 'hourly':hourly, 'as_of':datetime.now(TZ).isoformat(),
            'groups': sorted(by_group.values(), key=lambda g: (-g['minutes']['total_minutes'], g['name'])),
            'users': sorted(by_user.values(), key=lambda u: (-u['minutes']['total_minutes'], u['name'])),
            'days': [{'date': day, 'minutes': rounded(values)} for day, values in sorted(by_day.items())], 'rows': detail}


@app.get('/api/analytics')
def analytics(month: str, user_id: str | None = None, group_id: str | None = None, q: str='', status: str='', record_type: str='', months: str | None=None, period: str='MONTH', actor=Depends(current_user)):
    with connection() as c:
        return analytics_report(c, actor, month, user_id, group_id, query=q, status=status, record_type=record_type, months=months, period=period)


def personal_analytics_query(request):
    if any(key in request.query_params for key in ('user_id', 'group_id')):
        fail('Mi resumen sólo consulta al propietario de la sesión; no admite filtros de otras personas o Grupos.')


@app.get('/api/my-analytics')
def my_analytics(request: Request, month: str, q: str='', status: str='', record_type: str='', months: str | None=None, period: str='MONTH', actor=Depends(current_user)):
    personal_analytics_query(request)
    with connection() as c:
        return analytics_report(c, actor, month, personal=True, query=q, status=status, record_type=record_type, months=months, period=period)


@app.get('/api/analytics.xlsx')
def analytics_excel(month: str, user_id: str | None = None, group_id: str | None = None, q: str='', status: str='', record_type: str='', months: str | None=None, period: str='MONTH', actor=Depends(current_user)):
    with connection() as c:
        report = analytics_report(c, actor, month, user_id, group_id, query=q, status=status, record_type=record_type, months=months, period=period)
    return analytics_workbook(report)


@app.get('/api/my-analytics.xlsx')
def my_analytics_excel(request: Request, month: str, q: str='', status: str='', record_type: str='', months: str | None=None, period: str='MONTH', actor=Depends(current_user)):
    personal_analytics_query(request)
    with connection() as c:
        report = analytics_report(c, actor, month, personal=True, query=q, status=status, record_type=record_type, months=months, period=period)
    return analytics_workbook(report)


def analytics_workbook(report):
    month = report['month']
    wb = Workbook()
    safe = lambda value: "'" + str(value) if str(value).startswith(('=', '+', '-', '@')) else value
    labels = ['OT de emergencia (h)', 'OT programada (h)', 'OT histórica (h)', 'Guardia neta (h)', 'Total aprobado (h)']
    values = lambda totals: [round(totals[key] / 60, 4) for key in (*ANALYTICS_TYPES, 'total_minutes')]
    summary = wb.active
    summary.title = 'Resumen'
    summary.append(['Periodos' if len(report['months'])>1 else 'Mes', ', '.join(report['months'])]); summary.append(['Zona horaria', report['timezone']])
    summary.append(['Criterio', 'Tiempo real aprobado; OT única y Guardia neta. Grupo actual del Técnico.'])
    summary.append(['Tipo', 'Horas'])
    for label, amount in zip(labels, values(report['summary'])):
        summary.append([label, amount])
    groups = wb.create_sheet('Por grupo'); groups.append(['Código', 'Grupo', 'Técnicos', *labels])
    for group in report['groups']:
        groups.append([safe(group['code']), safe(group['name']), group['user_count'], *values(group['minutes'])])
    people = wb.create_sheet('Por técnico'); people.append(['Employee ID', 'Técnico', 'Grupo', *labels])
    for user in report['users']:
        people.append([safe(user['employee_id']), safe(user['name']), safe(user['group_name']), *values(user['minutes'])])
    detail = wb.create_sheet('Detalle'); detail.append(['Fecha', 'Employee ID', 'Técnico', 'Grupo', 'Tipo', 'Inicio', 'Fin', 'Horas', 'Referencia'])
    names = dict(zip(ANALYTICS_TYPES, ['OT de emergencia', 'OT programada', 'OT histórica', 'Guardia neta']))
    for row in report['rows']:
        detail.append([row['date'], safe(row['employee_id']), safe(row['name']), safe(row['group_name']), names[row['capture_type']], row['start'], row['end'], row['hours'], safe(row['reference'])])
    records=wb.create_sheet('Registros');records.append(['Employee ID','Técnico','Grupo','Tipo','Referencia','Descripción','Inicio','Fin','Estado'])
    for r in report['records']:
        records.append([safe(r['employee_id']),safe(r['name']),safe(r['group_name']),r['capture_type'],safe(r['reference']),safe(r['description']),r['from_at'],r['to_at'],r['display_status']])
    summary.append(['Búsqueda',safe(report['filters']['q'])]);summary.append(['Periodos',', '.join(report['months'])]);summary.append(['Registros encontrados',report['record_count']])
    if len(report['comparison'])>1:
        comparison=wb.create_sheet('Comparativa');comparison.append(['Mes',*labels])
        for item in report['comparison']:comparison.append([item['month'],*values(item['minutes'])])
    hourly=wb.create_sheet('Por hora');hourly.append(['Franja horaria',*labels])
    for item in report['hourly']:hourly.append([f"{item['hour']:02}:00–{item['hour']+1:02}:00",*values(item['minutes'])])
    for sheet in wb:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1C5185')
        sheet.freeze_panes = 'A2'
        if sheet.title != 'Resumen':
            sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 25
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '0.00'
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    period_label=report['months'][0]+'_'+report['months'][-1] if len(report['months'])>1 else month
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="Operational_Pulse_Analitica_{period_label}.xlsx"'})


@app.get('/api/export.xlsx')
def export_excel(month: str, user_id: str | None = None, group_id: str | None = None, kind: str | None = None, view: str | None = None, actor=Depends(current_user)):
    with connection() as c:
        visible=selected_users(c,actor,user_id,group_id)
        if view and view!='claim':fail('Vista de captura inválida.')
        claim=claim_accounting(c,visible,month) if view=='claim' else None
        result = claim['rows'] if claim else accounting(c,visible,month)
        if kind:
            if kind not in ('OT', 'GUARD'):
                fail('Tipo de exportación inválido.')
            result = [r for r in result if r['kind'] == kind]
        groups = {g['id']: g for g in rows(c, 'groups')}
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tiempo aprobado'
    ws.append(['Fecha', 'Employee ID', 'Técnico', 'Grupo', 'Tipo', 'Inicio', 'Fin', 'Duración HH:MM', 'Horas decimales', 'Referencia', 'Descripción'])
    for r in result:
        mins = round(r['minutes'])
        # Prevent spreadsheet formula execution from user-supplied labels.
        safe = lambda v: "'" + str(v) if str(v).startswith(('=', '+', '-', '@')) else v
        ws.append([r['date'], safe(r['employee_id']), safe(r['name']), safe(groups.get(r['group_id'], {}).get('name', '')), 'Onsite Overtime' if r['kind'] == 'OT' else 'Stand-by', r['start'], r['end'], f'{mins // 60:02}:{mins % 60:02}', r['hours'], safe(r['reference']), safe(r['description'])])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='173E48')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 24
    summary = wb.create_sheet('Resumen')
    summary.append(['Mes', month])
    summary.append(['Zona horaria', str(TZ)])
    summary.append(['Horas OT', sum(r['hours'] for r in result if r['kind'] == 'OT')])
    summary.append(['Horas Guardia neta', sum(r['hours'] for r in result if r['kind'] == 'GUARD')])
    summary.append(['Criterio', 'Sólo aprobados; OT única; Guardia menos OT; división por día local.'])
    if claim:
        from decimal import Decimal, ROUND_HALF_UP
        matrix=wb.create_sheet('Claim semanal',0)
        matrix.append(['Semana desde','Concepto','Total','Sábado','Domingo','Lunes','Martes','Miércoles','Jueves','Viernes'])
        cursor=date.fromisoformat(claim['period_start']);last=date.fromisoformat(claim['period_end'])
        while cursor<=last:
            for key,label in [('GUARD','Stand By'),('OT','OT')]:
                values=[(Decimal(str(sum(r['minutes'] for r in result if r['kind']==key and r['date']==(cursor+timedelta(days=i)).isoformat())))/Decimal(60)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP) for i in range(7)]
                matrix.append([cursor.isoformat(),label,float(sum(values)),*[float(v) for v in values]])
            cursor+=timedelta(days=7)
        for cell in matrix[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173E48')
        for row in matrix.iter_rows(min_row=2,min_col=3):
            for cell in row:cell.number_format='0.00'
        matrix.freeze_panes='D2';matrix.auto_filter.ref=matrix.dimensions
        for column in matrix.columns:matrix.column_dimensions[column[0].column_letter].width=18
        summary.append(['Periodo Claim completo',claim['period_start']+' / '+claim['period_end']])
        summary.append(['Unidad Claim','Horas decimales; 90 minutos = 1.5 horas; captura diaria a 2 decimales.'])
        summary.append(['OT del mes seleccionado',claim['month_ot_minutes']/60])
        summary.append(['Stand By del mes seleccionado',claim['month_guard_minutes']/60])
    if not claim:
        with connection() as c:
            append_monthly_capture(wb,c,result)
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="Operational_Pulse_{month}.xlsx"'})


from backend.time_recovery import install_recovery
append_monthly_capture, append_native_backup = install_recovery(app, globals())


app.mount('/static', StaticFiles(directory=ROOT / 'frontend' / 'static'), name='static')


@app.get('/')
def index():
    return FileResponse(ROOT / 'frontend' / 'index.html')


if __name__ == '__main__':
    import uvicorn
    uvicorn.run(app, host='127.0.0.1', port=5510)
