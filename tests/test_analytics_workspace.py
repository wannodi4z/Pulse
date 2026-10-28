"""Search drives both records and actual approved analytics, including multi-month comparisons."""
import io
import unittest
from openpyxl import load_workbook
import test_api


class AnalyticsWorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.f=test_api.PulseTests();self.f.setUp()
        a,g,self.m,self.mc,self.u,self.tc=self.f.organization()
        self.g=g;self.a=a
        self.g2=self.f.group(a,'G2',self.m['id'])
        self.other,self.other_client,_=self.f.person('USER','SECOND',group=self.g2)
        for month,hours in [('01',2),('02',3),('03',1)]:
            record=self.f.ot(self.tc,f'2026-{month}-05T10:00:00-06:00',f'2026-{month}-05T{10+hours:02}:00:00-06:00',reference='CHG-ALPHA',description='Análisis y respaldo')
            if month!='03':self.f.approve(self.mc,'ot',record)
        beta=self.f.ot(self.other_client,'2026-01-06T10:00:00-06:00','2026-01-06T14:00:00-06:00',reference='INC-BETA');self.f.approve(self.mc,'ot',beta)
        rule=self.f.post('/api/rules',{'name':'Cobertura NOC','effective_from':'2026-01-01','segments':[{'day':0,'start':'08:00','end':'16:00'}]},self.mc)['id']
        for _ in range(2):
            guard=self.f.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-05','end_date':'2026-01-05'},self.tc)['id'];self.f.approve(self.mc,'guard',guard)

    def tearDown(self):self.f.tearDown()

    def test_query_changes_records_and_totals_across_three_months(self):
        result=self.mc.get('/api/analytics',params={'month':'2026-01','months':'2026-02,2026-03','q':'chg-alpha'}).json()
        self.assertEqual(result['record_count'],3);self.assertEqual(result['record_counts'],{'APPROVED':2,'SUBMITTED':1})
        self.assertEqual(result['summary']['HISTORICAL'],300);self.assertEqual(result['summary']['GUARD'],0)
        self.assertEqual([m['minutes']['total_minutes'] for m in result['comparison']],[120,180,0])
        self.assertEqual(sum(r['minutes'] for r in result['rows']),300)
        accent=self.mc.get('/api/analytics',params={'month':'2026-01','q':'analisis','period':'ALL'}).json()
        self.assertEqual(accent['record_count'],3);self.assertEqual(accent['summary']['total_minutes'],300)
        pending=self.mc.get('/api/analytics',params={'month':'2026-01','q':'CHG-ALPHA','period':'ALL','status':'SUBMITTED'}).json()
        self.assertEqual(pending['record_count'],1);self.assertEqual(pending['summary']['total_minutes'],0)
        empty=self.mc.get('/api/analytics',params={'month':'2026-01','q':'DOES-NOT-EXIST','period':'ALL'}).json()
        self.assertEqual(empty['record_count'],0);self.assertEqual(empty['summary']['total_minutes'],0)

    def test_guard_search_keeps_ot_discount_and_deduplicates_guard_coverages(self):
        result=self.mc.get('/api/analytics',params={'month':'2026-01','q':'cobertura noc'}).json()
        self.assertEqual(result['record_count'],2)
        self.assertTrue(all(r['kind']=='guard' for r in result['records']))
        self.assertEqual(result['summary']['ot_minutes'],0);self.assertEqual(result['summary']['GUARD'],360)
        # Eight-hour coverage minus two hours of OT, not two copies of the coverage.
        all_records=self.mc.get('/api/analytics?month=2026-01').json()
        self.assertEqual(all_records['summary']['total_minutes'],720)

    def test_excel_has_matching_records_and_identical_month_comparison(self):
        response=self.mc.get('/api/analytics.xlsx',params={'month':'2026-01','months':'2026-02,2026-03','q':'CHG-ALPHA'})
        self.assertEqual(response.status_code,200)
        wb=load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames,['Resumen','Por grupo','Por técnico','Detalle','Registros','Comparativa'])
        self.assertEqual(len(list(wb['Registros'].values))-1,3)
        self.assertEqual([row[-1] for row in list(wb['Comparativa'].values)[1:]],[2,3,0])
        self.assertEqual(sum(row[7] for row in list(wb['Detalle'].values)[1:]),5)

    def test_comparison_and_search_keep_authorization_and_validate_periods(self):
        params={'month':'2026-01','months':'2026-02,2026-03','q':'CHG-ALPHA'}
        self.assertEqual(self.tc.get('/api/analytics',params=params).status_code,403)
        self.assertEqual(self.f.client.get('/api/analytics.xlsx',params=params).status_code,403)
        lead,lc,_=self.f.person('USER','BACKUP',group=self.g2,lead=True)
        delegation=self.f.post('/api/delegations',{'lead_id':lead['id']},self.mc)['id']
        self.assertEqual(lc.get('/api/analytics',params=params).json()['summary']['total_minutes'],300)
        self.mc.delete('/api/delegations/'+delegation)
        self.assertEqual(lc.get('/api/analytics',params=params).status_code,403)
        for bad in [{'months':'2026-99'},{'status':'INVALID'},{'record_type':'INVALID'},{'period':'INVALID'}]:
            self.assertEqual(self.mc.get('/api/analytics',params={'month':'2026-01',**bad}).status_code,422)
        deduplicated=self.mc.get('/api/analytics',params={'month':'2026-01','months':'2026-01,2026-01','q':'CHG-ALPHA'}).json()
        self.assertEqual(deduplicated['summary']['total_minutes'],120)
