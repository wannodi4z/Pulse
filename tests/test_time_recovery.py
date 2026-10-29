"""Excel recovery previews, ownership, deduplication and shared accounting."""
import base64
import io
import json
import unittest
from openpyxl import Workbook,load_workbook
import test_api

class TimeRecoveryTests(unittest.TestCase):
 def setUp(self):
  self.f=test_api.PulseTests();self.f.setUp();self.a,self.g,self.m,self.mc,self.u,self.uc=self.f.organization()
 def tearDown(self):self.f.tearDown()
 def encode(self,wb):
  out=io.BytesIO();wb.save(out);return {'content':base64.b64encode(out.getvalue()).decode()}
 def legacy(self,employee='T001'):
  w=Workbook();p=w.active;p.title='Tiempo aprobado';p.append(['Fecha','Employee ID','Tipo','Inicio','Fin','Duración HH:MM','Referencia','Descripción'])
  for kind,a,b,d,ref in [('Stand-by','00:00','03:00','03:00',''),('Onsite Overtime','03:00','04:30','01:30','CHG-RECOVERY'),('Stand-by','04:30','08:00','03:30','')]:p.append(['2026-01-03',employee,kind,a,b,d,ref,'Trabajo recuperado'])
  return w
 def preview(self,data,client=None):return self.f.post('/api/time-recovery/preview',data,client or self.uc)
 def confirm(self,p,client=None):return self.f.post('/api/time-recovery/'+p['id']+'/confirm',{'confirm':True},client or self.uc)
 def clear_time(self):
  with self.f.server.connection(True) as c:
   for t in ['time_recovery_items','time_recovery_drafts','ot_time_corrections','timers','ot_group_activities','ot_extensions','record_origins','overtime','guards','windows']:c.execute('DELETE FROM '+t)
 def test_legacy_preview_documentation_reapproval_and_net_hours(self):
  data=self.encode(self.legacy());p=self.preview(data)
  self.assertEqual(p['ready'],2);self.assertEqual(p['duplicates'],0);self.assertTrue(p['warnings'])
  self.assertEqual(sum(x['minutes'] for x in p['rows']),480)
  self.assertEqual(self.uc.get('/api/state').json()['overtime'],[])
  result=self.confirm(p);self.assertEqual((result['created'],result['documentation'],result['submitted']),(2,1,1))
  self.assertEqual(self.confirm(p),result)
  state=self.uc.get('/api/state').json();ot=state['overtime'][0];guard=state['guards'][0]
  self.assertEqual(ot['status'],'DOCUMENTATION');self.assertFalse(guard['rule_changed']);self.assertEqual(len(guard['timeline_days'][0]['segments']),2)
  self.assertEqual(self.mc.post('/api/approvals/ot/'+ot['id'],json={'decision':'APPROVED'}).status_code,422)
  self.assertEqual(self.uc.get('/api/accounting?month=2026-01').json()['ot_minutes'],0)
  updated=self.uc.put('/api/overtime/'+ot['id'],json={'reference':'CHG-RECOVERY','description':'Documentación revisada','justification':'Recuperación de trabajo','impact':'Cobertura de servicio','in_lieu':'Continuidad operativa','expected_revision':0});self.assertEqual(updated.status_code,200,updated.text)
  self.f.approve(self.mc,'ot',ot['id']);self.f.approve(self.mc,'guard',guard['id'])
  totals=self.uc.get('/api/accounting?month=2026-01').json();self.assertEqual((totals['ot_minutes'],totals['guard_minutes']),(90,390))
  repeat=self.preview(data);self.assertEqual((repeat['ready'],repeat['duplicates']),(0,2))
 def test_native_round_trip_and_capture_template(self):
  rule=self.f.post('/api/rules',{'name':'Cobertura','effective_from':'2026-01-01','segments':[{'day':5,'start':'00:00','end':'08:00'}]},self.mc)['id']
  guard=self.f.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-03','end_date':'2026-01-03'},self.uc)['id'];self.f.approve(self.mc,'guard',guard)
  ot=self.f.ot(self.uc,'2026-01-03T03:00:00-06:00','2026-01-03T04:30:00-06:00',reference='INC-RESTORE');self.f.approve(self.mc,'ot',ot)
  response=self.uc.get('/api/time-backup.xlsx?month=2026-01');self.assertEqual(response.status_code,200,response.text if response.status_code!=200 else '')
  w=load_workbook(io.BytesIO(response.content));self.assertEqual(w.sheetnames[0],'Captura');self.assertIn('Respaldo Pulse',w.sheetnames)
  self.assertEqual(w['Captura'].max_column,29);self.assertEqual(w['Captura']['N2'].value,'T001');self.assertIsNone(w['Captura']['W2'].value);self.assertIsNone(w['Captura']['AB2'].value)
  self.assertAlmostEqual(sum(r[8] for r in list(w['Captura'].values)[1:]),8)
  self.assertTrue(all(c.data_type!='f' for row in w['Captura'] for c in row))
  data={'content':base64.b64encode(response.content).decode()}
  self.assertEqual(self.preview(data)['duplicates'],2)
  self.clear_time();p=self.preview(data);result=self.confirm(p);self.assertEqual((result['created'],result['documentation']),(2,0))
  for key,table in [('ot','overtime'),('guard','guards')]:
   for r in self.uc.get('/api/state').json()[table]:self.f.approve(self.mc,key,r['id'])
  totals=self.uc.get('/api/accounting?month=2026-01').json();self.assertEqual((totals['ot_minutes'],totals['guard_minutes']),(90,390))
  claim=self.uc.get('/api/accounting?month=2026-01&view=claim').json();self.assertEqual((claim['month_ot_minutes'],claim['month_guard_minutes']),(90,390))
 def test_ownership_stale_state_invalid_file_and_conflicts(self):
  self.assertEqual(self.uc.post('/api/time-recovery/preview',json=self.encode(self.legacy('OTHER'))).status_code,422)
  self.assertEqual(self.mc.post('/api/time-recovery/preview',json=self.encode(self.legacy())).status_code,403)
  self.assertEqual(self.uc.post('/api/time-recovery/preview',json={'content':'invalid'}).status_code,422)
  p=self.preview(self.encode(self.legacy()))
  other,oc,_=self.f.person('USER','OTHER',group=self.g)
  self.assertEqual(oc.post('/api/time-recovery/'+p['id']+'/confirm',json={'confirm':True}).status_code,404)
  self.assertEqual(self.uc.post('/api/time-recovery/'+p['id']+'/confirm',json={'confirm':False}).status_code,422)
  self.f.ot(self.uc,'2026-01-03T04:00:00-06:00','2026-01-03T05:00:00-06:00',reference='CHG-CONFLICT')
  self.assertEqual(self.uc.post('/api/time-recovery/'+p['id']+'/confirm',json={'confirm':True}).status_code,409)
  fresh=self.preview(self.encode(self.legacy()));self.assertEqual(fresh['conflicts'],1);self.assertEqual(fresh['ready'],1)
 def test_recovered_guard_cutoff_is_preserved_in_calendar_and_backup(self):
  result=self.confirm(self.preview(self.encode(self.legacy())))
  g=self.uc.get('/api/state').json()['guards'][0]
  with self.f.server.connection(True) as c:
   snap=json.loads(c.execute('SELECT rule_snapshot FROM guards WHERE id=?',(g['id'],)).fetchone()[0]);snap['coverage_end_at']='2026-01-03T06:00:00-06:00';c.execute('UPDATE guards SET rule_snapshot=? WHERE id=?',(json.dumps(snap),g['id']))
  self.f.approve(self.mc,'guard',g['id'])
  self.assertEqual(self.uc.get('/api/accounting?month=2026-01').json()['guard_minutes'],270)
  state=self.uc.get('/api/state').json();self.assertTrue(state['guards'][0]['timeline_days'][-1]['end_at'].startswith('2026-01-03T06:00'))
  w=load_workbook(io.BytesIO(self.uc.get('/api/time-backup.xlsx?month=2026-01').content))
  payload=[json.loads(''.join(str(v or '') for v in row[1:])) for row in list(w['Respaldo Pulse'].values)[1:]]
  guard=next(x for x in payload if x['kind']=='GUARD');self.assertTrue(guard['segments'][-1]['end_at'].startswith('2026-01-03T06:00'))
 def test_native_multiple_segments_preserve_waits_and_future_ot_rejected(self):
  w=Workbook();p=w.active;p.title='Respaldo Pulse';p.append(['Formato','Registro JSON'])
  item={'employee_id':'T001','kind':'OT','source_id':'SOURCE','reference':'CHG-PARTS','description':'Actividad','justification':'Motivo','impact':'Impacto','in_lieu':'Razón','segments':[{'start_at':'2026-01-03T00:00:00-06:00','end_at':'2026-01-03T01:00:00-06:00'},{'start_at':'2026-01-03T03:00:00-06:00','end_at':'2026-01-03T04:00:00-06:00'}]}
  p.append(['PULSE-TIME-1',json.dumps(item)]);r=self.confirm(self.preview(self.encode(w)));self.f.approve(self.mc,'ot',r['ids'][0]);self.assertEqual(self.uc.get('/api/accounting?month=2026-01').json()['ot_minutes'],120)
  item['segments']=[{'start_at':'2099-01-01T00:00:00-06:00','end_at':'2099-01-01T01:00:00-06:00'}];p['B2']=json.dumps(item);self.assertEqual(self.uc.post('/api/time-recovery/preview',json=self.encode(w)).status_code,422)
