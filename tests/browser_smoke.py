"""Real Chromium smoke test. Creates an isolated database and local server automatically.
Install requirements-dev.txt and run `python -m playwright install chromium` first.
Optional PULSE_CHROMIUM_EXECUTABLE selects an already installed Chromium binary.
"""
import json
from datetime import datetime, timedelta, timezone
from openpyxl import load_workbook
import os
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]


def run():
    with tempfile.TemporaryDirectory() as tmp:
        with socket.socket() as sock:
            sock.bind(('127.0.0.1', 0)); port = sock.getsockname()[1]
        env = {**os.environ, 'PULSE_DB_PATH': str(Path(tmp)/'browser-test.db')}
        log = open(Path(tmp)/'server.log', 'w+')
        process = subprocess.Popen([sys.executable, '-m', 'uvicorn', 'server:app', '--host', '127.0.0.1', '--port', str(port)], cwd=ROOT, env=env, stdout=log, stderr=log)
        base = f'http://127.0.0.1:{port}'
        try:
            for _ in range(100):
                try:
                    urllib.request.urlopen(base+'/api/health', timeout=1); break
                except Exception:
                    time.sleep(.1)
            else:
                raise RuntimeError('El servidor de pruebas no inició.')
            with sync_playwright() as p:
                launch = {'headless': True}
                if os.environ.get('PULSE_CHROMIUM_EXECUTABLE'):
                    launch.update(executable_path=os.environ['PULSE_CHROMIUM_EXECUTABLE'], args=['--no-sandbox','--disable-dev-shm-usage'])
                browser = p.chromium.launch(**launch)
                context = browser.new_context(viewport={'width':1440,'height':1000}, accept_downloads=True)
                page = context.new_page(); errors=[];page.on('pageerror',lambda error: errors.append(str(error)))
                page.on('dialog',lambda dialog: dialog.accept())
                page.goto(base)
                def fill(name,value):
                    element=page.locator('dialog[open] [name='+name+']');element.evaluate('(el)=>PulseLayout.reveal(el)');element.fill(value)
                def save():
                    # A resize rebuilds form steps on the next frame; interact after that layout finishes.
                    page.evaluate('()=>new Promise(resolve=>requestAnimationFrame(()=>requestAnimationFrame(resolve)))')
                    for _ in range(30):
                        nxt=page.locator('[data-step-next]:visible:not(:disabled)')
                        if not nxt.count():break
                        nxt.click()
                    page.locator('dialog[open] button[type=submit]').click()
                    page.wait_for_function('()=>!busy')
                def action(name,within=None):
                    page.wait_for_function('()=>!busy')
                    selector=f'[data-action="{name}"]'
                    target=within if within is not None else page.locator('dialog[open]') if page.locator('dialog[open] '+selector).count() else page
                    if not target.locator(selector).count():
                        kinds={'manager-new':'account','account-delete':'account','account-edit':'account','group-transfer':'group','group-edit':'group','group-delete':'group','person-transfer':'person','person-offboard':'manager','person-reactivate':'manager'}
                        popup=target.locator('[data-action=entity-menu]'+('[data-id^="'+kinds[name]+':"]' if name in kinds and within is None else '')).first
                        popup.evaluate('(el)=>PulseLayout.reveal(el)');popup.click();target=page.locator('.entity-popup')
                    element=target.locator(selector).first;element.evaluate('(el)=>{el.closest("details")?.setAttribute("open","");PulseLayout.reveal(el)}');element.click()
                def go(target):
                    element=page.locator('[data-page='+target+']').first
                    element.evaluate('(el)=>el.closest("details")?.setAttribute("open", "")');element.click()
                def login(username,password='Password2026!'):
                    page.locator('[data-form=login] [name=identity]').fill(username)
                    page.locator('[data-form=login] [name=password]').fill(password)
                    page.locator('[data-form=login] button[type=submit]').click()
                    page.locator('.sidebar').wait_for()
                def activate(username,code):
                    action('activate');fill('identity',username);fill('code',code);fill('password','Password2026!');fill('confirm','Password2026!');save();expect(page.locator('.access-result-body')).to_be_visible();action('access-login');expect(page.locator('#modal')).not_to_be_visible()
                # Login visibility preserves the value and never submits the form.
                page.locator('[data-form=login]').wait_for()
                page.locator('#login-password').fill('VisiblePassword2026!')
                page.locator('[data-id=login-password]').click()
                expect(page.locator('#login-password')).to_have_attribute('type','text')
                expect(page.locator('#login-password')).to_have_value('VisiblePassword2026!')
                expect(page.locator('[data-id=login-password]')).to_have_attribute('aria-label','Ocultar contraseña')
                expect(page.locator('[data-form=login]')).to_be_visible()
                page.locator('[data-id=login-password]').click()
                expect(page.locator('#login-password')).to_have_attribute('type','password')
                expect(page.locator('[data-id=login-password]')).to_have_attribute('aria-pressed','false')
                expect(page.locator('.login-card h2')).to_have_text('Control de OT y guardias')
                page.locator('#login-password').fill('')
                # Access dialogs: layout, feedback, visibility and a real recovery cycle.
                action('activate')
                expect(page.locator('#modal [name=identity]')).to_be_focused()
                fill('password','Password2026!');fill('confirm','Different2026!')
                expect(page.locator('[data-password-check=length]')).to_have_class('satisfied')
                expect(page.locator('[data-password-check=match]')).not_to_have_class('satisfied')
                page.locator('#modal [data-action=password-visibility]').first.click()
                expect(page.locator('#access-password')).to_have_attribute('type','text')
                expect(page.locator('#modal [data-action=password-visibility]').first).to_have_attribute('aria-pressed','true')
                page.locator('#modal [data-action=password-visibility]').first.click()
                expect(page.locator('#access-password')).to_have_attribute('type','password')
                for width,height in [(1440,1000),(390,844)]:
                    page.set_viewport_size({'width':width,'height':height})
                    assert page.locator('#modal').evaluate('(el)=>el.scrollWidth<=el.clientWidth'),'Access dialog horizontal overflow'
                    assert page.locator('#modal').evaluate('(el)=>el.getBoundingClientRect().height<=innerHeight'),'Access dialog vertical overflow'
                page.set_viewport_size({'width':1440,'height':1000})
                if os.environ.get('PULSE_SCREENSHOT_DIR'):
                    dest=Path(os.environ['PULSE_SCREENSHOT_DIR']);dest.mkdir(parents=True,exist_ok=True);page.screenshot(path=str(dest/'activate-v1.0.1.png'))
                page.keyboard.press('Escape');expect(page.locator('#modal')).not_to_be_visible()
                action('recovery-request');fill('identity','admin2')
                if os.environ.get('PULSE_SCREENSHOT_DIR'):page.screenshot(path=str(Path(os.environ['PULSE_SCREENSHOT_DIR'])/'recovery-v1.0.1.png'))
                save();expect(page.locator('.access-result-body')).to_contain_text('Si tu cuenta está activa')
                expect(page.locator('#modal h2')).to_have_text('Solicitud recibida')
                action('recovery-complete');expect(page.locator('#modal [name=identity]')).to_have_value('admin2')
                action('recovery-request');expect(page.locator('#modal [name=identity]')).to_have_value('admin2')
                action('close')
                # Branding and credits work before and after authentication.
                action('about');expect(page.locator('.about-project')).to_contain_text('Francisco Cervantes Marquez')
                expect(page.locator('.about-project')).to_contain_text('Versión 1.0.1')
                expect(page.locator('.author-list')).to_contain_text('Samuel Alejandro Díaz Chávez')
                expect(page.locator('.author-list article')).to_have_count(2)
                page.set_viewport_size({'width':390,'height':844})
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'),'Credit dialog overflow on mobile'
                assert page.locator('#modal').evaluate('(el)=>el.getBoundingClientRect().right<=innerWidth')
                page.set_viewport_size({'width':1440,'height':1000})
                assert page.locator('.about-project img.kyndryl-logo').evaluate('(img)=>img.complete&&img.naturalWidth>0')
                action('close')
                login('admin1','admin12026')
                expect(page.locator('.footer-version')).to_contain_text('v1.0.1')
                expect(page.locator('.sidebar [data-page=analytics]')).to_have_count(0)
                action('about');expect(page.locator('.author-list')).to_contain_text('Francisco Cervantes Marquez');action('close')
                go('recoveries');action('recovery-issue')
                expect(page.locator('.secret-code')).to_be_visible();old_recovery_code=page.locator('.secret-code').inner_text();action('close')
                action('recovery-view');expect(page.locator('.secret-code')).to_have_text(old_recovery_code);action('close')
                action('recovery-issue');expect(page.locator('.secret-code')).to_be_visible();recovery_code=page.locator('.secret-code').inner_text();assert recovery_code!=old_recovery_code;action('close');action('logout')
                action('recovery-request');fill('identity','admin2');action('recovery-complete')
                fill('code',old_recovery_code);fill('password','Recovered2026!');fill('confirm','Mismatch2026!');save()
                expect(page.locator('#access-error')).to_contain_text('Las contraseñas no coinciden')
                expect(page.locator('#access-error')).to_be_focused()
                fill('confirm','Recovered2026!');save();expect(page.locator('#access-error')).to_contain_text('Código inválido o vencido')
                expect(page.locator('#modal [name=identity]')).to_have_value('admin2')
                fill('code',recovery_code);save();expect(page.locator('#modal h2')).to_have_text('Contraseña actualizada')
                action('access-login');expect(page.locator('[data-form=login] [name=identity]')).to_have_value('admin2')
                login('admin2','Recovered2026!');action('logout');login('admin1','admin12026')
                go('organization')
                expect(page.locator('h1')).to_have_text('Organización')
                action('account-new');fill('code','MX');fill('name','Operación México');save();expect(page.locator('.account-block')).to_have_count(1)
                account_id=page.locator('.account-block').get_attribute('data-account')
                if os.environ.get('PULSE_SCREENSHOT_DIR'):
                    preview=Path(os.environ['PULSE_SCREENSHOT_DIR']);preview.mkdir(parents=True,exist_ok=True)
                    page.wait_for_function('()=>Number(getComputedStyle(document.querySelector("#toast")).opacity)===0')
                    page.screenshot(path=str(preview/'enterprise-empty-account.png'))
                    page.locator('[data-action=entity-menu][data-id="account:'+account_id+'"]:visible').click()
                    page.screenshot(path=str(preview/'enterprise-account-menu.png'))
                    page.keyboard.press('Escape')
                for code in ['G1','G2','G3']:
                    action('group-new');fill('code',code);fill('name','Grupo '+code);save();expect(page.locator('.group-card')).to_have_count(int(code[-1]))
                gids=page.locator('.group-card').evaluate_all('(nodes)=>nodes.map(n=>n.dataset.group)')
                action('manager-new')
                for name,value in [('first','María'),('last','Responsable'),('employee_id','M001'),('email','manager@example.test')]:fill(name,value)
                page.locator('[name=group_ids]').nth(0).check();page.locator('[name=group_ids]').nth(1).check();save();expect(page.locator('.secret-code')).to_be_visible();manager_code=page.locator('.secret-code').inner_text();action('close')
                expect(page.locator('.manager-block')).to_have_count(1)
                expect(page.locator('.manager-block .group-card')).to_have_count(2)
                expect(page.locator('.manager-meta')).to_contain_text('Multigrupo · 2 grupos')
                expect(page.locator('.unassigned .group-card')).to_have_count(1)
                action('manager-new')
                for name,value in [('first','José'),('last','Segundo'),('employee_id','M002'),('email','manager2@example.test')]:fill(name,value)
                expect(page.locator('[name=group_ids]')).to_have_count(1);page.locator('[name=group_ids]').check();save();expect(page.locator('.secret-code')).to_be_visible();manager2_code=page.locator('.secret-code').inner_text();action('close')
                expect(page.locator('.manager-block')).to_have_count(2)
                codes={}
                for emp,lead,gid in [('T001',False,gids[0]),('L001',True,gids[0]),('T002',False,gids[0]),('L002',False,gids[1])]:
                    action('tech-new',page.locator('[data-group="'+gid+'"]'))
                    for name,value in [('first','Persona'),('last',emp),('employee_id',emp),('email',emp+'@example.test')]:fill(name,value)
                    if lead:page.locator('[name=is_lead]').check()
                    if emp in ['T002','L002']:
                        expect(page.locator('#lead-option')).not_to_be_visible();expect(page.locator('#lead-note')).to_contain_text('L001')
                    save();expect(page.locator('.secret-code')).to_be_visible();codes[emp]=page.locator('.secret-code').inner_text();action('close')
                expect(page.locator('[data-group="'+gids[0]+'"] .technician')).to_have_count(3)
                leadrow=page.locator('.technician').filter(has_text='L001');action('person-edit',leadrow)
                expect(page.locator('[name=is_lead]')).to_be_checked();action('close')
                # Navigation stays on the selected page after sync and can return to Organization.
                for target,title in [('dashboard','Panel de administración'),('admins','Administradores'),('recoveries','Recuperación de acceso'),('profile','Mi perfil')]:
                    go(target);expect(page.locator('h1')).to_have_text(title)
                    page.wait_for_timeout(150)
                    expect(page.locator('.account-block')).to_have_count(0)
                    go('organization');expect(page.locator('h1')).to_have_text('Organización')
                page.wait_for_timeout(5200);expect(page.locator('h1')).to_have_text('Organización');assert not page.url.endswith('#organization')
                # Secondary actions remain available through a keyboard-accessible anchored menu.
                more=page.locator('[data-action=entity-menu][data-id="account:'+account_id+'"]')
                more.evaluate('(el)=>PulseLayout.reveal(el)');more.click()
                expect(page.locator('.entity-popup')).to_be_visible()
                expect(page.locator('.entity-popup [data-action=manager-new]')).to_be_focused()
                page.keyboard.press('ArrowDown');expect(page.locator('.entity-popup [data-action=account-edit]')).to_be_focused()
                menu_geometry=page.locator('.entity-popup').evaluate('(el)=>{const r=el.getBoundingClientRect();return r.left>=0&&r.top>=0&&r.right<=innerWidth&&r.bottom<=innerHeight}')
                assert menu_geometry,'Action menu must stay in the viewport'
                page.keyboard.press('Escape');expect(page.locator('.entity-popup')).to_have_count(0);expect(more).to_be_focused()
                more.click();page.locator('h1').click();expect(page.locator('.entity-popup')).to_have_count(0)
                # Summary and hierarchy share a single workspace; branding and navigation are distinct.
                expect(page.locator('.organization-summary')).to_be_visible()
                expect(page.locator('.account-block')).to_be_visible()
                expect(page.locator('.viewport-sections')).not_to_be_visible()
                expect(page.locator('.topbar .corporate-signature .kyndryl-logo')).to_be_visible()
                expect(page.locator('.sidebar .product-signature')).to_have_count(0)
                assert page.locator('.sidebar').evaluate('(el)=>getComputedStyle(el).backgroundColor')=='rgb(255, 255, 255)'
                assert page.locator('.sidebar .nav-item.active .nav-text').evaluate('(el)=>getComputedStyle(el).color')=='rgb(39, 89, 186)'
                expect(page.locator('.nav-settings summary use')).to_have_attribute('href','/static/icons.svg#settings')
                go('dashboard');expect(page.locator('h1')).to_have_text('Panel de administración')
                expect(page.locator('.welcome-banner')).to_have_count(0)
                expect(page.locator('.admin-account-row')).to_have_count(1)
                action('organization-open');expect(page.locator('.account-block')).to_be_visible()
                # Account with groups cannot be removed.
                action('account-delete');save();expect(page.locator('#toast')).to_contain_text('contiene Grupos');expect(page.locator('.account-block')).to_have_count(1)
                # Desktop and mobile layout must not overflow the viewport.
                for width,height in [(1440,1000),(390,844)]:
                    page.set_viewport_size({'width':width,'height':height});page.wait_for_timeout(100)
                    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'),f'Horizontal overflow at {width}'
                page.set_viewport_size({'width':1440,'height':1000})
                screenshot=os.environ.get('PULSE_SCREENSHOT_DIR')
                if screenshot:
                    Path(screenshot).mkdir(parents=True,exist_ok=True);page.wait_for_function('()=>Number(getComputedStyle(document.querySelector("#toast")).opacity)===0');page.screenshot(path=str(Path(screenshot)/'organization.png'),full_page=True)
                action('logout');activate('L001',codes['L001']);activate('L002',codes['L002']);activate('T001',codes['T001']);activate('M001',manager_code);login('M001')
                go('team');expect(page.locator('h1')).to_have_text('Mi equipo')
                expect(page.locator('[data-action=tech-new]')).to_have_count(0);expect(page.locator('[data-action=person-edit]')).to_have_count(0)
                search=page.locator('#team-search');search.evaluate('(el)=>PulseLayout.reveal(el)');search.fill('T001');expect(page.locator('tbody tr:visible')).to_have_count(1);assert search.evaluate('(el)=>el===document.activeElement');search.fill('')
                # Delegate inline from a person's row; replacing/revoking preserves Lead capability.
                for employee in ['L001']:
                    row=page.locator('tbody tr').filter(has_text=employee)
                    action('delegation-inline',row);expect(page.locator('#modal')).not_to_be_visible()
                    expect(page.locator('#delegation-editor')).to_contain_text(employee)
                    expect(page.locator('#delegation-editor input[type=datetime-local]')).to_have_count(0)
                    expect(page.locator('#delegation-editor input[name=group_ids]')).to_have_count(0)
                    if employee=='L002':expect(page.locator('#delegation-editor')).to_contain_text('sustituirá')
                    if screenshot and employee=='L001':page.screenshot(path=str(Path(screenshot)/'inline-delegation-v1.0.1.png'),full_page=True)
                    page.locator('#delegation-editor button[type=submit]').click()
                    expect(page.locator('#delegation-editor')).to_have_count(0)
                    expect(page.locator('.delegation-summary')).to_contain_text(employee)
                delegated=page.request.get(base+'/api/state').json()
                assert len([d for d in delegated['delegations'] if not d['revoked_at']])==1
                assert len([p for p in delegated['people'] if p['is_lead']])==1
                action('delegation-revoke');expect(page.locator('.delegation-summary')).to_contain_text('Sin delegación asignada')
                # Capture the requested 13-interval week, including overnight and full Sunday.
                go('timeRules');action('rule-new');fill('name','Guardia semanal');fill('effective_from','2026-01-01')
                for day in range(6):
                    page.locator('[data-action=rule-day-open][data-id="'+str(day)+'"]').click()
                    section=page.locator('[data-rule-day="'+str(day)+'"]')
                    section.locator('[name=days]').check()
                    section.locator('[data-start]').first.fill('02:00' if day<5 else '00:00')
                    section.locator('[data-end]').first.fill('04:00' if day<5 else '08:00')
                    action('rule-add',section)
                    section.locator('[data-start]').nth(1).fill('20:00' if day<5 else '15:00')
                    section.locator('[data-end]').nth(1).fill('00:00')
                page.locator('[data-action=rule-day-open][data-id="0"]').click()
                monday=page.locator('[data-rule-day="0"]');action('rule-add',monday)
                expect(monday.locator('[data-rule-segment]')).to_have_count(3)
                monday.locator('[data-action=rule-remove]').last.click();expect(monday.locator('[data-rule-segment]')).to_have_count(2)
                page.locator('[data-action=rule-day-open][data-id="6"]').click()
                sunday=page.locator('[data-rule-day="6"]');sunday.locator('[name=days]').check();sunday.locator('[data-full-day]').check()
                expect(sunday.locator('.rule-full-note')).to_be_visible()
                page.set_viewport_size({'width':390,'height':844});assert page.locator('#modal').evaluate('(el)=>el.scrollWidth<=el.clientWidth')
                page.set_viewport_size({'width':1440,'height':1000})
                if screenshot:page.locator('#modal').evaluate('(el)=>el.scrollTop=0');page.screenshot(path=str(Path(screenshot)/'weekly-intervals-v1.0.1.png'))
                save();expect(page.locator('.rules-grid')).to_contain_text('24 horas')
                assert len(page.request.get(base+'/api/state').json()['rules'][0]['segments'])==13
                # Independent technician session provides a real active event for Manager drilldowns.
                tech_context=browser.new_context();tc=tech_context.request
                assert tc.post(base+'/api/login',data={'identity':'T001','password':'Password2026!'}).status==200
                timer=tc.post(base+'/api/timers/start',data={}).json()['id']
                pending=tc.post(base+'/api/overtime',data={'capture_type':'HISTORICAL','scheduled_start':'2025-12-01T02:00:00-06:00','scheduled_end':'2025-12-01T04:00:00-06:00','reference':'INC-CARD','description':'Actividad desde tarjeta','justification':'Servicio','impact':'Cobertura','in_lieu':'No aplica'})
                assert pending.status==200
                page.reload();page.locator('[data-action=dashboard-drill][data-id=running]').wait_for()
                page.locator('#month').fill('2026-01');page.locator('#month').dispatch_event('change')
                page.locator('[data-action=dashboard-drill][data-id=running]').click()
                expect(page.locator('[name=filter-status]')).to_have_value('RUNNING')
                expect(page.locator('[name=filter-period]')).to_have_value('ALL')
                expect(page.locator('tbody tr')).to_have_count(1);expect(page.locator('.timer-card')).to_have_count(1)
                action('timer-open');expect(page.locator('#modal')).to_contain_text('T001');expect(page.locator('#modal .badge')).to_contain_text('En curso');action('close')
                go('dashboard')
                expect(page.locator('.priority-list')).to_contain_text('INC-CARD')
                page.locator('[data-action=dashboard-drill][data-id=pending]').click()
                expect(page.locator('[name=filter-period]')).to_have_value('ALL');expect(page.locator('tbody tr')).to_have_count(1)
                action('record-detail');action('approve');expect(page.locator('#modal')).not_to_be_visible();expect(page.locator('tbody tr')).to_have_count(0)
                assert tc.post(base+'/api/timers/'+timer+'/stop',data={}).status==200
                tech_context.close()
                action('logout');login('T001')
                go('time');expect(page.locator('h1')).to_have_text('Mi tiempo')
                action('ot-new');page.locator('[name=capture_type]').select_option('HISTORICAL')
                for name,value in [('scheduled_start','2026-01-05T23:00'),('scheduled_end','2026-01-06T01:00'),('reference','CHG12345'),('description','Cambio de prueba'),('justification','Mantenimiento'),('impact','Disponibilidad'),('in_lieu','Ventana de servicio')]:fill(name,value)
                save();expect(page.locator('#modal')).not_to_be_visible()
                page.locator('#month').fill('2026-01');page.locator('#month').dispatch_event('change');expect(page.locator('.record-list').first).to_contain_text('CHG12345')
                action('guard-new');fill('start_date','2026-01-05');fill('end_date','2026-01-05');save();expect(page.locator('#modal')).not_to_be_visible()
                action('timer-start');save();expect(page.locator('.timer-card')).to_have_count(1);page.wait_for_timeout(1100);action('timer-stop');expect(page.locator('.timer-card')).to_have_count(0)
                action('logout');login('M001');go('teamTime')
                page.locator('#month').fill('2026-01');page.locator('#month').dispatch_event('change');expect(page.locator('tbody tr')).to_have_count(2)
                action('approve');expect(page.locator('[data-action=approve]')).to_have_count(1);action('approve');expect(page.locator('[data-action=approve]')).to_have_count(0)
                with page.expect_download() as download_info:page.locator('a[download]').click()
                download=download_info.value;path=Path(tmp)/'export.xlsx';download.save_as(path);assert path.stat().st_size>1000
                # Team analytics reconcile with approved Jan OT (2h) and net guard (5h).
                go('analytics');expect(page.locator('h1')).to_have_text('Analítica del equipo')
                report=page.request.get(base+'/api/analytics?month=2026-01').json()
                assert report['summary']['HISTORICAL']==120 and report['summary']['GUARD']==300
                expect(page.locator('.analytics-metrics button')).to_have_count(5)
                expect(page.locator('#analytics-users tbody tr')).to_have_count(4)
                page.locator('[data-action=analytics-records][data-id=HISTORICAL]').evaluate('(el)=>PulseLayout.reveal(el)');page.locator('[data-action=analytics-records][data-id=HISTORICAL]').click();expect(page.locator('#analytics-records tbody tr')).to_have_count(2)
                page.locator('[data-action=analytics-records][data-id=EMERGENCY]').evaluate('(el)=>PulseLayout.reveal(el)');page.locator('[data-action=analytics-records][data-id=EMERGENCY]').click();expect(page.locator('#analytics-records')).to_contain_text('Sin horas aprobadas de este tipo')
                page.locator('[data-action=analytics-group][data-id="'+gids[0]+'"]').first.evaluate('(el)=>PulseLayout.reveal(el)');page.locator('[data-action=analytics-group][data-id="'+gids[0]+'"]').first.click()
                expect(page.locator('[name=filter-group]')).to_have_value(gids[0]);expect(page.locator('#analytics-users tbody tr')).to_have_count(3)
                tech_row=page.locator('#analytics-users tbody tr').filter(has_text='T001');action('analytics-user',tech_row)
                expect(page.locator('#analytics-records tbody tr')).to_have_count(4)
                for width,height in [(1440,1000),(390,844)]:
                    page.set_viewport_size({'width':width,'height':height});assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
                page.set_viewport_size({'width':1440,'height':1000})
                if screenshot:page.evaluate('window.scrollTo(0,0)');page.screenshot(path=str(Path(screenshot)/'analytics-team-v1.0.1.png'),full_page=True)
                with page.expect_download() as analytics_download:page.locator('a[download]').click()
                analytics_file=Path(tmp)/'analytics.xlsx';analytics_download.value.save_as(analytics_file)
                workbook=load_workbook(analytics_file);assert workbook.sheetnames==['Resumen','Por grupo','Por técnico','Detalle']
                assert list(workbook['Por técnico'].values)[1][-1]==7
                other_person=next(p['id'] for p in report['users'] if p['employee_id']=='T002')
                lead_person=next(p['id'] for p in report['users'] if p['employee_id']=='L001')
                action('logout');login('T001')
                expect(page.locator('.sidebar [data-page=analytics]')).to_have_count(0)
                go('myAnalytics');expect(page.locator('h1')).to_have_text('Mi resumen mensual')
                expect(page.locator('.analytics-filters')).to_have_count(0);expect(page.locator('#analytics-users')).to_have_count(0)
                personal=page.request.get(base+'/api/my-analytics?month=2026-01').json();assert len(personal['users'])==1 and personal['users'][0]['employee_id']=='T001'
                assert personal['summary']['total_minutes']==420
                assert page.request.get(base+'/api/analytics?month=2026-01').status==403
                assert page.request.get(base+'/api/my-analytics?month=2026-01&user_id='+other_person).status==422
                for target in ['claim','adp']:
                    go(target);expect(page.locator('h1')).to_have_text(target.upper() if target=='adp' else 'Claim');expect(page.locator('tbody tr')).to_have_count(4)
                action('logout');login('L001')
                expect(page.locator('.sidebar [data-page=time]')).to_have_count(1);expect(page.locator('.sidebar [data-page=teamTime]')).to_have_count(0)
                expect(page.locator('.sidebar [data-page=analytics]')).to_have_count(0)
                manager_context=browser.new_context();manager_api=manager_context.request
                assert manager_api.post(base+'/api/login',data={'identity':'M001','password':'Password2026!'}).status==200
                current=datetime.now(timezone.utc)
                delegation=manager_api.post(base+'/api/delegations',data={'lead_id':lead_person,'group_ids':[gids[1]],'start_at':(current-timedelta(minutes=1)).isoformat(),'end_at':(current+timedelta(hours=1)).isoformat()})
                assert delegation.status==200
                page.reload();page.locator('.sidebar [data-page=analytics]').wait_for()
                expect(page.locator('.sidebar [data-page=recoveries]')).to_have_count(0)
                go('analytics')
                expect(page.locator('#analytics-users tbody tr')).to_have_count(4);expect(page.locator('#analytics-users')).to_contain_text('L002')
                expect(page.locator('[name=filter-group] option')).to_have_count(3)
                assert page.request.get(base+'/api/analytics?month=2026-01&group_id='+gids[0]).status==200
                assert manager_api.delete(base+'/api/delegations/'+delegation.json()['id']).status==200
                page.locator('#month').fill('2026-02');page.locator('#month').dispatch_event('change')
                expect(page.locator('h1')).to_have_text('Mi resumen mensual');expect(page.locator('.sidebar [data-page=analytics]')).to_have_count(0)
                manager_context.close()
                go('time');expect(page.locator('h1')).to_have_text('Mi tiempo')
                # Editing and deleting the template cannot rewrite existing guard requests.
                action('logout');login('M001');go('timeRules')
                existing=page.request.get(base+'/api/state').json();rule_id=existing['rules'][0]['id']
                original_snapshot=existing['guards'][0]['rule_snapshot']
                action('rule-edit');fill('name','Guardia semanal actualizada');save()
                expect(page.locator('.rules-grid')).to_contain_text('Guardia semanal actualizada')
                updated=page.request.get(base+'/api/state').json()
                assert updated['guards'][0]['rule_snapshot']==original_snapshot
                action('rule-delete');expect(page.locator('#modal')).to_contain_text('Las Guardias existentes conservan');save()
                assert page.request.get(base+'/api/state').json()['rules']==[]
                # Admin can transfer active people/groups independently from access withdrawal.
                action('logout');activate('M002',manager2_code);login('admin1','admin12026')
                go('organization')
                source=page.request.get(base+'/api/state').json()
                tech=next(x for x in source['people'] if x['employee_id']=='T001')
                manager1=next(x for x in source['people'] if x['employee_id']=='M001')
                manager2=next(x for x in source['people'] if x['employee_id']=='M002')
                row=page.locator('[data-person="'+tech['id']+'"]')
                action('person-transfer',row);expect(page.locator('.form-instructions')).to_contain_text('conserva su estado')
                page.locator('[name=group_id]').select_option(gids[1]);fill('transfer_reason','Reorganización de equipo activo');save()
                expect(page.locator('[data-group="'+gids[1]+'"] [data-person="'+tech['id']+'"]')).to_have_count(1)
                assert next(x for x in page.request.get(base+'/api/state').json()['people'] if x['id']==tech['id'])['status']=='ACTIVE'
                action('group-transfer',page.locator('[data-group="'+gids[1]+'"]'))
                page.locator('[name=manager_id]').select_option(manager2['id']);fill('transfer_reason','Traspaso de equipo activo');save()
                expect(page.locator('[data-manager="'+manager2['id']+'"] [data-group="'+gids[1]+'"]')).to_have_count(1)
                action('person-offboard',page.locator('[data-manager="'+manager1['id']+'"] > .manager-header'))
                page.locator('[name="transfer:'+gids[0]+'"]').select_option(manager2['id'])
                fill('reason','Sustitución de Manager');page.locator('[name=confirm]').check()
                for width,height in [(1440,760),(1366,700),(1280,650)]:
                    page.set_viewport_size({'width':width,'height':height})
                    geometry=page.locator('#modal').evaluate('(el)=>({width:el.scrollWidth,client:el.clientWidth,height:el.getBoundingClientRect().height,bodyScroll:el.querySelector(".form-body").scrollHeight,bodyClient:el.querySelector(".form-body").clientHeight})')
                    print('100% offboard layout',width,height,geometry,flush=True)
                    assert geometry['width']<=geometry['client']
                    assert geometry['height']<=height
                    assert geometry['bodyScroll']<=geometry['bodyClient']+1,'Baja requires scrolling at 100%'
                page.set_viewport_size({'width':1440,'height':760})
                if screenshot:page.screenshot(path=str(Path(screenshot)/'manager-offboard-v1.0.1.png'))
                save();expect(page.locator('#modal')).not_to_be_visible()
                final=page.request.get(base+'/api/state').json()
                assert next(x for x in final['people'] if x['id']==manager1['id'])['status']=='DISABLED'
                assert next(x for x in final['people'] if x['id']==tech['id'])['status']=='ACTIVE'
                expect(page.locator('[data-manager="'+manager2['id']+'"] .group-card')).to_have_count(3)
                action('person-reactivate',page.locator('[data-manager="'+manager1['id']+'"] > .manager-header'))
                fill('reason','Nuevo acceso con código de activación');save();expect(page.locator('.secret-code')).to_be_visible();action('close')
                assert next(x for x in page.request.get(base+'/api/state').json()['people'] if x['id']==manager1['id'])['status']=='PENDING_ACTIVATION'
                assert errors==[],errors
                print(json.dumps({'result':'PASS','browser':'Chromium','checks':['Kyndryl brand and credits before/after login','version 1.0.1','login password visibility and OT/guard copy','access dialogs: mobile/focus/visibility/validation/success','real password recovery through UI','initial login','account/groups/two managers','physical nesting','multigroup label','technicians and Lead selector','activation Manager/Technician/Lead','sidebar exit and polling','responsive desktop/mobile','search focus','13 weekly intervals and full Sunday','inline unique Manager Lead delegation/revocation','dashboard running/pending drilldowns and detail approvals','rule and guard request','historical OT','emergency timer start/stop','manager approvals','download Excel','team analytics group/user/type/month drilldowns and Excel','personal analytics cannot request other users','delegated analytics scoped and revoked live','Admin-only people management','view/regenerate recovery codes','Claim/ADP rows','Lead has no implicit approval authority','one Lead across all Manager groups','guided expanded confirmation dialogs','active Technician and Group transfers','Manager offboard and reactivation via UI','edit/delete guard rules with preserved request snapshots','simplified sidebar and user menu','native viewport sections and list pagination','organization summary combined with hierarchy','actionable Admin dashboard opens selected Account','corporate header Kyndryl and distinct configuration gear','release version 1.0.1','Enterprise light shell, hierarchy and secondary action menus','action menu keyboard focus/arrows/Escape and outside dismissal'],'javascript_errors':errors},ensure_ascii=False,indent=2))
                browser.close()
        except Exception:
            log.flush();log.seek(0);print(log.read()[-7000:]);raise
        finally:
            process.terminate()
            try:process.wait(timeout=8)
            except subprocess.TimeoutExpired:process.kill();process.wait()
            log.close()


if __name__=='__main__':run()
