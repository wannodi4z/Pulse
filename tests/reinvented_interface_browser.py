"""Real navigation and operational regression for the alternative presentation shells.
Run this file inside a packaged Orbit/Prisma project; each test uses an isolated DB.
"""
import io,os,shutil,socket,subprocess,sys,tempfile,time,urllib.request
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
R=Path(__file__).resolve().parents[1]
with tempfile.TemporaryDirectory() as d:
 db=Path(d)/'ui.db';shutil.copy(R/'fixtures/banorte_pruebas.sqlite3',db)
 with socket.socket() as sk:sk.bind(('127.0.0.1',0));port=sk.getsockname()[1]
 base=f'http://127.0.0.1:{port}';proc=subprocess.Popen([sys.executable,'-m','uvicorn','server:app','--port',str(port)],cwd=R,env={**os.environ,'PULSE_DB_PATH':str(db)},stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
 try:
  for _ in range(100):
   try:urllib.request.urlopen(base+'/api/health');break
   except Exception:time.sleep(.1)
  with sync_playwright() as p:
   browser=p.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);ctx=browser.new_context(viewport={'width':1440,'height':1000},accept_downloads=True,reduced_motion='reduce');page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
   page.goto(base);page.locator('[data-form=login]').wait_for();kind=page.locator('html').get_attribute('data-experience');assert kind in ('orbit','prisma')
   for w,h in [(1440,1000),(1024,600),(390,844),(320,568)]:
    page.set_viewport_size({'width':w,'height':h});assert page.evaluate('document.documentElement.scrollWidth<=innerWidth+1'),('login',w)
    for selector in ['.login-submit','[data-action=activate]','[data-action=recovery-request]']:
     page.locator(selector).scroll_into_view_if_needed();expect(page.locator(selector)).to_be_visible()
   page.set_viewport_size({'width':1440,'height':1000})
   def login(emp,pw):
    ctx.clear_cookies();assert ctx.request.post(base+'/api/login',data={'identity':emp,'password':pw}).ok;page.goto(base);page.locator('#main').wait_for()
   def nav(view):
    if kind=='prisma':
     page.locator('.prisma-menu-toggle').click();expect(page.locator('#experience-menu')).to_be_visible();page.locator('#experience-menu [data-page='+view+']').click();expect(page.locator('#experience-menu')).not_to_be_visible()
    else:page.locator('.orbit-navigation [data-page='+view+']').click()
    page.wait_for_function('()=>!busy');expect(page.locator('#main')).to_have_attribute('data-experience-page',view)
   def close():page.locator('#modal [data-action=close]').first.click()
   def fill(name,value):page.locator('#modal [name='+name+']').fill(value)
   def submit():page.locator('#modal [type=submit]').click();page.wait_for_function('()=>!busy')
   checks=0
   roles=[('admin1','admin12026',['dashboard','organization','admins','recoveries']),('MGR0001','manager12026',['dashboard','team','teamTime','analytics','timeRules','recoveries']),('SL0001','lead120261',['dashboard','time','myAnalytics','claim','adp']),('951357','tecnico12026',['dashboard','time','myAnalytics','claim','adp'])]
   for emp,pw,pages in roles:
    login(emp,pw)
    for w,h in [(1440,1000),(1024,600),(800,540),(390,844)]:
     page.set_viewport_size({'width':w,'height':h})
     for view in pages:
      nav(view);assert page.evaluate('document.documentElement.scrollWidth<=innerWidth+1'),(kind,emp,view,w);expect(page.locator('#main h1')).to_be_visible();checks+=1
    if kind=='prisma':
     page.locator('.prisma-menu-toggle').click();expect(page.locator('#experience-menu')).to_be_visible();assert page.evaluate('document.querySelector("#experience-menu").contains(document.activeElement)');page.keyboard.press('Escape');expect(page.locator('#experience-menu')).not_to_be_visible()
   # Create and edit a real Technician using the existing production actions.
   page.set_viewport_size({'width':1440,'height':1000});login('MGR0001','manager12026');nav('team')
   rr=ctx.request.post(base+'/api/rules',data={'name':'Regla UI','effective_from':'2026-01-01','segments':[{'day':i,'start':'20:00','end':'00:00'} for i in range(7)]});assert rr.ok;rule=rr.json()['id']
   before=ctx.request.get(base+'/api/state').json();own=before['groups'];source=own[0]['id'];target=own[1]['id']
   page.locator('[data-team-destination="'+source+'"] [data-action=tech-new]').click()
   for k,v in [('first','Prueba'),('last','Interfaz'),('employee_id','UI-NEW'),('email','ui-new@example.test')]:fill(k,v)
   submit();expect(page.locator('.secret-code')).to_be_visible();code=page.locator('.secret-code').inner_text();close()
   public=browser.new_context();assert public.request.post(base+'/api/activate',data={'identity':'UI-NEW','code':code,'password':'TechUI2026!'}).ok
   state=ctx.request.get(base+'/api/state').json();u=next(x for x in state['people'] if x['employee_id']=='UI-NEW');card=page.locator('[data-team-person="'+u['id']+'"]')
   card.locator('[data-action=team-details]').last.click();page.locator('#modal [data-action=person-edit]').click();fill('first','Editado');submit();expect(card).to_contain_text('Editado')
   # Native drag must still open the real move confirmation, not write immediately.
   card.drag_to(page.locator('[data-team-destination="'+target+'"]'));expect(page.locator('#dialog-title')).to_have_text('Mover Técnico');expect(page.locator('#modal [name=group_id]')).to_have_value(target);fill('reason','Prueba de redistribución visual');submit();assert next(x for x in ctx.request.get(base+'/api/state').json()['people'] if x['id']==u['id'])['group_id']==target
   # Create historical OT in the UI, and approve it through the Manager UI.
   login('UI-NEW','TechUI2026!');nav('time');page.locator('[data-action=ot-new]').first.click()
   for k,v in [('scheduled_start','2026-10-05T21:00'),('scheduled_end','2026-10-05T22:30'),('reference','CHG106501'),('description','Validación de interfaz'),('justification','Actividad requerida'),('impact','Continuidad'),('in_lieu','Cobertura fuera de horario')]:fill(k,v)
   for ww,hh in [(1440,1000),(1024,600),(390,844)]:
    page.set_viewport_size({'width':ww,'height':hh});page.wait_for_timeout(100);assert page.locator('#modal').evaluate('e=>{let r=e.getBoundingClientRect();return r.left>=0&&r.right<=innerWidth+1&&r.top>=0&&r.bottom<=innerHeight+1}');assert page.locator('#modal .dialog-footer').evaluate('e=>e.getBoundingClientRect().bottom<=innerHeight+1')
   page.set_viewport_size({'width':1440,'height':1000})
   submit();expect(page.locator('#modal')).not_to_be_visible();state=ctx.request.get(base+'/api/state').json();ot=next(x for x in state['overtime'] if x['reference']=='CHG106501');assert ot['status']=='SUBMITTED'
   page.locator('#month').fill('2026-10');page.locator('#month').dispatch_event('change');page.wait_for_function('()=>!busy')
   start=page.locator('[data-calendar-date="2026-10-12"] .calendar-date-button');end=page.locator('[data-calendar-date="2026-10-18"] .calendar-date-button');start.evaluate('e=>e.scrollIntoView({block:"center",behavior:"instant"})');aa=start.bounding_box();bb=end.bounding_box();page.mouse.move(aa['x']+10,aa['y']+10);page.mouse.down();page.mouse.move(bb['x']+10,bb['y']+10,steps=15);page.mouse.up();expect(page.locator('#dialog-title')).to_have_text('Solicitar Guardia');page.locator('#modal [name=rule_id]').select_option(rule);submit();assert len(ctx.request.get(base+'/api/state').json()['guards'])==1
   login('MGR0001','manager12026');nav('teamTime');page.locator('[data-action=record-detail][data-id="ot:'+ot['id']+'"]').first.click();page.locator('#modal [data-action=approve]').click();page.wait_for_function('()=>!busy');assert next(x for x in ctx.request.get(base+'/api/state').json()['overtime'] if x['id']==ot['id'])['status']=='APPROVED'
   nav('analytics');page.locator('#month').fill('2026-10');page.locator('#month').dispatch_event('change');page.wait_for_function('()=>!busy');page.locator('[data-action=chart-expand][data-id=hourly]').click();expect(page.locator('#modal.expanded-chart-dialog')).to_be_visible();close()
   login('UI-NEW','TechUI2026!');nav('claim');page.locator('#month').fill('2026-10');page.locator('#month').dispatch_event('change');page.wait_for_function('()=>!busy');expect(page.locator('#main')).to_contain_text('1.5')
   nav('adp');expect(page.locator('#main')).to_contain_text('01:30')
   nav('time');page.locator('[data-action=timer-start]').first.click();page.locator('#modal [type=submit]').click();page.wait_for_function('()=>!busy');expect(page.locator('#modal [data-action=timer-stop]')).to_be_visible();page.locator('#modal [data-action=timer-stop]').click();page.wait_for_function('()=>!busy');expect(page.locator('#dialog-title')).to_have_text('Documentar OT');close();expect(page.locator('[data-documentation-id]').first).to_be_visible()
   page.locator('[data-action=time-recovery]').click()
   with page.expect_download() as dl:page.locator('.recovery-download').click()
   assert dl.value.suggested_filename.endswith('.pulse');file=Path(d)/'backup.pulse';dl.value.save_as(file);w=load_workbook(io.BytesIO(file.read_bytes()));assert 'Respaldo Pulse' in w.sheetnames
   page.locator('[name=backup_file]').set_input_files(file);submit();expect(page.locator('#modal')).to_contain_text('duplicado');expect(page.locator('#modal [type=submit]')).to_be_disabled();close()
   assert not errors,errors
   public.close();ctx.close();browser.close();print(f'PASS {kind}: {checks} real navigation/page/viewport checks; responsive access; launcher keyboard focus/Escape; Technician create/activation/edit/native drag; historical OT and approval; drag-selected Guard; emergency Timer/documentation; responsive modal footers; expanded chart; Claim/ADP; .pulse and duplicate recovery; no JavaScript errors')
 finally:proc.terminate();proc.wait(timeout=10)
