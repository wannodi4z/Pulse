"""Hourly accounting, scoped personal comparisons, and weekly overlap protection."""
import io
import unittest
from openpyxl import load_workbook
import test_analytics_workspace

class Dashboard107Tests(unittest.TestCase):
    def setUp(self):
        self.fixture=test_analytics_workspace.AnalyticsWorkspaceTests();self.fixture.setUp()
        self.f=self.fixture.f;self.mc=self.fixture.mc;self.tc=self.fixture.tc
    def tearDown(self):self.fixture.tearDown()
    def test_hourly_totals_match_net_accounting_and_scope(self):
        r=self.mc.get('/api/analytics?month=2026-01').json()
        self.assertEqual(len(r['hourly']),24)
        self.assertEqual(sum(x['minutes']['total_minutes'] for x in r['hourly']),r['summary']['total_minutes'])
        self.assertEqual(r['hourly'][10]['minutes']['HISTORICAL'],120)
        self.assertEqual(r['hourly'][10]['minutes']['GUARD'],0)
        self.assertEqual(r['hourly'][8]['minutes']['GUARD'],60)
        scoped=self.mc.get('/api/analytics',params={'month':'2026-01','group_id':self.fixture.g}).json()
        self.assertEqual(scoped['summary']['total_minutes'],480)
        self.assertEqual(len(scoped['users']),1)
        queried=self.mc.get('/api/analytics?month=2026-01&q=INC-BETA').json()
        self.assertEqual(sum(x['minutes']['total_minutes'] for x in queried['hourly']),240)
    def test_personal_comparison_search_and_export_never_include_peers(self):
        r=self.tc.get('/api/my-analytics?month=2026-01&months=2026-02,2026-03&q=CHG-ALPHA').json()
        self.assertEqual(r['mode'],'PERSONAL');self.assertEqual(r['summary']['total_minutes'],300)
        self.assertEqual([m['minutes']['total_minutes'] for m in r['comparison']],[120,180,0])
        other=self.tc.get('/api/my-analytics?month=2026-01&q=INC-BETA&period=ALL').json()
        self.assertEqual(other['records'],[])
        self.assertEqual(self.tc.get('/api/my-analytics?month=2026-01&user_id=other').status_code,422)
        export=self.tc.get('/api/my-analytics.xlsx?month=2026-01&months=2026-02&q=CHG-ALPHA')
        wb=load_workbook(io.BytesIO(export.content));self.assertIn('Por hora',wb.sheetnames);self.assertIn('Comparativa',wb.sheetnames)
    def test_weekly_wrap_overlap_rejected_but_adjacent_ranges_accepted(self):
        payload={'name':'Cruce semanal','effective_from':'2026-01-01','segments':[{'day':6,'start':'22:00','end':'02:00'},{'day':0,'start':'01:00','end':'03:00'}]}
        self.assertEqual(self.mc.post('/api/rules',json=payload).status_code,422)
        payload['segments'][1]['start']='02:00'
        self.assertEqual(self.mc.post('/api/rules',json=payload).status_code,200)
        payload['segments']=[{'day':0,'start':'00:00','end':'00:00'},{'day':0,'start':'02:00','end':'03:00'}]
        self.assertEqual(self.mc.post('/api/rules',json=payload).status_code,422)
    def test_hourly_cross_midnight_fractional_time(self):
        ot=self.f.ot(self.tc,'2026-01-07T23:30:00-06:00','2026-01-08T00:15:00-06:00',reference='INC-MIDNIGHT');self.f.approve(self.mc,'ot',ot)
        r=self.mc.get('/api/analytics?month=2026-01&q=INC-MIDNIGHT').json()
        self.assertEqual(r['hourly'][23]['minutes']['total_minutes'],30)
        self.assertEqual(r['hourly'][0]['minutes']['total_minutes'],15)
