"""Functional tests run against a disposable SQLite database, never the user's local data."""
import importlib.util
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


class PulseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        with patch.dict(os.environ, {'PULSE_DB_PATH': str(Path(self.temp.name) / 'test.db')}):
            spec = importlib.util.spec_from_file_location('pulse_test_server', ROOT / 'server.py')
            self.server = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(self.server)
        self.client = TestClient(self.server.app)
        self.login(self.client, 'admin1', 'admin12026')

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def login(self, client, username, password='Password2026!'):
        result = client.post('/api/login', json={'identity': username, 'password': password})
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()['user']

    def post(self, path, data, client=None):
        result = (client or self.client).post(path, json=data)
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    def account(self, code='AC'):
        return self.post('/api/accounts', {'code': code, 'name': 'Cuenta ' + code})['id']

    def group(self, account, code='G1', manager=None):
        return self.post('/api/groups', {'code': code, 'name': 'Grupo ' + code, 'account_id': account, 'manager_id': manager})['id']

    def person(self, role, emp, account=None, group=None, groups=None, lead=False):
        data = {'role': role, 'employee_id': emp, 'username': emp, 'first': 'Persona', 'last': emp, 'email': emp+'@example.test', 'account_id': account, 'group_id': group, 'group_ids': groups or [], 'is_lead': lead}
        creator=None
        if role=='USER' and not lead and group:
            with self.server.connection() as db:
                owner=db.execute('SELECT u.username FROM groups g JOIN users u ON u.id=g.manager_id WHERE g.id=?',(group,)).fetchone()
            if owner:
                creator=TestClient(self.server.app);self.login(creator,owner['username'])
        if role=='USER' and not lead and group and not creator:
            identity=self.server.uid()
            with self.server.connection(True) as db:
                db.execute('INSERT INTO users(id,role,username,employee_id,first,last,email,group_id,status,password_hash,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',(identity,'USER',emp,emp,'Persona',emp,emp+'@example.test',group,'ACTIVE',self.server.password_hash('Password2026!'),self.server.now()))
            legacy=TestClient(self.server.app);return self.login(legacy,emp),legacy,data
        r = self.post('/api/people', data, creator)
        if creator:creator.post('/api/logout');creator.close()
        c = TestClient(self.server.app)
        self.post('/api/activate', {'identity': emp, 'code': r['activation_code'], 'password': 'Password2026!'}, c)
        user = self.login(c, emp)
        return user, c, data

    def organization(self):
        a = self.account()
        g = self.group(a)
        m, mc, md = self.person('MANAGER', 'M001', account=a, groups=[g])
        u, uc, ud = self.person('USER', 'T001', group=g)
        return a, g, m, mc, u, uc

    def ot(self, client, start='2026-01-05T23:00:00-06:00', end='2026-01-06T01:00:00-06:00', **extra):
        return self.post('/api/overtime', {'capture_type': 'HISTORICAL', 'scheduled_start': start, 'scheduled_end': end, 'reference': 'CHG123456', 'description': 'Trabajo aprobado', 'justification': 'Cambio requerido', 'impact': 'Servicio', 'in_lieu': 'Cobertura', **extra}, client)['id']

    def approve(self, manager, kind, identity):
        return self.post(f'/api/approvals/{kind}/{identity}', {'decision': 'APPROVED'}, manager)

    def test_clean_bootstrap_and_reinitialization(self):
        with self.server.connection() as c:
            users = self.server.rows(c, 'users')
            self.assertEqual({u['username'] for u in users}, {'admin1', 'admin2'})
            for table in ['accounts', 'groups', 'overtime', 'guards', 'timers', 'rules', 'windows', 'delegations', 'recoveries']:
                self.assertEqual(c.execute(f'SELECT count(*) FROM {table}').fetchone()[0], 0, table)
        a = self.account()
        self.server.init_db()
        self.assertEqual(self.client.get('/api/state').json()['accounts'][0]['id'], a)
        text = self.client.get('/api/state').text
        self.assertNotIn('password_hash', text)
        self.assertNotIn('activation_hash', text)
        self.assertEqual(TestClient(self.server.app).get('/api/state').status_code, 401)

    def test_hierarchy_ownership_cross_account_and_delete(self):
        a, b = self.account(), self.account('OTHER')
        g1, g2, g3 = self.group(a), self.group(a, 'G2'), self.group(b, 'G3')
        m, mc, data = self.person('MANAGER', 'M001', account=a, groups=[g1, g2])
        m2, _, d2 = self.person('MANAGER', 'M002', account=a)
        r = self.client.put('/api/people/'+m2['id'], json={**d2,'group_ids':[g1]})
        self.assertEqual(r.status_code, 422)
        r = self.client.put('/api/people/'+m['id'], json={**data,'group_ids':[g1,g3]})
        self.assertEqual(r.status_code, 422)
        self.assertEqual(self.client.delete('/api/accounts/'+a).status_code, 422)
        empty = self.account('EMPTY')
        self.assertEqual(self.client.delete('/api/accounts/'+empty).status_code, 200)
        state = self.client.get('/api/state').json()
        self.assertEqual(len([p for p in state['people'] if p['role']=='MANAGER' and p['account_id']==a]),2)
        self.assertEqual(len([g for g in state['groups'] if g['manager_id']==m['id']]),2)
        self.assertEqual(mc.post('/api/accounts',json={'name':'Bad','code':'BAD'}).status_code,403)

    def test_delegated_backup_covers_all_current_manager_groups(self):
        a,g1,m,mc,_,_=self.organization()
        g2=self.group(a,'SECOND',m['id'])
        lead,lc,_=self.person('USER','BACKUP',group=g1,lead=True)
        _,tc,_=self.person('USER','SECONDTECH',group=g2)
        d=self.post('/api/delegations',{'lead_id':lead['id'],'group_ids':[g1]},mc)['id']
        self.assertEqual(set(lc.get('/api/state').json()['scope']),{g1,g2})
        self.approve(lc,'ot',self.ot(tc))
        rejected=self.ot(tc,'2026-01-07T23:00:00-06:00','2026-01-08T01:00:00-06:00')
        self.assertEqual(lc.post('/api/approvals/ot/'+rejected,json={'decision':'REJECTED','reason':'Corregir documentación'}).status_code,200)
        g3=self.group(a,'NEWGROUP',m['id'])
        self.assertEqual(set(lc.get('/api/state').json()['scope']),{g1,g2,g3})
        foreign=self.account('OTHER');fg=self.group(foreign,'FG')
        _,fc,_=self.person('MANAGER','OTHER_MANAGER',account=foreign,groups=[fg])
        _,ft,_=self.person('USER','FOREIGN_TECH',group=fg)
        foreign_ot=self.ot(ft)
        self.assertEqual(lc.post('/api/approvals/ot/'+foreign_ot,json={'decision':'APPROVED'}).status_code,403)
        self.assertEqual(mc.delete('/api/delegations/'+d).status_code,200)
        self.assertFalse(lc.get('/api/state').json()['can_approve'])

    def test_delegation_has_no_expiry_and_requires_manual_revocation(self):
        _,gid,_,mc,_,_=self.organization()
        lead,lc,_=self.person('USER','BACKUP',group=gid,lead=True)
        d=self.post('/api/delegations',{'lead_id':lead['id'],'group_ids':[gid]},mc)['id']
        row=mc.get('/api/state').json()['delegations'][0]
        self.assertEqual(row['end_at'],'')
        with self.server.connection() as c:
            self.assertEqual(len(self.server.effective_delegations(c)),1)
        with patch.object(self.server,'now',return_value='2030-01-01T00:00:00+00:00'):
            with self.server.connection() as c:
                self.assertEqual(len(self.server.effective_delegations(c)),1)
        self.assertTrue(lc.get('/api/state').json()['can_approve'])
        self.assertEqual(mc.delete('/api/delegations/'+d).status_code,200)
        self.assertFalse(lc.get('/api/state').json()['can_approve'])
        self.assertTrue(lc.get('/api/state').json()['user']['is_lead'])

    def test_one_lead_moves_and_independent_delegation(self):
        a=self.account();g1=self.group(a);g2=self.group(a,'G2')
        m,mc,_=self.person('MANAGER','M001',account=a,groups=[g1,g2])
        lead,lc,d1=self.person('USER','L001',group=g1,lead=True)
        other,oc,d2=self.person('USER','L002',group=g2)
        self.assertEqual(self.client.put('/api/people/'+other['id'],json={**d2,'is_lead':True,'transfer_reason':'Designar backup directo'}).status_code,422)
        tech,tc,d3=self.person('USER','T003',group=g1)
        self.assertEqual(self.client.put('/api/people/'+tech['id'],json={**d3,'is_lead':True,'transfer_reason':'Designar backup directo'}).status_code,422)
        self.assertEqual(self.client.put('/api/people/'+lead['id'],json={**d1,'group_id':g2}).status_code,200)
        self.assertEqual(self.client.put('/api/people/'+lead['id'],json={**d1,'group_id':g2,'transfer_reason':'Mover backup dentro del Manager'}).status_code,200)
        self.assertEqual(self.client.put('/api/people/'+lead['id'],json={**d1,'transfer_reason':'Regresar backup'}).status_code,200)
        self.login(lc,'L001')
        self.assertFalse(lc.get('/api/state').json()['can_approve'])
        future=(self.server.datetime.now(self.server.timezone.utc)+self.server.timedelta(hours=3)).isoformat()
        start=(self.server.datetime.now(self.server.timezone.utc)-self.server.timedelta(minutes=1)).isoformat()
        d=self.post('/api/delegations',{'lead_id':lead['id'],'group_ids':[g1],'start_at':start,'end_at':future},mc)
        self.assertTrue(lc.get('/api/state').json()['can_approve'])
        own=self.ot(lc)
        self.assertEqual(lc.post('/api/approvals/ot/'+own,json={'decision':'APPROVED'}).status_code,403)
        target=self.ot(tc)
        self.approve(lc,'ot',target)
        denied=self.ot(oc)
        self.approve(lc,'ot',denied)
        self.assertEqual(mc.delete('/api/delegations/'+d['id']).status_code,200)
        s=lc.get('/api/state').json();self.assertFalse(s['can_approve']);self.assertTrue(s['user']['is_lead'])
        self.assertEqual(self.client.put('/api/people/'+lead['id'],json={**d1,'is_lead':False,'transfer_reason':'Asignar como Técnico'}).status_code,200)
        self.assertEqual(self.client.put('/api/people/'+tech['id'],json={**d3,'is_lead':True,'transfer_reason':'Designar backup directo'}).status_code,200)

    def test_ot_type_is_derived_from_period_on_create_and_edit(self):
        _,_,_,_,_,tc=self.organization()
        current=self.server.datetime.now(self.server.timezone.utc)
        past_start=(current-self.server.timedelta(hours=4)).isoformat()
        past_end=(current-self.server.timedelta(hours=3)).isoformat()
        future_start=(current+self.server.timedelta(hours=3)).isoformat()
        future_end=(current+self.server.timedelta(hours=4)).isoformat()
        past=self.ot(tc,past_start,past_end,capture_type='PLANNED')
        future=self.ot(tc,future_start,future_end,capture_type='HISTORICAL')
        records={o['id']:o for o in tc.get('/api/state').json()['overtime']}
        self.assertEqual(records[past]['capture_type'],'HISTORICAL');self.assertEqual(records[past]['execution'],'COMPLETED')
        self.assertEqual(records[future]['capture_type'],'PLANNED');self.assertIsNone(records[future]['start_at'])
        data={'scheduled_start':(current-self.server.timedelta(hours=2)).isoformat(),'scheduled_end':(current-self.server.timedelta(hours=1)).isoformat(),'reference':'CHGAUTO','description':'Periodo corregido','justification':'Servicio','impact':'Continuidad','in_lieu':'No aplica'}
        self.assertEqual(tc.put('/api/overtime/'+future,json=data).status_code,200)
        changed=next(o for o in tc.get('/api/state').json()['overtime'] if o['id']==future)
        self.assertEqual(changed['capture_type'],'HISTORICAL');self.assertEqual(changed['execution'],'COMPLETED');self.assertEqual(changed['status'],'SUBMITTED')
        invalid=tc.post('/api/overtime',json={**data,'scheduled_start':(current-self.server.timedelta(minutes=5)).isoformat(),'scheduled_end':future_end})
        self.assertEqual(invalid.status_code,422);self.assertIn('Timer',invalid.json()['detail'])

    def test_guard_rule_effective_date_is_compared_with_requested_start(self):
        _,_,_,mc,_,tc=self.organization()
        rule=self.post('/api/rules',{'name':'Reglas_Mainframe_v1','effective_from':'2026-10-26','segments':[{'day':0,'start':'02:00','end':'04:00'},{'day':0,'start':'20:00','end':'00:00'}]},mc)['id']
        invalid=tc.post('/api/guards',json={'rule_id':rule,'start_date':'2026-09-26','end_date':'2026-10-02'})
        self.assertEqual(invalid.status_code,422)
        self.assertIn('2026-10-26',invalid.json()['detail']);self.assertIn('2026-09-26',invalid.json()['detail'])
        self.assertEqual(len(tc.get('/api/state').json()['guards']),0)
        self.post('/api/guards',{'rule_id':rule,'start_date':'2026-10-26','end_date':'2026-11-01'},tc)
        self.assertEqual(len(tc.get('/api/state').json()['guards']),1)

    def test_multiple_weekday_intervals_full_sunday_and_net_guard(self):
        _, _, _, mc, u, uc = self.organization()
        segments = [{'day': day, 'start': start, 'end': end} for day in range(5) for start, end in [('02:00', '04:00'), ('20:00', '00:00')]]
        segments += [{'day': 5, 'start': '00:00', 'end': '08:00'}, {'day': 5, 'start': '15:00', 'end': '00:00'}, {'day': 6, 'start': '00:00', 'end': '00:00'}]
        rule = self.post('/api/rules', {'name': 'Cobertura por intervalos', 'effective_from': '2026-01-01', 'segments': segments}, mc)['id']
        saved = next(r for r in mc.get('/api/state').json()['rules'] if r['id'] == rule)
        self.assertEqual(saved['segments'], segments)
        self.assertEqual(len(saved['segments']), 13)
        guard = self.post('/api/guards', {'rule_id': rule, 'start_date': '2026-01-05', 'end_date': '2026-01-11'}, uc)['id']
        self.approve(mc, 'guard', guard)
        result = uc.get('/api/accounting?month=2026-01').json()
        self.assertEqual(result['guard_minutes'], 71*60)
        self.assertEqual(sum(r['minutes'] for r in result['rows'] if r['date']=='2026-01-11'), 24*60)
        overtime = self.ot(uc, '2026-01-05T20:30:00-06:00', '2026-01-05T22:30:00-06:00')
        self.approve(mc, 'ot', overtime)
        result = uc.get('/api/accounting?month=2026-01').json()
        self.assertEqual(result['ot_minutes'], 120)
        self.assertEqual(result['guard_minutes'], 69*60)
        self.assertEqual(mc.post('/api/rules', json={'name':'Invalid', 'effective_from':'2026-01-01', 'segments':[{'day':0,'start':'30:00','end':'04:00'}]}).status_code,422)
        self.assertEqual(mc.post('/api/rules', json={'name':'Invalid', 'effective_from':'2026-01-01', 'segments':'bad'}).status_code,422)

    def test_single_backup_atomic_replacement_and_history(self):
        a = self.account(); g1 = self.group(a); g2 = self.group(a, 'G2')
        manager, mc, _ = self.person('MANAGER', 'M001', account=a, groups=[g1,g2])
        first, c1, d1 = self.person('USER', 'L001', group=g1, lead=True)
        second, c2, d2 = self.person('USER', 'L002', group=g2)
        start=(self.server.datetime.now(self.server.timezone.utc)-self.server.timedelta(minutes=1)).isoformat()
        end=(self.server.datetime.now(self.server.timezone.utc)+self.server.timedelta(hours=2)).isoformat()
        data={'lead_id':first['id'],'group_ids':[g1,g2],'start_at':start,'end_at':end}
        old=self.post('/api/delegations',data,mc)['id']
        self.assertTrue(c1.get('/api/state').json()['can_approve'])
        self.assertEqual(mc.post('/api/delegations',json={**data,'lead_id':second['id']}).status_code,422)
        # Invalid replacements roll back without revoking the current backup.
        self.assertEqual(mc.post('/api/delegations',json={**data,'lead_id':'missing','replace_current':True}).status_code,404)
        self.assertTrue(c1.get('/api/state').json()['can_approve'])
        # Admin changes the unique Lead; removing the capability revokes the old delegation.
        self.assertEqual(self.client.put('/api/people/'+first['id'],json={**d1,'is_lead':False,'transfer_reason':'Asignar como Técnico'}).status_code,200)
        self.assertEqual(self.client.put('/api/people/'+second['id'],json={**d2,'is_lead':True,'transfer_reason':'Designar backup directo'}).status_code,200)
        latest=self.post('/api/delegations',{**data,'lead_id':second['id'],'replace_current':True},mc)['id']
        self.login(c1,'L001');self.assertFalse(c1.get('/api/state').json()['can_approve'])
        self.login(c2,'L002');self.assertTrue(c2.get('/api/state').json()['can_approve'])
        records=mc.get('/api/state').json()['delegations']
        self.assertEqual(len(records),2)
        self.assertTrue(next(d for d in records if d['id']==old)['revoked_at'])
        self.assertEqual(len([d for d in records if not d['revoked_at']]),1)
        self.assertEqual(c1.post('/api/delegations',json=data).status_code,403)
        self.assertEqual(mc.delete('/api/delegations/'+latest).status_code,200)
        self.assertFalse(c2.get('/api/state').json()['can_approve'])
        self.assertFalse(c1.get('/api/state').json()['user']['is_lead'])
        self.assertTrue(c2.get('/api/state').json()['user']['is_lead'])

    def test_legacy_multiple_backups_never_grant_simultaneous_authority(self):
        a=self.account();g1=self.group(a);g2=self.group(a,'G2')
        manager,mc,_=self.person('MANAGER','M001',account=a,groups=[g1,g2])
        first,c1,_=self.person('USER','L001',group=g1,lead=True)
        second,c2,_=self.person('USER','L002',group=g2)
        # Reproduce a legacy pre-1.3 database, not a permitted new API assignment.
        with self.server.connection(True) as c:
            c.execute('UPDATE users SET is_lead=1 WHERE id=?',(second['id'],))
        start=(self.server.datetime.now(self.server.timezone.utc)-self.server.timedelta(minutes=1)).isoformat()
        end=(self.server.datetime.now(self.server.timezone.utc)+self.server.timedelta(hours=2)).isoformat()
        import json
        with self.server.connection(True) as c:
            for identity,lead in [('old',first),('latest',second)]:
                c.execute('INSERT INTO delegations VALUES(?,?,?,?,?,?,NULL)',(identity,manager['id'],lead['id'],json.dumps([g1,g2]),start,end))
        self.server.init_db()
        self.assertEqual(len(mc.get('/api/state').json()['delegations']),2)
        self.login(c1,'L001');self.assertFalse(c1.get('/api/state').json()['can_approve'])
        self.login(c2,'L002');self.assertTrue(c2.get('/api/state').json()['can_approve'])
        self.assertEqual(mc.delete('/api/delegations/latest').status_code,200)
        self.login(c1,'L001');self.assertFalse(c1.get('/api/state').json()['can_approve'])
        self.assertFalse(c2.get('/api/state').json()['can_approve'])
        self.assertEqual(len(mc.get('/api/state').json()['delegations']),2)

    def test_accounting_midnight_guard_net_unique_ot_and_excel(self):
        _,_,_,mc,u,uc=self.organization()
        rule=self.post('/api/rules',{'name':'Noche','effective_from':'2026-01-01','segments':[{'day':0,'start':'22:00','end':'06:00'}]},mc)['id']
        guard=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-05','end_date':'2026-01-05'},uc)['id']
        guard2=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-05','end_date':'2026-01-05'},uc)['id']
        self.assertNotEqual(guard,guard2)
        self.approve(mc,'guard',guard);self.approve(mc,'guard',guard2)
        ot=self.ot(uc);self.approve(mc,'ot',ot)
        duplicate=uc.post('/api/overtime',json={'capture_type':'HISTORICAL','scheduled_start':'2026-01-05T23:00:00-06:00','scheduled_end':'2026-01-06T01:00:00-06:00','reference':'INC123','description':'Duplicada'})
        self.assertEqual(duplicate.status_code,409)
        r=uc.get('/api/accounting?month=2026-01').json()
        self.assertEqual(r['ot_minutes'],120);self.assertEqual(r['guard_minutes'],360)
        self.assertEqual(len([x for x in r['rows'] if x['kind']=='OT']),2)
        self.assertEqual({x['date'] for x in r['rows']},{'2026-01-05','2026-01-06'})
        ot2=self.ot(uc,'2026-01-06T00:30:00-06:00','2026-01-06T02:00:00-06:00');self.approve(mc,'ot',ot2)
        r=uc.get('/api/accounting?month=2026-01').json();self.assertEqual(r['ot_minutes'],180);self.assertEqual(r['guard_minutes'],300)
        response=uc.get('/api/export.xlsx?month=2026-01');self.assertEqual(response.status_code,200)
        wb=load_workbook(io.BytesIO(response.content));self.assertEqual(wb.sheetnames,['Tiempo aprobado','Resumen'])
        values=list(wb['Tiempo aprobado'].values)[1:];self.assertAlmostEqual(sum(row[8] for row in values),8)
        filtered=uc.get('/api/export.xlsx?month=2026-01&kind=OT')
        wb2=load_workbook(io.BytesIO(filtered.content));self.assertTrue(all(row[4]=='Onsite Overtime' for row in list(wb2.active.values)[1:]))

    def test_timer_emergency_documentation_reapproval_and_history(self):
        _,_,_,mc,u,uc=self.organization()
        with patch.object(self.server,'now',return_value='2026-01-06T05:00:00+00:00'):
            timer=self.post('/api/timers/start',{},uc)['id']
        self.assertEqual(uc.post('/api/timers/start',json={}).status_code,422)
        self.assertEqual(mc.post(f'/api/timers/{timer}/stop',json={}).status_code,403)
        with patch.object(self.server,'now',return_value='2026-01-06T07:00:00+00:00'):
            self.post(f'/api/timers/{timer}/stop',{},uc)
        state=uc.get('/api/state').json();o=state['overtime'][0];self.assertEqual(o['status'],'DOCUMENTATION');self.assertFalse(state['timers'])
        self.assertEqual(mc.post('/api/approvals/ot/'+o['id'],json={'decision':'APPROVED'}).status_code,422)
        data={'reference':'INC1000','description':'Emergencia','justification':'Servicio','impact':'Interrupción','in_lieu':'No aplica'}
        self.assertEqual(uc.put('/api/overtime/'+o['id'],json=data).status_code,200)
        self.approve(mc,'ot',o['id'])
        self.assertEqual(uc.get('/api/accounting?month=2026-01').json()['ot_minutes'],120)
        self.assertEqual(uc.put('/api/overtime/'+o['id'],json={**data,'description':'Actualizada'}).status_code,200)
        self.assertEqual(uc.get('/api/accounting?month=2026-01').json()['ot_minutes'],0)
        self.approve(mc,'ot',o['id'])
        self.assertTrue(any(a['action']=='EDIT_OT_REOPEN_APPROVAL' for a in uc.get('/api/state').json()['audit']))
        self.assertTrue(any(a['action']=='APPROVED' for a in uc.get('/api/state').json()['audit']))

    def test_shared_planned_individual_participations(self):
        _,gid,_,mc,u,uc=self.organization()
        u2,c2,_=self.person('USER','T002',group=gid)
        current=self.server.datetime.now(self.server.timezone.utc)
        a=(current-self.server.timedelta(minutes=5)).isoformat();b=(current+self.server.timedelta(hours=1)).isoformat()
        # Existing publications remain usable after removing Manager programming authority.
        payload={'group_id':gid,'scheduled_start':a,'scheduled_end':b,'reference':'CHG9000','description':'Cambio previo'}
        for client in (mc,uc,self.client):
            self.assertEqual(client.post('/api/windows',json=payload).status_code,403)
        w=self.server.uid()
        with self.server.connection(True) as c:
            owner=mc.get('/api/state').json()['user']['id']
            c.execute('INSERT INTO windows VALUES(?,?,?,?,?,?,?,?,?)',(w,gid,owner,a,b,'CHG9000','Cambio previo','APPROVED',self.server.now()))
        one=self.post(f'/api/windows/{w}/join',{},uc)['id'];two=self.post(f'/api/windows/{w}/join',{},c2)['id'];self.assertNotEqual(one,two)
        self.assertEqual(uc.post(f'/api/windows/{w}/join',json={}).status_code,409)
        t1=self.post('/api/timers/start',{'overtime_id':one},uc)['id'];t2=self.post('/api/timers/start',{'overtime_id':two},c2)['id']
        self.assertEqual(len(mc.get('/api/state').json()['timers']),2)
        self.post(f'/api/timers/{t1}/stop',{},uc);self.post(f'/api/timers/{t2}/stop',{},c2)
        self.assertTrue(all(o['status']=='SUBMITTED' for o in mc.get('/api/state').json()['overtime']))

    def test_activation_recovery_password_and_session_invalidation(self):
        _,_,_,mc,u,uc=self.organization()
        self.post('/api/recoveries',{'identity':'T001'},TestClient(self.server.app))
        r=mc.get('/api/state').json()['recoveries'][0]
        code=self.post('/api/recoveries/'+r['id']+'/issue',{},mc)['code']
        public=TestClient(self.server.app)
        self.post('/api/recoveries/complete',{'identity':'T001','code':code,'password':'NewPassword2026!'},public)
        self.assertEqual(uc.get('/api/state').status_code,401)
        self.assertEqual(public.post('/api/recoveries/complete',json={'identity':'T001','code':code,'password':'Different2026!'}).status_code,422)
        self.login(uc,'T001','NewPassword2026!')
        self.assertEqual(uc.put('/api/password',json={'current':'wrong','password':'FinalPassword2026!'}).status_code,422)
        self.assertEqual(uc.put('/api/password',json={'current':'NewPassword2026!','password':'FinalPassword2026!'}).status_code,200)
        self.assertEqual(uc.get('/api/state').status_code,401)
        self.login(uc,'T001','FinalPassword2026!')

    def analytics_fixture(self):
        a,g1,m,mc,u1,c1=self.organization()
        g2=self.group(a,'G2')
        result=self.client.put('/api/groups/'+g2,json={'account_id':a,'code':'G2','name':'Grupo G2','manager_id':m['id']})
        self.assertEqual(result.status_code,200,result.text)
        u2,c2,_=self.person('USER','T002',group=g2)
        self.person('USER','T003',group=g1)
        events=[(c1,'2026-01-05T23:00:00-06:00','2026-01-06T01:00:00-06:00','EMERGENCY'),
                (c1,'2026-01-06T00:30:00-06:00','2026-01-06T02:30:00-06:00','PLANNED'),
                (c1,'2026-01-06T03:00:00-06:00','2026-01-06T04:00:00-06:00','HISTORICAL'),
                (c2,'2026-01-06T10:00:00-06:00','2026-01-06T13:00:00-06:00','HISTORICAL')]
        # Known persisted categories in a disposable DB. Execution flows are tested separately.
        for client,start,end,category in events:
            identity=self.ot(client,start,end);self.approve(mc,'ot',identity)
            with self.server.connection(True) as c:c.execute('UPDATE overtime SET capture_type=? WHERE id=?',(category,identity))
        self.ot(c1,'2026-01-06T15:00:00-06:00','2026-01-06T16:00:00-06:00')  # pending, no hours
        scheduled=self.ot(c1,'2026-01-07T10:00:00-06:00','2026-01-07T11:00:00-06:00');self.approve(mc,'ot',scheduled)
        with self.server.connection(True) as c:c.execute("UPDATE overtime SET capture_type='PLANNED',execution='SCHEDULED',start_at=NULL,end_at=NULL WHERE id=?",(scheduled,))
        rule=self.post('/api/rules',{'name':'Noche','effective_from':'2026-01-01','segments':[{'day':0,'start':'22:00','end':'06:00'}]},mc)['id']
        guard=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-05','end_date':'2026-01-05'},c1)['id'];self.approve(mc,'guard',guard)
        return a,g1,g2,m,mc,u1,c1,u2,c2

    def test_analytics_breakdown_reconciles_groups_users_and_payroll(self):
        _,g1,g2,_,mc,u1,c1,u2,c2=self.analytics_fixture()
        response=mc.get('/api/analytics?month=2026-01');self.assertEqual(response.status_code,200,response.text)
        report=response.json();summary=report['summary']
        self.assertEqual(summary,{'EMERGENCY':120,'PLANNED':90,'HISTORICAL':240,'GUARD':210,'ot_minutes':450,'total_minutes':660})
        self.assertEqual(sum(g['minutes']['total_minutes'] for g in report['groups']),660)
        self.assertEqual(sum(u['minutes']['total_minutes'] for u in report['users']),660)
        self.assertEqual(sum(d['minutes']['total_minutes'] for d in report['days']),660)
        self.assertEqual(len(report['users']),3)
        self.assertEqual(next(g for g in report['groups'] if g['id']==g1)['minutes']['total_minutes'],480)
        self.assertEqual(next(g for g in report['groups'] if g['id']==g2)['minutes']['total_minutes'],180)
        accounting=mc.get('/api/accounting?month=2026-01').json()
        self.assertEqual(summary['ot_minutes'],accounting['ot_minutes']);self.assertEqual(summary['GUARD'],accounting['guard_minutes'])
        personal=c1.get('/api/my-analytics?month=2026-01').json()
        self.assertEqual(personal['mode'],'PERSONAL');self.assertEqual(personal['summary']['total_minutes'],480)
        self.assertEqual([u['id'] for u in personal['users']],[u1['id']])
        self.assertTrue(all(r['user_id']==u1['id'] for r in personal['rows']))
        filtered=mc.get('/api/analytics?month=2026-01&group_id='+g2).json()
        self.assertEqual(filtered['summary']['total_minutes'],180);self.assertEqual(len(filtered['users']),1)
        empty=mc.get('/api/analytics?month=2026-02').json();self.assertEqual(empty['summary']['total_minutes'],0);self.assertEqual(empty['rows'],[])
        self.assertEqual(mc.get('/api/analytics?month=invalid').status_code,422)
        self.assertEqual(mc.get('/api/analytics?month=').status_code,422)

    def test_analytics_scope_export_delegation_manual_revocation(self):
        a,g1,g2,m,mc,u1,c1,u2,c2=self.analytics_fixture()
        self.assertEqual(c1.get('/api/analytics?month=2026-01').status_code,403)
        self.assertEqual(c1.get('/api/analytics.xlsx?month=2026-01').status_code,403)
        self.assertEqual(self.client.get('/api/analytics?month=2026-01').status_code,403)
        self.assertEqual(c1.get('/api/my-analytics?month=2026-01&user_id='+u2['id']).status_code,422)
        self.assertEqual(c1.get('/api/my-analytics.xlsx?month=2026-01&group_id='+g2).status_code,422)
        foreign=self.account('FOREIGN');foreign_group=self.group(foreign,'FG')
        self.assertEqual(mc.get('/api/analytics?month=2026-01&group_id='+foreign_group).status_code,403)
        lead,lc,_=self.person('USER','L001',group=g1,lead=True)
        self.assertEqual(lc.get('/api/analytics?month=2026-01').status_code,403)
        current=self.server.datetime.now(self.server.timezone.utc)
        d=self.post('/api/delegations',{'lead_id':lead['id'],'group_ids':[g2],'start_at':(current-self.server.timedelta(minutes=1)).isoformat(),'end_at':(current+self.server.timedelta(hours=2)).isoformat()},mc)['id']
        report=lc.get('/api/analytics?month=2026-01').json()
        self.assertEqual({g['id'] for g in report['groups']},{g1,g2});self.assertIn(u2['id'],[u['id'] for u in report['users']])
        self.assertEqual(lc.get('/api/analytics?month=2026-01&group_id='+g1).status_code,200)
        self.assertEqual(lc.get('/api/analytics?month=2026-01&user_id='+lead['id']).status_code,200)
        self.assertEqual(lc.get('/api/my-analytics?month=2026-01').status_code,200)
        with patch.object(self.server,'now',return_value=(current+self.server.timedelta(hours=3)).isoformat()):
            self.assertEqual(lc.get('/api/analytics?month=2026-01').status_code,200)
            self.assertEqual(lc.get('/api/analytics.xlsx?month=2026-01').status_code,200)
            self.assertEqual(lc.get('/api/my-analytics?month=2026-01').status_code,200)
        self.assertEqual(mc.delete('/api/delegations/'+d).status_code,200)
        self.assertEqual(lc.get('/api/analytics?month=2026-01').status_code,403)
        exported=mc.get('/api/analytics.xlsx?month=2026-01');self.assertEqual(exported.status_code,200)
        wb=load_workbook(io.BytesIO(exported.content));self.assertEqual(wb.sheetnames,['Resumen','Por grupo','Por técnico','Detalle','Registros'])
        self.assertEqual(sum(row[-1] for row in list(wb['Por grupo'].values)[1:]),11)
        personal=c1.get('/api/my-analytics.xlsx?month=2026-01');self.assertEqual(personal.status_code,200)
        wb=load_workbook(io.BytesIO(personal.content));values=list(wb['Por técnico'].values)[1:]
        self.assertEqual(len(values),1);self.assertEqual(values[0][0],'T001');self.assertEqual(values[0][-1],8)

    def test_analytics_respects_month_boundaries_and_zero_hours(self):
        _,_,_,mc,u,uc=self.organization()
        ot=self.ot(uc,'2026-01-31T23:00:00-06:00','2026-02-01T01:00:00-06:00');self.approve(mc,'ot',ot)
        with self.server.connection(True) as c:c.execute("UPDATE overtime SET capture_type='EMERGENCY' WHERE id=?",(ot,))
        rule=self.post('/api/rules',{'name':'Fin de mes','effective_from':'2026-01-01','segments':[{'day':5,'start':'22:00','end':'06:00'}]},mc)['id']
        guard=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-31','end_date':'2026-01-31'},uc)['id'];self.approve(mc,'guard',guard)
        january=mc.get('/api/analytics?month=2026-01').json();february=mc.get('/api/analytics?month=2026-02').json()
        self.assertEqual(january['summary']['EMERGENCY'],60);self.assertEqual(january['summary']['GUARD'],60)
        self.assertEqual(february['summary']['EMERGENCY'],60);self.assertEqual(february['summary']['GUARD'],300)
        self.assertEqual({r['date'] for r in january['rows']},{'2026-01-31'});self.assertEqual({r['date'] for r in february['rows']},{'2026-02-01'})

    def test_manager_manages_technicians_admin_manages_higher_roles(self):
        a,g,m,mc,u,uc=self.organization()
        data={'role':'USER','employee_id':'T009','first':'Persona','last':'Nueva','email':'new@example.test','group_id':g}
        self.assertEqual(mc.post('/api/people',json=data).status_code,200)
        self.assertEqual(mc.put('/api/people/'+u['id'],json={**data,'employee_id':'T001'}).status_code,200)
        self.assertEqual(uc.put('/api/people/'+u['id'],json={**data,'employee_id':'T001'}).status_code,403)
        self.assertEqual(self.client.put('/api/people/'+u['id'],json={**data,'employee_id':'T001'}).status_code,403)
        # Personal contact details remain self-service; hierarchy and Employee ID are not profile fields.
        self.assertEqual(uc.put('/api/profile',json={'first':'Mi','last':'Perfil','email':'own@example.test'}).status_code,200)
        self.assertEqual(uc.get('/api/state').json()['user']['group_id'],g)

    def test_recovery_routes_technicians_to_manager_managers_to_admin(self):
        a,g,m,mc,u,uc=self.organization()
        public=TestClient(self.server.app)
        self.post('/api/recoveries',{'identity':'T001'},public)
        tech_request=next(r for r in mc.get('/api/state').json()['recoveries'] if r['user_id']==u['id'])
        self.assertNotIn(tech_request['id'],{r['id'] for r in self.client.get('/api/state').json()['recoveries']})
        self.assertEqual(self.client.post('/api/recoveries/'+tech_request['id']+'/issue',json={}).status_code,403)
        lead,lc,_=self.person('USER','L001',group=g,lead=True)
        current=self.server.datetime.now(self.server.timezone.utc)
        self.post('/api/delegations',{'lead_id':lead['id'],'group_ids':[g],'start_at':(current-self.server.timedelta(minutes=1)).isoformat(),'end_at':(current+self.server.timedelta(hours=2)).isoformat()},mc)
        self.assertTrue(lc.get('/api/state').json()['can_approve'])
        self.assertEqual(lc.get('/api/state').json()['recoveries'],[])
        self.assertEqual(lc.post('/api/recoveries/'+tech_request['id']+'/issue',json={}).status_code,403)
        self.post('/api/recoveries',{'identity':'M001'},public)
        manager_request=next(r for r in self.client.get('/api/state').json()['recoveries'] if r['user_id']==m['id'])
        self.assertNotIn(manager_request['id'],{r['id'] for r in mc.get('/api/state').json()['recoveries']})
        self.post('/api/recoveries/'+manager_request['id']+'/issue',{},self.client)
        self.assertEqual(mc.post('/api/recoveries/'+manager_request['id']+'/view',json={}).status_code,403)
        self.post('/api/recoveries',{'identity':'admin1'},public)
        admin2=TestClient(self.server.app);self.login(admin2,'admin2','admin22026')
        own_request=next(r for r in admin2.get('/api/state').json()['recoveries'] if r['user_id']==self.client.get('/api/state').json()['user']['id'])
        self.assertEqual(self.client.post('/api/recoveries/'+own_request['id']+'/issue',json={}).status_code,403)
        self.post('/api/recoveries/'+own_request['id']+'/issue',{},admin2)
        orphan_group=self.group(a,'UNASSIGNED')
        orphan,_,_=self.person('USER','T099',group=orphan_group)
        self.post('/api/recoveries',{'identity':'T099'},public)
        orphan_request=next(r for r in self.client.get('/api/state').json()['recoveries'] if r['user_id']==orphan['id'])
        self.post('/api/recoveries/'+orphan_request['id']+'/issue',{},self.client)

    def test_recovery_code_view_regenerate_encryption_expiry_and_single_use(self):
        _,_,_,mc,u,uc=self.organization();public=TestClient(self.server.app)
        self.post('/api/recoveries',{'identity':'T001'},public)
        r=mc.get('/api/state').json()['recoveries'][0];path='/api/recoveries/'+r['id']
        first=self.post(path+'/issue',{},mc)
        shown=self.post(path+'/view',{},mc)
        self.assertEqual(first['code'],shown['code']);self.assertEqual(first['expires_at'],shown['expires_at'])
        self.assertNotIn(first['code'],mc.get('/api/state').text)
        self.assertTrue(mc.get('/api/state').json()['recoveries'][0]['code_available'])
        self.assertEqual(uc.post(path+'/view',json={}).status_code,403)
        self.assertEqual(self.client.post(path+'/view',json={}).status_code,403)
        with self.server.connection() as c:
            encrypted=c.execute('SELECT encrypted_code FROM recovery_code_vault WHERE recovery_id=?',(r['id'],)).fetchone()[0]
            self.assertNotIn(first['code'].encode(),encrypted)
        self.assertTrue(self.server.RECOVERY_KEY.is_file())
        self.assertEqual(self.server.RECOVERY_KEY.stat().st_mode & 0o777,0o600)
        self.post('/api/recoveries',{'identity':'T001'},public)
        self.assertEqual(len(mc.get('/api/state').json()['recoveries']),1)
        second=self.post(path+'/issue',{},mc);self.assertNotEqual(first['code'],second['code'])
        self.assertEqual(public.post('/api/recoveries/complete',json={'identity':'T001','code':first['code'],'password':'NewPassword2026!'}).status_code,422)
        self.assertEqual(self.post(path+'/view',{},mc)['code'],second['code'])
        # Legacy codes remain usable by their owners but cannot be recovered from hashes.
        with self.server.connection(True) as c:c.execute('DELETE FROM recovery_code_vault WHERE recovery_id=?',(r['id'],))
        self.assertFalse(mc.get('/api/state').json()['recoveries'][0]['code_available'])
        self.assertEqual(mc.post(path+'/view',json={}).status_code,409)
        third=self.post(path+'/issue',{},mc)
        with self.server.connection(True) as c:c.execute('UPDATE recoveries SET expires_at=? WHERE id=?',((self.server.datetime.now(self.server.timezone.utc)-self.server.timedelta(minutes=1)).isoformat(),r['id']))
        self.assertEqual(mc.post(path+'/view',json={}).status_code,422)
        fourth=self.post(path+'/issue',{},mc)
        self.post('/api/recoveries/complete',{'identity':'T001','code':fourth['code'],'password':'NewPassword2026!'},public)
        self.assertEqual(mc.post(path+'/view',json={}).status_code,422)
        self.assertEqual(public.post('/api/recoveries/complete',json={'identity':'T001','code':fourth['code'],'password':'AnotherPassword2026!'}).status_code,422)
        with self.server.connection() as c:self.assertEqual(c.execute('SELECT count(*) FROM recovery_code_vault').fetchone()[0],0)
        self.assertEqual(uc.get('/api/state').status_code,401)
        self.login(uc,'T001','NewPassword2026!')

    def test_authorization_no_state_overwrite_and_external_origin(self):
        a,g,m,mc,u,uc=self.organization()
        a2=self.account('SECRET');g2=self.group(a2,'SECRET-G')
        outsider,oc,_=self.person('USER','PRIVATE',group=g2)
        self.assertNotIn(outsider['id'],{p['id'] for p in uc.get('/api/state').json()['people']})
        self.assertEqual(uc.get('/api/accounting?month=2026-01&user_id='+outsider['id']).status_code,403)
        self.assertEqual(uc.put('/api/state',json={}).status_code,405)
        self.assertEqual(self.client.post('/api/accounts',json={'code':'BAD','name':'Bad'},headers={'origin':'https://evil.test'}).status_code,403)
        self.assertEqual(self.client.get('/').headers['x-frame-options'],'DENY')
        for p in ['server.py','operational_pulse.db','operational_pulse.key','requirements.txt']:
            self.assertEqual(self.client.get('/'+p).status_code,404)

    def preview(self, identity):
        r=self.client.get('/api/admin/people/'+identity+'/impact')
        self.assertEqual(r.status_code,200,r.text)
        return r.json()

    def test_active_transfers_manager_lead_conflicts_and_authorization(self):
        a,g,m,mc,u,uc=self.organization()
        g2=self.group(a,'G2'); m2,c2,_=self.person('MANAGER','M002',account=a,groups=[g2])
        lead,lc,ld=self.person('USER','L001',group=g,lead=True)
        lead2,l2c,l2d=self.person('USER','L002',group=g2,lead=True)
        original=uc.get('/api/state').json()['user']
        # A different Group under the same Manager cannot have a second Lead.
        g3=self.group(a,'G3',m['id'])
        blocked=self.client.post('/api/people',json={'role':'USER','employee_id':'L003','first':'Bad','last':'Lead','email':'bad@test','group_id':g3,'is_lead':True})
        self.assertEqual(blocked.status_code,422)
        request={**original,'group_id':g2,'transfer_reason':'Reorganización activa'}
        self.assertEqual(mc.put('/api/people/'+u['id'],json=request).status_code,403)
        self.assertEqual(self.client.put('/api/people/'+u['id'],json=request).status_code,403)
        transfer=self.post('/api/team/transfers',{'user_id':u['id'],'target_manager_id':m2['id'],'reason':'Reorganización activa'},mc)['id']
        self.post('/api/team/transfers/'+transfer+'/decision',{'decision':'APPROVED','group_id':g2,'reason':'Aceptado por destino'},c2)
        self.assertEqual(uc.get('/api/state').status_code,401)
        self.login(uc,'T001');moved=uc.get('/api/state').json()['user']
        self.assertEqual((moved['id'],moved['status'],moved['group_id']),(u['id'],'ACTIVE',g2))
        # Both person and whole-Group transfers reject duplicate Manager backups atomically.
        self.assertEqual(self.client.put('/api/people/'+lead['id'],json={**ld,'group_id':g2,'transfer_reason':'Mover Lead'}).status_code,422)
        payload={'account_id':a,'code':'G1','name':'Grupo G1','manager_id':m2['id'],'transfer_reason':'Mover equipo'}
        self.assertEqual(self.client.put('/api/groups/'+g,json=payload).status_code,200)
        impact=self.preview(m['id'])
        self.assertEqual(self.client.post('/api/admin/people/'+m['id']+'/offboard',json={'expected':impact['expected'],'confirm':True,'reason':'Traspaso con conflicto','transfers':{g:m2['id'],g3:m2['id']}}).status_code,422)
        self.assertEqual(mc.get('/api/state').status_code,200)
        self.assertEqual(next(x for x in self.client.get('/api/state').json()['groups'] if x['id']==g)['manager_id'],m2['id'])
        self.assertEqual(self.client.put('/api/people/'+lead2['id'],json={**l2d,'is_lead':False,'transfer_reason':'Asignar como Técnico'}).status_code,200)
        self.assertEqual(self.client.put('/api/groups/'+g,json=payload).status_code,200)
        self.assertEqual(lc.get('/api/state').json()['user']['status'],'ACTIVE')
        self.assertEqual(mc.get('/api/admin/people/'+u['id']+'/impact').status_code,403)
        self.assertEqual(mc.post('/api/admin/people/'+u['id']+'/reactivate',json={'reason':'No'}).status_code,403)

    def test_manager_offboard_transfers_live_team_recovery_and_delegation(self):
        a,g,m,mc,u,uc=self.organization()
        g2=self.group(a,'G2',m['id']); other,oc,_=self.person('USER','T002',group=g2)
        m2,c2,_=self.person('MANAGER','M002',account=a)
        lead,lc,_=self.person('USER','L001',group=g,lead=True)
        current=self.server.datetime.now(self.server.timezone.utc)
        self.post('/api/delegations',{'lead_id':lead['id'],'group_ids':[g,g2],'start_at':(current-self.server.timedelta(minutes=1)).isoformat(),'end_at':(current+self.server.timedelta(hours=1)).isoformat()},mc)
        active=self.post('/api/timers/start',{},uc)
        overtime=self.ot(oc)
        self.post('/api/recoveries',{'identity':'T002'},oc)
        rid=mc.get('/api/state').json()['recoveries'][0]['id']
        reset=self.post('/api/recoveries/'+rid+'/issue',{},mc)
        before=self.preview(m['id'])
        request={'lead_manager_id':m2['id'],'reason':'Cambio de responsable','confirm':True,'expected':before['expected'],'transfers':{g:m2['id'],g2:m2['id']}}
        self.assertEqual(self.client.post('/api/admin/people/'+m['id']+'/offboard',json={**request,'transfers':{g:m2['id']}}).status_code,422)
        self.post('/api/admin/people/'+m['id']+'/offboard',request)
        self.assertEqual(mc.get('/api/state').status_code,401)
        target=c2.get('/api/state').json()
        self.assertEqual({x['id'] for x in target['groups']},{g,g2})
        self.assertTrue(all(x['status']=='ACTIVE' for x in target['people'] if x['role']=='USER'))
        self.assertTrue(any(t['id']==active['id'] and not t['end_at'] for t in target['timers']))
        self.login(lc,'L001')
        self.assertFalse(lc.get('/api/state').json()['can_approve'])
        self.assertEqual(next(r for r in target['recoveries'] if r['id']==rid)['status'],'PENDING')
        self.assertEqual(oc.post('/api/recoveries/complete',json={'identity':'T002','code':reset['code'],'password':'AnotherPassword2026!'}).status_code,422)
        self.approve(c2,'ot',overtime)
        self.assertEqual(len(uc.get('/api/state').json()['timers']),1)
        self.assertEqual(self.client.post('/api/login',json={'identity':'M001','password':'Password2026!'}).status_code,401)

    def test_offboard_stale_preview_new_manager_activation_and_self_protection(self):
        a,g,m,mc,u,uc=self.organization()
        replacement=self.post('/api/people',{'role':'MANAGER','employee_id':'M002','first':'Nuevo','last':'Manager','email':'new@test','account_id':a})
        snapshot=self.preview(m['id']);payload={'expected':snapshot['expected'],'reason':'Reemplazo','confirm':True,'transfers':{g:replacement['id']}}
        path='/api/admin/people/'+m['id']+'/offboard'
        self.assertEqual(self.client.post(path,json=payload).status_code,422)
        self.assertEqual(mc.get('/api/state').status_code,200)
        self.post('/api/activate',{'identity':'M002','code':replacement['activation_code'],'password':'Password2026!'})
        # Adding a team member invalidates the reviewed impact.
        self.person('USER','T002',group=g)
        self.assertEqual(self.client.post(path,json=payload).status_code,409)
        payload['expected']=self.preview(m['id'])['expected'];self.post(path,payload)
        mine=self.client.get('/api/state').json()['user']['id'];impact=self.preview(mine)
        self.assertEqual(self.client.post('/api/admin/people/'+mine+'/offboard',json={'expected':impact['expected'],'reason':'Self','confirm':True,'transfers':{}}).status_code,422)

    def test_pending_access_and_legacy_lead_regularization(self):
        a,g,m,mc,u,uc=self.organization()
        g2=self.group(a,'G2',m['id']);g3=self.group(a,'G3',m['id'])
        first,lc,ld=self.person('USER','L001',group=g,lead=True)
        second,sc,sd=self.person('USER','L002',group=g2)
        third,tc,td=self.person('USER','L003',group=g3)
        with self.server.connection(True) as c:
            c.execute('UPDATE users SET is_lead=1 WHERE id IN (?,?)',(second['id'],third['id']))
        self.server.init_db()
        # Existing data survives, and Admin can reduce even three legacy designations to one.
        self.assertEqual(self.client.put('/api/people/'+second['id'],json={**sd,'is_lead':False,'transfer_reason':'Asignar como Técnico'}).status_code,200)
        self.assertEqual(self.client.put('/api/people/'+third['id'],json={**td,'is_lead':False,'transfer_reason':'Asignar como Técnico'}).status_code,200)
        pending=self.post('/api/people',{'role':'USER','employee_id':'NEW','first':'Nueva','last':'Persona','email':'new@test','group_id':g},mc)
        impact=self.preview(pending['id'])
        self.post('/api/admin/people/'+pending['id']+'/offboard',{'reason':'Alta cancelada','confirm':True,'expected':impact['expected'],'transfers':{}})
        self.assertEqual(self.client.post('/api/activate',json={'identity':'NEW','code':pending['activation_code'],'password':'Password2026!'}).status_code,422)
        revived=self.post('/api/admin/people/'+pending['id']+'/reactivate',{'reason':'Alta retomada'})
        self.assertNotEqual(revived['activation_code'],pending['activation_code'])
        self.post('/api/activate',{'identity':'NEW','code':revived['activation_code'],'password':'Password2026!'})

    def test_technician_offboard_timer_guard_cutoff_and_reactivation(self):
        a,g,m,mc,u,uc=self.organization()
        with patch.object(self.server,'now',return_value='2026-01-06T05:00:00+00:00'):
            rule=self.post('/api/rules',{'name':'Noche','effective_from':'2026-01-01','segments':[{'day':0,'start':'22:00','end':'06:00'},{'day':1,'start':'22:00','end':'06:00'}]},mc)['id']
            guard=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-05','end_date':'2026-01-05'},uc)['id'];self.approve(mc,'guard',guard)
            future=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-06','end_date':'2026-01-06'},uc)['id'];self.approve(mc,'guard',future)
            timer=self.post('/api/timers/start',{},uc)
        with patch.object(self.server,'now',return_value='2026-01-06T07:00:00+00:00'):
            impact=self.preview(u['id']);self.assertEqual(len(impact['documents']),1)
            payload={'reason':'Baja de personal','confirm':True,'expected':impact['expected'],'transfers':{}}
            path='/api/admin/people/'+u['id']+'/offboard'
            self.assertEqual(self.client.post(path,json=payload).status_code,422)
            payload['stop_timers']=True
            self.assertEqual(self.client.post(path,json=payload).status_code,422)
            with self.server.connection() as c:self.assertIsNone(c.execute('SELECT end_at FROM timers WHERE id=?',(timer['id'],)).fetchone()[0])
            oid=impact['documents'][0]['id']
            payload['documents']=[{'id':oid,'reference':'INC123456','description':'Incidente','justification':'Servicio','impact':'Operación','in_lieu':'Sin cobertura'}]
            self.post(path,payload)
            self.assertEqual(uc.get('/api/state').status_code,401)
            self.approve(mc,'ot',oid)
            result=mc.get('/api/accounting?month=2026-01&user_id='+u['id']).json()
            self.assertEqual((result['ot_minutes'],result['guard_minutes']),(120,60))
            state=mc.get('/api/state').json()
            self.assertEqual(next(r for r in state['guards'] if r['id']==future)['status'],'CANCELLED')
            revived=self.post('/api/admin/people/'+u['id']+'/reactivate',{'reason':'Recontratación'})
            self.assertEqual(uc.post('/api/login',json={'identity':'T001','password':'Password2026!'}).status_code,401)
            self.post('/api/activate',{'identity':'T001','code':revived['activation_code'],'password':'NewPassword2026!'},uc)
            self.login(uc,'T001','NewPassword2026!')
            restored=uc.get('/api/state').json();self.assertFalse(restored['user']['is_lead'])
            self.assertEqual(restored['user']['id'],u['id'])
            self.assertFalse(any(not t['end_at'] for t in restored['timers']))
            self.assertEqual(uc.get('/api/accounting?month=2026-01').json()['guard_minutes'],60)


    def test_rule_edit_delete_preserves_existing_guard_snapshots(self):
        a,g,m,mc,u,uc=self.organization()
        m2,c2,_=self.person('MANAGER','M002',account=a)
        original={'name':'Cobertura original','effective_from':'2026-01-01','segments':[{'day':0,'start':'22:00','end':'06:00'}]}
        rule=self.post('/api/rules',original,mc)['id']
        guard=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-05','end_date':'2026-01-05'},uc)['id']
        self.approve(mc,'guard',guard)
        changed={**original,'name':'Cobertura actualizada','segments':[{'day':0,'start':'20:00','end':'00:00'},{'day':1,'start':'02:00','end':'04:00'}]}
        for client in [self.client,uc,c2]:
            self.assertEqual(client.put('/api/rules/'+rule,json=changed).status_code,403)
            self.assertEqual(client.delete('/api/rules/'+rule).status_code,403)
        self.assertEqual(mc.put('/api/rules/'+rule,json={**changed,'segments':[]}).status_code,422)
        self.assertEqual(mc.put('/api/rules/'+rule,json=changed).status_code,200)
        before=uc.get('/api/accounting?month=2026-01').json()
        self.assertEqual(before['guard_minutes'],480)
        old=next(r for r in uc.get('/api/state').json()['guards'] if r['id']==guard)
        self.assertEqual(old['rule_snapshot']['segments'],original['segments'])
        newer=self.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-12','end_date':'2026-01-12'},uc)['id']
        self.assertEqual(next(r for r in uc.get('/api/state').json()['guards'] if r['id']==newer)['rule_snapshot']['segments'],changed['segments'])
        self.assertEqual(mc.delete('/api/rules/'+rule).status_code,200)
        self.assertEqual(uc.get('/api/state').json()['rules'],[])
        self.assertEqual(uc.get('/api/accounting?month=2026-01').json()['guard_minutes'],480)
        self.assertEqual(uc.post('/api/guards',json={'rule_id':rule,'start_date':'2026-01-19','end_date':'2026-01-19'}).status_code,404)
        audit=mc.get('/api/state').json()['audit']
        self.assertTrue(any(r['action']=='EDIT_RULE' for r in audit));self.assertTrue(any(r['action']=='DELETE_RULE' for r in audit))


    def test_dedicated_transfer_preserves_fresh_profile_and_rejects_stale_origin(self):
        a,g,m,mc,u,uc=self.organization()
        g2=self.group(a,'G2',m['id'])
        self.assertEqual(uc.put('/api/profile',json={'first':'Nombre actualizado','last':'Apellido vigente','email':'fresh@example.test'}).status_code,200)
        payload={'expected_group_id':g,'group_id':g2,'transfer_reason':'Reorganización','first':'Dato obsoleto','last':'Viejo','email':'old@test','status':'DISABLED','is_lead':True}
        path='/api/team/people/'+u['id']+'/move'
        self.assertEqual(self.client.post(path,json=payload).status_code,403)
        self.assertEqual(uc.post(path,json=payload).status_code,403)
        self.post(path,payload,mc)
        moved=next(p for p in self.client.get('/api/state').json()['people'] if p['id']==u['id'])
        self.assertEqual((moved['first'],moved['last'],moved['email']),('Nombre actualizado','Apellido vigente','fresh@example.test'))
        self.assertEqual((moved['group_id'],moved['status'],moved['is_lead']),(g2,'ACTIVE',False))
        self.assertEqual(uc.get('/api/state').status_code,401)
        self.assertEqual(mc.post(path,json={**payload,'group_id':g}).status_code,409)
        self.assertEqual(next(p for p in self.client.get('/api/state').json()['people'] if p['id']==u['id'])['group_id'],g2)
        self.assertEqual(mc.post(path,json={'expected_group_id':g2,'group_id':g,'transfer_reason':''}).status_code,422)

    def test_dedicated_transfer_timer_and_lead_constraints_are_atomic(self):
        a,g,m,mc,u,uc=self.organization();g2=self.group(a,'G2',m['id'])
        path='/api/team/people/'+u['id']+'/move';payload={'expected_group_id':g,'group_id':g2,'transfer_reason':'Movimiento solicitado'}
        self.post('/api/timers/start',{},uc)
        self.assertEqual(mc.post(path,json=payload).status_code,422)
        self.assertEqual(next(p for p in self.client.get('/api/state').json()['people'] if p['id']==u['id'])['group_id'],g)
        lead,lc,_=self.person('USER','L001',group=g,lead=True)
        self.assertEqual(mc.post('/api/team/people/'+lead['id']+'/move',json=payload).status_code,403)
        after=next(p for p in self.client.get('/api/state').json()['people'] if p['id']==lead['id'])
        self.assertIsNone(after['group_id']);self.assertEqual(after['manager_id'],m['id']);self.assertTrue(after['is_lead'])


if __name__=='__main__':
    unittest.main(verbosity=2)
