"""Personal backup and legacy recovery with explicit review, no automatic approval."""
import io,os,socket,subprocess,sys,time,urllib.request
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
import test_time_recovery
root=Path(__file__).resolve().parents[1];fixture=test_time_recovery.TimeRecoveryTests();fixture.setUp();proc=None
try:
 source=Path(fixture.f.temp.name)/'respaldo.xlsx';fixture.legacy().save(source)
 with socket.socket() as sock:sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
 base=f'http://127.0.0.1:{port}'
 proc=subprocess.Popen([sys.executable,'-m','uvicorn','server:app','--port',str(port)],cwd=root,env={**os.environ,'PULSE_DB_PATH':str(fixture.f.server.DB)},stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
 for _ in range(100):
  try:urllib.request.urlopen(base+'/api/health',timeout=1);break
  except Exception:time.sleep(.1)
 with sync_playwright() as p:
  b=p.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);ctx=b.new_context(viewport={'width':1440,'height':780},accept_downloads=True);page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  assert ctx.request.post(base+'/api/login',data={'identity':'T001','password':'Password2026!'}).ok
  page.goto(base);page.locator('.sidebar [data-page=time]').click();page.wait_for_function('()=>!busy')
  page.locator('#month').fill('2026-01');page.locator('#month').dispatch_event('change');page.wait_for_function('()=>!busy')
  page.locator('[data-action=time-recovery]').click();page.locator('[name=backup_file]').set_input_files(source);page.locator('#modal [type=submit]').click();page.wait_for_function('()=>!busy')
  expect(page.locator('#modal')).to_contain_text('2 registros disponibles');expect(page.locator('#modal')).to_contain_text('Falta documentación de OT')
  assert fixture.uc.get('/api/state').json()['overtime']==[]
  for width,height in [(1440,780),(390,844)]:
   page.set_viewport_size({'width':width,'height':height});page.wait_for_timeout(150)
   assert page.locator('#modal .dialog-footer').evaluate('el=>el.getBoundingClientRect().bottom<=innerHeight+1')
   assert page.evaluate('document.documentElement.scrollWidth<=innerWidth+1')
  page.locator('[name=confirm]').check();page.locator('#modal [type=submit]').click();page.wait_for_function('()=>!busy');expect(page.locator('#modal')).not_to_be_visible()
  expect(page.locator('[data-action=ot-edit]').first).to_be_visible()
  state=fixture.uc.get('/api/state').json();assert state['overtime'][0]['status']=='DOCUMENTATION';assert state['guards'][0]['status']=='SUBMITTED'
  assert fixture.uc.get('/api/accounting?month=2026-01').json()['ot_minutes']==0
  page.locator('[data-action=time-recovery]').click()
  with page.expect_download() as pending:page.locator('.recovery-download').click()
  dest=Path(fixture.f.temp.name)/'download.xlsx';pending.value.save_as(dest);w=load_workbook(dest);assert w.sheetnames[0]=='Captura';assert w['Respaldo Pulse'].max_row==3
  page.locator('[name=backup_file]').set_input_files(source);page.locator('#modal [type=submit]').click();page.wait_for_function('()=>!busy');expect(page.locator('#modal')).to_contain_text('2 duplicados');expect(page.locator('#modal [type=submit]')).to_be_disabled()
  assert not errors,errors
  b.close();print('PASS: personal backup download with Captura first, legacy preview before persistence, owner confirmation, required OT documentation, pending Guard, duplicate protection, mobile layout and no JS errors')
finally:
 if proc:proc.terminate();proc.wait(timeout=10)
 fixture.tearDown()
