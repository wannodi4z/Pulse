"""Claim uses full Saturday-Friday weeks and the same approved net accounting as ADP."""
import io,unittest
from openpyxl import load_workbook
import test_api
class ClaimCaptureTests(unittest.TestCase):
 def setUp(self):
  self.f=test_api.PulseTests();self.f.setUp();self.a,self.g,self.m,self.mc,self.u,self.uc=self.f.organization()
  rule=self.f.post('/api/rules',{'name':'Cobertura Claim','effective_from':'2025-01-01','segments':[{'day':d,'start':'00:00','end':'08:00'} for d in range(7)]},self.mc)['id']
  guard=self.f.post('/api/guards',{'rule_id':rule,'start_date':'2026-01-03','end_date':'2026-01-09'},self.uc)['id'];self.f.approve(self.mc,'guard',guard)
  for start,end,ref in [('2026-01-03T03:00:00-06:00','2026-01-03T04:30:00-06:00','CHG-90MIN'),('2026-01-31T23:30:00-06:00','2026-02-01T01:00:00-06:00','CHG-CROSS'),('2025-12-27T10:00:00-06:00','2025-12-27T10:30:00-06:00','CHG-PREV')]:
   ot=self.f.ot(self.uc,start=start,end=end,reference=ref);self.f.approve(self.mc,'ot',ot)
  self.f.ot(self.uc,start='2026-01-04T06:00:00-06:00',end='2026-01-04T07:00:00-06:00',reference='CHG-PENDING')
 def tearDown(self):self.f.tearDown()
 def test_full_weeks_decimal_units_and_month_totals(self):
  r=self.uc.get('/api/accounting?month=2026-01&view=claim');self.assertEqual(r.status_code,200,r.text);data=r.json()
  self.assertEqual((data['period_start'],data['period_end']),('2025-12-27','2026-02-06'))
  self.assertEqual(data['month_ot_minutes'],120);self.assertEqual(data['ot_minutes'],210)
  self.assertEqual(data['month_guard_minutes'],3270)
  sat=[r for r in data['rows'] if r['date']=='2026-01-03']
  self.assertEqual(sum(r['hours'] for r in sat if r['kind']=='OT'),1.5)
  self.assertEqual(sum(r['hours'] for r in sat if r['kind']=='GUARD'),6.5)
  self.assertNotIn('CHG-PENDING',[r['reference'] for r in data['rows']])
  self.assertIn('2026-02-01',[r['date'] for r in data['rows']])
 def test_export_matrix_agrees_and_adp_accounting_is_unchanged(self):
  normal=self.uc.get('/api/accounting?month=2026-01').json();self.assertEqual(normal['ot_minutes'],120);self.assertEqual(normal['guard_minutes'],3270)
  wb=load_workbook(io.BytesIO(self.uc.get('/api/export.xlsx?month=2026-01&view=claim').content))
  self.assertEqual(wb.sheetnames[0],'Claim semanal')
  matrix=list(wb['Claim semanal'].values);self.assertEqual(matrix[0][3:],('Sábado','Domingo','Lunes','Martes','Miércoles','Jueves','Viernes'))
  standby=next(r for r in matrix if r[0]=='2026-01-03' and r[1]=='Stand By');ot=next(r for r in matrix if r[0]=='2026-01-03' and r[1]=='OT')
  self.assertEqual(standby[2],54.5);self.assertEqual(standby[3],6.5);self.assertEqual(ot[3],1.5)
  wb=load_workbook(io.BytesIO(self.uc.get('/api/export.xlsx?month=2026-01').content));self.assertEqual(wb.sheetnames,['Tiempo aprobado','Resumen'])
 def test_scope_and_invalid_view(self):
  other,client,_=self.f.person('USER','CLAIM-PRIVATE',group=self.g)
  self.assertEqual(self.uc.get('/api/accounting?month=2026-01&view=claim&user_id='+other['id']).status_code,403)
  self.assertEqual(self.uc.get('/api/export.xlsx?month=2026-01&view=claim&user_id='+other['id']).status_code,403)
  self.assertEqual(client.get('/api/accounting?month=2026-01&view=claim').json()['rows'],[])
  self.assertEqual(self.uc.get('/api/accounting?month=2026-13&view=claim').status_code,422)
  self.assertEqual(self.uc.get('/api/accounting?month=2026-01&view=unknown').status_code,422)

 def test_entry_totals_sum_rounded_daily_cells_without_changing_actual_minutes(self):
  for day in ('2026-01-17','2026-01-18'):
   ot=self.f.ot(self.uc,start=day+'T10:00:00-06:00',end=day+'T10:01:00-06:00',reference='CHG-MINUTE-'+day);self.f.approve(self.mc,'ot',ot)
  wb=load_workbook(io.BytesIO(self.uc.get('/api/export.xlsx?month=2026-01&view=claim').content))
  entry=next(r for r in list(wb['Claim semanal'].values) if r[0]=='2026-01-17' and r[1]=='OT')
  self.assertEqual(entry[2:5],(0.04,0.02,0.02))
  actual=self.uc.get('/api/accounting?month=2026-01').json();self.assertEqual(actual['ot_minutes'],122)
